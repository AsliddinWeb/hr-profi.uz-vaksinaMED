"""Hisobotlar markazi: hisobot katalogi (har biri o'z ruxsati bilan), Excel/PDF eksport, jadval (cron), yetkazish."""

import asyncio
import io
import re
from datetime import UTC, date, datetime, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.deps import AuthContext, error
from app.models.auth import User
from app.models.employees import Employee
from app.models.org import Branch, Department
from app.models.reports import ReportRun, ReportSchedule
from app.rbac import matches
from app.services.notifications import notify, register_template
from app.services.xlsx_safe import harden
from app.timeutil import local_date, tz

MAX_DAYS = 62
PERIODS = ("yesterday", "last_7d", "last_week", "this_month", "last_month")
FORMATS = ("xlsx", "pdf")
RETENTION_DAYS = 90

register_template("reports.ready", {"uz": "Hisobot tayyor", "ru": "Отчёт готов"}, {"uz": "{name}", "ru": "{name}"}, "info")
register_template("reports.failed", {"uz": "Jadvalli hisobot bajarilmadi", "ru": "Отчёт по расписанию не выполнен"}, {"uz": "{name}", "ru": "{name}"}, "danger")

# type -> (sarlavha uz, ru, ruxsatlardan biri, davr kerakmi)
CATALOG: dict[str, tuple[str, str, tuple[str, ...], bool]] = {
    "attendance": ("Davomat (filiallar kesimida)", "Посещаемость (по филиалам)", ("attendance.read",), True),
    "headcount": ("Xodimlar soni (filial va bo'lim)", "Численность (филиал и отдел)", ("employees.read",), False),
    "tasks": ("Topshiriqlar (ijrochilar kesimida)", "Поручения (по исполнителям)", ("tasks.read", "tasks.manage"), True),
    "needs": ("Ehtiyojlar (kategoriyalar kesimida)", "Потребности (по категориям)", ("needs.dispatch", "needs.manage"), True),
    "tickets": ("AyTi ticketlar (ijrochilar kesimida)", "ИТ-заявки (по исполнителям)", ("helpdesk.read", "helpdesk.manage", "helpdesk.work"), True),
    "visits": ("Mehmonlar (kunlar kesimida)", "Гости (по дням)", ("reception.read", "reception.desk"), True),
}


def allowed(auth_perms: frozenset[str], rtype: str) -> bool:
    return rtype in CATALOG and any(matches(auth_perms, p) for p in CATALOG[rtype][2])


def resolve_period(period: str, today: date) -> tuple[date, date]:
    if period == "yesterday":
        d = today - timedelta(days=1)
        return d, d
    if period == "last_7d":
        return today - timedelta(days=7), today - timedelta(days=1)
    if period == "last_week":
        monday = today - timedelta(days=today.weekday())
        return monday - timedelta(days=7), monday - timedelta(days=1)
    if period == "this_month":
        return today.replace(day=1), today - timedelta(days=1) if today.day > 1 else today
    if period == "last_month":
        first = today.replace(day=1)
        last = first - timedelta(days=1)
        return last.replace(day=1), last
    raise error(400, "bad_period")


def next_run(s: ReportSchedule, after: datetime) -> datetime:
    """`after` dan keyingi ishga tushish (mahalliy vaqt bo'yicha, UTC qaytariladi)."""
    local = after.astimezone(tz())
    cand = local.replace(hour=s.hour, minute=s.minute, second=0, microsecond=0)
    if s.frequency == "daily":
        if cand <= local:
            cand += timedelta(days=1)
    elif s.frequency == "weekly":
        cand += timedelta(days=((s.weekday or 0) - cand.weekday()) % 7)
        if cand <= local:
            cand += timedelta(days=7)
    else:
        day = min(max(s.month_day or 1, 1), 28)
        cand = cand.replace(day=day)
        if cand <= local:
            y, m = (cand.year + 1, 1) if cand.month == 12 else (cand.year, cand.month + 1)
            cand = cand.replace(year=y, month=m)
    return cand.astimezone(UTC)


# --- Hisobot jadvallari -----------------------------------------------------------------


def _pct(v: float | None) -> str:
    return "—" if v is None else f"{v * 100:.1f}%"


def _hours(v: float | None) -> str:
    return "—" if v is None else str(v)


async def build_table(db: AsyncSession, rtype: str, start: date, end: date) -> dict[str, Any]:
    """{'summary': [(uz, ru, qiymat)], 'columns': [(uz, ru)], 'rows': [[...]]}"""
    if rtype == "attendance":
        from app.api.attendance import _dashboard

        d = await _dashboard(db, start, end, q=None, department_id=None, branch_id=None, group_id=None)
        k = d["kpis"]
        return {
            "summary": [("Xodimlar", "Сотрудники", k["employees"]), ("Davomat foizi", "Посещаемость", _pct(k["rate"])), ("Kelmagan kunlar", "Дни неявок", k["absent"]), ("Kechikishlar", "Опоздания", k["late_count"])],
            "columns": [("Filial", "Филиал"), ("Xodimlar", "Сотрудники"), ("Kelgan kunlar", "Дни явки"), ("Kelmagan kunlar", "Дни неявок"), ("Kechikish (marta)", "Опоздания (раз)"), ("Kechikish (daq)", "Опоздания (мин)"), ("Davomat %", "Посещаемость %")],
            "rows": [[r["name"] or "—", r["employees"], r["attended"], r["absent"], r["late_count"], r["late_minutes"], _pct(r["rate"])] for r in d["by_branch"]],
        }  # fmt: skip
    if rtype == "headcount":
        brs = {b.id: (f"№{b.number} {b.name}" if b.number else b.name) for b in (await db.scalars(select(Branch))).all()}
        deps = {x.id: x.name for x in (await db.scalars(select(Department))).all()}
        rows = (await db.execute(select(Employee.branch_id, Employee.department_id, func.count()).where(Employee.status == "active").group_by(Employee.branch_id, Employee.department_id))).all()
        dismissed = await db.scalar(select(func.count()).select_from(Employee).where(Employee.status != "active")) or 0
        total = sum(r[2] for r in rows)
        return {
            "summary": [("Faol xodimlar", "Активные сотрудники", total), ("Nofaol (ishdan ketgan)", "Неактивные (уволены)", dismissed)],
            "columns": [("Filial", "Филиал"), ("Bo'lim", "Отдел"), ("Xodimlar", "Сотрудники")],
            "rows": sorted(([brs.get(r[0] or 0, "—"), deps.get(r[1] or 0, "—"), r[2]] for r in rows), key=lambda x: (x[0], -x[2])),
        }  # fmt: skip
    if rtype == "tasks":
        from app.services import tasks

        d = await tasks.analytics(db, start, end, None)
        k = d["kpis"]
        return {
            "summary": [("Yaratilgan", "Создано", k["created"]), ("Bajarilgan", "Выполнено", k["completed"]), ("Ochiq", "Открыто", k["open"]), ("Muddati o'tgan", "Просрочено", k["overdue"]), ("Muddatida", "В срок", _pct(k["on_time_rate"]))],
            "columns": [("Ijrochi", "Исполнитель"), ("Berilgan", "Назначено"), ("Bajargan", "Выполнено"), ("Ochiq", "Открыто"), ("Muddati o'tgan", "Просрочено"), ("O'rtacha soat", "Среднее, ч")],
            "rows": [[r["user"]["name"] or "—", r["assigned"], r["done"], r["open"], r["overdue"], _hours(r["avg_hours"])] for r in d["by_assignee"]],
        }  # fmt: skip
    if rtype == "needs":
        from app.services import needs

        d = await needs.analytics(db, start, end, None)
        k = d["kpis"]
        return {
            "summary": [("So'ralgan", "Запрошено", k["created"]), ("Bajarilgan", "Выполнено", k["completed"]), ("Ochiq", "Открыто", k["open"]), ("Muddati o'tgan", "Просрочено", k["overdue"]), ("Muddatida", "В срок", _pct(k["on_time_rate"]))],
            "columns": [("Kategoriya", "Категория"), ("So'ralgan", "Запрошено"), ("Bajarilgan", "Выполнено"), ("Ochiq", "Открыто"), ("Muddati o'tgan", "Просрочено"), ("O'rtacha soat", "Среднее, ч"), ("Baho", "Оценка")],
            "rows": [[r["name"] or "—", r["created"], r["done"], r["open"], r["overdue"], _hours(r["avg_hours"]), r["rating"] if r["rating"] is not None else "—"] for r in d["by_category"]],
        }  # fmt: skip
    if rtype == "tickets":
        from app.services import helpdesk

        d = await helpdesk.analytics(db, start, end)
        k = d["kpis"]
        return {
            "summary": [("Ochilgan", "Открыто", k["created"]), ("Hal qilingan", "Решено", k["resolved"]), ("SLA ga ulgurgan", "Уложились в SLA", _pct(k["sla_rate"])), ("Muddati o'tgan", "Просрочено", k["overdue"]), ("O'rtacha hal qilish (soat)", "Среднее решение (ч)", _hours(k["avg_resolution_hours"]))],
            "columns": [("Ijrochi", "Исполнитель"), ("Ochilgan", "Открыто"), ("Hal qilingan", "Решено"), ("Ochiq", "В работе"), ("Muddati o'tgan", "Просрочено"), ("O'rtacha soat", "Среднее, ч"), ("Baho", "Оценка")],
            "rows": [[r["name"] or "—", r["created"], r["done"], r["open"], r["overdue"], _hours(r["avg_hours"]), r["rating"] if r["rating"] is not None else "—"] for r in d["by_executor"]],
        }  # fmt: skip
    if rtype == "visits":
        from app.services import reception

        d = await reception.report(db, start, end, None)
        return {
            "summary": [("Jami tashrif", "Всего визитов", d["total"]), ("Kelmaganlar", "Не пришли", d["no_show"]), ("O'rtacha turish (daq)", "Среднее пребывание (мин)", d["avg_minutes"] if d["avg_minutes"] is not None else "—")],
            "columns": [("Sana", "Дата"), ("Mehmonlar", "Гости")],
            "rows": [[x["date"].strftime("%d.%m.%Y"), x["count"]] for x in d["trend"]],
        }  # fmt: skip
    raise error(400, "unknown_report")


# --- Eksport ----------------------------------------------------------------------------


def _l(pair: tuple[str, str], lang: str) -> str:
    return pair[1] if lang == "ru" else pair[0]


def period_text(rtype: str, start: date, end: date) -> str:
    return f"{start:%d.%m.%Y} – {end:%d.%m.%Y}" if CATALOG[rtype][3] else f"{local_date(datetime.now(UTC)):%d.%m.%Y}"


def build_xlsx(rtype: str, table: dict[str, Any], start: date, end: date, lang: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _l(CATALOG[rtype][:2], lang)[:31].replace("/", "-")
    ws.append([_l(CATALOG[rtype][:2], lang)])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([("Davr" if lang != "ru" else "Период"), period_text(rtype, start, end)])
    for uz, ru, val in table["summary"]:
        ws.append([ru if lang == "ru" else uz, val])
    ws.append([])
    head = ws.max_row + 1
    ws.append([_l(c, lang) for c in table["columns"]])
    for cell in ws[head]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in table["rows"]:
        ws.append(row)
    for i in range(1, len(table["columns"]) + 1):
        width = max((len(str(c.value)) for c in ws[get_column_letter(i)] if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 48)
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    buf = io.BytesIO()
    harden(wb)
    wb.save(buf)
    return buf.getvalue()


def build_pdf(rtype: str, table: dict[str, Any], start: date, end: date, lang: str) -> bytes:
    from app.services.attendance_pdf import _fonts

    regular, bold = _fonts()
    ss = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=ss["Title"], fontName=bold, fontSize=15, alignment=0)
    body = ParagraphStyle("b", parent=ss["Normal"], fontName=regular, fontSize=9, leading=12)
    cell = ParagraphStyle("c", parent=body, fontSize=8, leading=10)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4) if len(table["columns"]) > 4 else A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    story: list[Any] = [Paragraph(_l(CATALOG[rtype][:2], lang), title), Paragraph(f"{'Период' if lang == 'ru' else 'Davr'}: {period_text(rtype, start, end)}", body), Spacer(1, 4 * mm)]
    for uz, ru, val in table["summary"]:
        story.append(Paragraph(f"<b>{ru if lang == 'ru' else uz}:</b> {val}", body))
    story.append(Spacer(1, 4 * mm))
    data = [[Paragraph(f"<b>{_l(c, lang)}</b>", cell) for c in table["columns"]]] + [[Paragraph(str(v), cell) for v in r] for r in table["rows"]]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E0E7FF")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")])]))
    story.append(t)
    doc.build(story)
    return buf.getvalue()


MIME = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "pdf": "application/pdf"}


def file_name(rtype: str, start: date, end: date, fmt: str) -> str:
    base = re.sub(r"[^a-z0-9-]+", "-", f"hisobot-{rtype}-{start}-{end}")
    return f"{base}.{fmt}"


async def generate(db: AsyncSession, storage, *, rtype: str, start: date, end: date, fmt: str, lang: str, user_id: int | None, schedule: ReportSchedule | None = None, recipients: list[int] | None = None, moment: datetime | None = None) -> ReportRun:
    """Hisobotni tuzadi, faylni saqlaydi va ReportRun yozadi (xato bo'lsa status=failed bilan)."""
    if rtype not in CATALOG:
        raise error(400, "unknown_report")
    if fmt not in FORMATS:
        raise error(400, "bad_format")
    if CATALOG[rtype][3]:
        if end < start or (end - start).days + 1 > MAX_DAYS:
            raise error(400, "range_invalid", max_days=MAX_DAYS)
    else:
        start = end = local_date(moment or datetime.now(UTC))
    run = ReportRun(schedule_id=schedule.id if schedule else None, report_type=rtype, start_date=start, end_date=end, format=fmt, lang=lang, status="ok", recipient_ids=recipients or [], created_by=user_id)
    db.add(run)
    await db.flush()
    try:
        table = await build_table(db, rtype, start, end)
        builder = build_xlsx if fmt == "xlsx" else build_pdf
        data = await asyncio.to_thread(builder, rtype, table, start, end, lang)
        run.filename = file_name(rtype, start, end, fmt)
        run.file_key = f"reports/{run.id}/{run.filename}"
        await storage.put(run.file_key, data, MIME[fmt])
        run.file_size, run.row_count = len(data), len(table["rows"])
    except Exception as exc:  # noqa: BLE001 - jadval bo'yicha ishga tushirishda bitta hisobot boshqalarni to'xtatmasin
        if hasattr(exc, "status_code"):
            raise
        run.status, run.error = "failed", f"{type(exc).__name__}: {exc}"[:500]
    return run


def can_download(auth: AuthContext, run: ReportRun) -> bool:
    return matches(auth.permissions, "reports.manage") or run.created_by == auth.user.id or auth.user.id in (run.recipient_ids or [])


async def run_schedule(db: AsyncSession, storage, s: ReportSchedule, moment: datetime):
    """Jadvalni hozir ishga tushiradi: hisobot, bildirishnoma, keyingi vaqt. (run, notifications, telegram_user_ids)"""
    start, end = resolve_period(s.period, local_date(moment))
    owner = await db.get(User, s.created_by) if s.created_by else None
    run = await generate(db, storage, rtype=s.report_type, start=start, end=end, fmt=s.format, lang=s.lang, user_id=s.created_by, schedule=s, recipients=list(s.recipient_ids or []), moment=moment)
    s.last_run_at, s.last_status, s.next_run_at = moment, run.status, next_run(s, moment)
    notes: list = []
    tg: list[int] = []
    if run.status == "ok":
        notes = await notify(db, list(s.recipient_ids or []), "reports.ready", link=f"/reports/history?run={run.id}", params={"name": s.name})
        if "telegram" in (s.channels or []):
            rows = (await db.scalars(select(User.id).where(User.id.in_(list(s.recipient_ids or [])), User.is_active.is_(True), User.telegram_chat_id.is_not(None)))).all()
            tg = list(rows)
    elif owner:
        notes = await notify(db, [owner.id], "reports.failed", link="/reports/schedules", params={"name": s.name})
    return run, notes, tg


async def purge_runs(db: AsyncSession, storage, moment: datetime) -> int:
    """90 kundan eski fayllar va yozuvlar o'chiriladi."""
    old = (await db.scalars(select(ReportRun).where(ReportRun.created_at < moment - timedelta(days=RETENTION_DAYS)))).all()
    for r in old:
        if r.file_key:
            try:
                await storage.delete(r.file_key)
            except Exception:  # noqa: BLE001
                pass
        await db.delete(r)
    return len(old)


async def read_file(storage, key: str) -> bytes:
    return b"".join([chunk async for chunk in storage.stream(key)])
