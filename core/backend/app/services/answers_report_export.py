"""«Javob olish holati» eksporti: Excel (jadval + xodimlar reytingi) va matnli PDF (kartochkalar, tarix bilan)."""

import io
from datetime import date, datetime
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.answers_report import local_dt
from app.services.attendance_pdf import _fonts
from app.services.xlsx_safe import harden

L: dict[str, dict[str, Any]] = {
    "uz": {
        "title": "Javob olish holati", "total": "ta so'rov", "people": "kishi", "page": "sahifa", "month": "Oy", "filters": "Filtr",
        "status": {"pending_coordinator": "KOORDINATORDA", "pending_hr": "HR KUTMOQDA", "approved": "TASDIQLANGAN", "rejected": "RAD ETILGAN", "cancelled": "BEKOR QILINGAN"},
        "status_low": {"pending_coordinator": "Koordinatorda", "pending_hr": "HR kutmoqda", "approved": "Tasdiqlangan", "rejected": "Rad etilgan", "cancelled": "Bekor qilingan"},
        "kind": {"day": "Kunlik", "hourly": "Soatlik"},
        "reason": {"sick": "Kasallik", "family": "Oilaviy", "personal": "Shaxsiy", "business": "Xizmat", "other": "Boshqa"},
        "branch": "Filial", "shift": "Smena", "coordinator": "Koordinator", "note": "Izoh", "history": "Tarix (kim / qachon)",
        "day_unit": "kun", "hour_unit": "soat", "min_unit": "daqiqa",
        "ev": {"submitted": "So'rov yuborildi", "escalated": "Koordinator javob bermadi — HR ga o'tdi", "approved": "tasdiqladi", "rejected": "rad etdi", "cancelled": "So'rov bekor qilindi"},
        "level": {"coordinator": "Koordinator", "hr": "HR"}, "awaiting_hr": "HR yakuniy ruxsati kutilmoqda", "awaiting_coord": "Koordinator qarori kutilmoqda",
        "cols": ["№", "Xodim", "Filial", "Koordinator", "Boshlanish", "Tugash", "Turi", "Muddat", "Sabab turi", "Izoh", "Holat", "Yuborilgan", "Qaror", "Qaror qildi", "Qaror izohi"],
        "rank": "Xodimlar reytingi", "rcols": ["Xodim", "Filial", "Marta", "Kunlik (kun)", "Soatlik (marta)", "Soatlik jami"],
        "truncated": "Ro'yxat qisqartirildi: filtrni toraytiring", "generated": "Tayyorlandi",
    },
    "ru": {
        "title": "Статус заявок на отсутствие", "total": "заявок", "people": "чел.", "page": "стр.", "month": "Месяц", "filters": "Фильтр",
        "status": {"pending_coordinator": "У КООРДИНАТОРА", "pending_hr": "ЖДЁТ HR", "approved": "ОДОБРЕНО", "rejected": "ОТКЛОНЕНО", "cancelled": "ОТМЕНЕНО"},
        "status_low": {"pending_coordinator": "У координатора", "pending_hr": "Ждёт HR", "approved": "Одобрено", "rejected": "Отклонено", "cancelled": "Отменено"},
        "kind": {"day": "Дневная", "hourly": "Почасовая"},
        "reason": {"sick": "Болезнь", "family": "Семейная", "personal": "Личная", "business": "Служебная", "other": "Другая"},
        "branch": "Филиал", "shift": "Смена", "coordinator": "Координатор", "note": "Комментарий", "history": "История (кто / когда)",
        "day_unit": "дн.", "hour_unit": "ч", "min_unit": "мин",
        "ev": {"submitted": "Заявка отправлена", "escalated": "Координатор не ответил — передано HR", "approved": "одобрил(а)", "rejected": "отклонил(а)", "cancelled": "Заявка отменена"},
        "level": {"coordinator": "Координатор", "hr": "HR"}, "awaiting_hr": "Ожидается итоговое решение HR", "awaiting_coord": "Ожидается решение координатора",
        "cols": ["№", "Сотрудник", "Филиал", "Координатор", "Начало", "Конец", "Тип", "Длительность", "Причина", "Комментарий", "Статус", "Отправлена", "Решение", "Решил(а)", "Комментарий к решению"],
        "rank": "Рейтинг сотрудников", "rcols": ["Сотрудник", "Филиал", "Раз", "Дневные (дн.)", "Почасовые (раз)", "Почасовые всего"],
        "truncated": "Список сокращён: сузьте фильтр", "generated": "Сформировано",
    },
}
STAMP = {"pending_coordinator": "#b45309", "pending_hr": "#c2410c", "approved": "#15803d", "rejected": "#b91c1c", "cancelled": "#6b7280"}
FILL = {"pending_coordinator": "#fff7ed", "pending_hr": "#fff7ed", "approved": "#f0fdf4", "rejected": "#fef2f2", "cancelled": "#f3f4f6"}


def _t(lang: str) -> dict[str, Any]:
    return L["ru" if lang == "ru" else "uz"]


def dur_text(it: dict[str, Any], t: dict[str, Any]) -> str:
    if it["minutes"] is not None:
        h, m = divmod(it["minutes"], 60)
        return " ".join(p for p in (f"{h} {t['hour_unit']}" if h else "", f"{m} {t['min_unit']}" if m else "") if p) or f"0 {t['min_unit']}"
    return f"{it['days']} {t['day_unit']}"


def when_text(it: dict[str, Any]) -> str:
    d1 = it["start_date"].strftime("%d.%m.%Y")
    if it["kind"] == "hourly" and it["start_time"]:
        return f"{d1} {it['start_time']}–{it['end_time']}"
    return d1 if it["end_date"] == it["start_date"] else f"{d1} – {it['end_date'].strftime('%d.%m.%Y')}"


def fmt_dt(moment: datetime | None) -> str:
    m = local_dt(moment)
    return m.strftime("%d.%m.%Y, %H:%M") if m else "—"


def hm(minutes: int, t: dict[str, Any]) -> str:
    h, m = divmod(minutes, 60)
    return f"{h} {t['hour_unit']} {m} {t['min_unit']}" if minutes else "—"


def build_xlsx(items: list[dict[str, Any]], summary: dict[str, Any], lang: str) -> bytes:
    t = _t(lang)
    wb = Workbook()
    ws = wb.active
    ws.title = t["title"][:31]
    ws.append(t["cols"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEF0FB")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, it in enumerate(items, 1):
        ws.append([
            i, it["employee"]["full_name"], it["branch"]["name"] if it["branch"] else "", it["coordinator"] or "",
            it["start_date"], it["end_date"], t["kind"][it["kind"]], dur_text(it, t), t["reason"][it["reason_type"]], it["reason"] or "",
            t["status_low"][it["status"]], fmt_dt(it["submitted_at"]), fmt_dt(it["decided_at"]) if it["decided_at"] else "", it["decided_by"] or "", it["decision_note"] or "",
        ])  # fmt: skip
    for col, w in zip("ABCDEFGHIJKLMNO", (5, 28, 22, 22, 12, 12, 10, 14, 12, 30, 16, 18, 18, 22, 30), strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for c in row:
            c.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    rk = wb.create_sheet(t["rank"][:31])
    rk.append(t["rcols"])
    for c in rk[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEF0FB")
    for p in summary["people"]:
        rk.append([p["employee"]["full_name"], p["branch"]["name"] if p["branch"] else "", p["count"], p["days"], p["hourly"], hm(p["hourly_minutes"], t)])
    for col, w in zip("ABCDEF", (30, 24, 8, 14, 16, 18), strict=True):
        rk.column_dimensions[col].width = w
    buf = io.BytesIO()
    harden(wb).save(buf)
    return buf.getvalue()


def _history(it: dict[str, Any], t: dict[str, Any]) -> list[str]:
    lines = []
    for ev in it["events"]:
        a, who, lvl = ev["action"], ev["actor"], ev["level"]
        if a in ("approved", "rejected"):
            text = f"{t['level'].get(lvl or '', '')} {t['ev'][a]}".strip()
        else:
            text = t["ev"].get(a, a)
        tail = f" · {who}" if who else ""
        note = f" ({ev['note']})" if ev.get("note") and a != "submitted" else ""
        lines.append(f"{text}{note} — {fmt_dt(ev['at'])}{tail}")
    if it["status"] == "pending_hr":
        lines.append(f"{t['awaiting_hr']} — —")
    elif it["status"] == "pending_coordinator":
        lines.append(f"{t['awaiting_coord']} — —")
    return lines


def build_pdf(items: list[dict[str, Any]], summary: dict[str, Any], lang: str, filters_text: str, generated: str, truncated: bool = False) -> bytes:
    t = _t(lang)
    regular, bold = _fonts()
    ss = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=ss["Normal"], fontName=regular, fontSize=8.5, leading=11.5)
    h1 = ParagraphStyle("h1", parent=base, fontName=bold, fontSize=17, leading=21)
    name = ParagraphStyle("n", parent=base, fontName=bold, fontSize=11.5, leading=14)
    muted = ParagraphStyle("m", parent=base, textColor=colors.HexColor("#6b7280"), fontSize=8)
    lab = ParagraphStyle("l", parent=base, fontName=bold, fontSize=8.5)
    stamp = ParagraphStyle("s", parent=base, fontName=bold, fontSize=9, alignment=2)
    width = A4[0] - 28 * mm

    def box(title: str, lines: list[str], border: str, bg: str) -> Table:
        paras = [Paragraph(escape(title), lab), *[Paragraph(escape(x), base) for x in lines]]
        tb = Table([[paras]], colWidths=[width - 4 * mm])
        tb.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(border)), ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(bg)),
                                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))  # fmt: skip
        return tb

    story: list[Any] = [Paragraph(t["title"], h1), Spacer(1, 2)]
    k = summary["kpis"]
    story.append(Paragraph(f"{summary['month']} · {len(items)} {t['total']} · {k['people']} {t['people']}" + (f" · {filters_text}" if filters_text else ""), muted))
    story.append(Paragraph(f"{t['generated']}: {generated}", muted))
    if truncated:
        story.append(Paragraph(t["truncated"], ParagraphStyle("w", parent=muted, textColor=colors.HexColor("#b91c1c"))))
    story.append(Spacer(1, 8))
    for it in items:
        col = STAMP[it["status"]]
        head = Table([[Paragraph(escape(it["employee"]["full_name"]), name), Paragraph(t["status"][it["status"]], ParagraphStyle("st", parent=stamp, textColor=colors.HexColor(col)))]], colWidths=[width * 0.68, width * 0.32 - 4 * mm])
        head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))  # fmt: skip
        info = f"{when_text(it)} · {t['kind'][it['kind']]} · {dur_text(it, t)} · {t['reason'][it['reason_type']]}"
        block: list[Any] = [head, Paragraph(escape(info), muted), Spacer(1, 3),
                            box(f"{t['branch']} / {t['shift']} / {t['coordinator']}", [f"{t['branch']}: {it['branch']['name'] if it['branch'] else '—'}", f"{t['shift']}: {it['shift'] or '—'}", f"{t['coordinator']}: {it['coordinator'] or '—'}"], "#c7cbe0", "#f4f5fb")]  # fmt: skip
        if it["reason"]:
            block += [Spacer(1, 3), box(t["note"], [it["reason"]], "#fbbf24", "#fffbeb")]
        block += [Spacer(1, 3), box(t["history"], _history(it, t), "#86efac", "#f0fdf4")]
        card = Table([[block]], colWidths=[width])
        card.setStyle(TableStyle([("LINEBEFORE", (0, 0), (0, -1), 2.5, colors.HexColor(col)), ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))  # fmt: skip
        story += [KeepTogether([card]), Spacer(1, 7)]

    def footer(canvas, doc) -> None:  # noqa: ANN001
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawRightString(A4[0] - 14 * mm, 8 * mm, f"{t['page']} {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title=t["title"]).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def today_text() -> str:
    return date.today().strftime("%d.%m.%Y")
