"""Xodimlarni Excel/CSV dan o'qish va Excelga yozish."""

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services import employees as svc
from app.services.xlsx_safe import harden

HEADERS_UZ = [
    "Familiya", "Ism", "Otasining ismi", "Tug'ilgan sana", "Jinsi", "Telefon", "Email", "Manzil",
    "PINFL", "Pasport", "Bo'lim", "Lavozim", "Filial №", "Ishga olingan sana",
]  # fmt: skip
KEYS = [
    "last_name", "first_name", "middle_name", "birth_date", "gender", "phone", "email", "address",
    "pinfl", "passport_no", "department", "position", "branch", "hire_date",
]  # fmt: skip

ALIASES: dict[str, set[str]] = {
    "full_name": {"f.i.sh", "fish", "f.i.o", "fio", "ф.и.о", "фио", "xodim", "full name", "full_name", "to'liq ism"},
    "last_name": {"familiya", "фамилия", "last name", "last_name", "surname"},
    "first_name": {"ism", "имя", "first name", "first_name", "name"},
    "middle_name": {"otasining ismi", "sharif", "отчество", "middle name", "middle_name"},
    "birth_date": {"tug'ilgan sana", "tugilgan sana", "tug'ilgan", "дата рождения", "birth date", "birth_date", "birthday"},
    "gender": {"jinsi", "jins", "пол", "gender", "sex"},
    "phone": {"telefon", "tel", "телефон", "phone"},
    "email": {"email", "e-mail", "pochta", "почта"},
    "address": {"manzil", "адрес", "address"},
    "pinfl": {"pinfl", "пинфл", "jshshir", "жшшир"},
    "passport_no": {"pasport", "passport", "паспорт", "pasport raqami", "passport_no"},
    "department": {"bo'lim", "bolim", "отдел", "department"},
    "position": {"lavozim", "должность", "position"},
    "branch": {"filial", "filial №", "filial raqami", "филиал", "филиал №", "branch", "dorixona"},
    "hire_date": {"ishga olingan sana", "ishga olingan", "ishga kirgan sana", "дата приёма", "дата приема", "hire date", "hire_date"},
}  # fmt: skip

GENDER_MAP = {
    "m": "m", "male": "m", "erkak": "m", "er": "m", "м": "m", "муж": "m", "мужской": "m",
    "f": "f", "female": "f", "ayol": "f", "ж": "f", "жен": "f", "женский": "f",
}  # fmt: skip

DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y")


def _clean(v: Any) -> Any:
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def parse_date(v: Any) -> date | None:
    v = _clean(v)
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(v, fmt).date()
            except ValueError:
                continue
    raise ValueError("date_invalid")


def _stringify(v: Any) -> str | None:
    v = _clean(v)
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # Excel raqamni float qilib beradi (PINFL, telefon)
    return str(v)


def read_rows(filename: str, data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    """Fayldan qatorlarni o'qiydi: ({kalit: qiymat}, noma'lum ustunlar). Birinchi qator sarlavha."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        raw = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    elif ext in ("csv", "txt"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
        raw = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValueError("file_format")
    raw = [r for r in raw if any(c not in (None, "") for c in r)]
    if len(raw) < 2:
        return [], []

    headers = [str(h or "").strip().lower() for h in raw[0]]
    mapping: dict[int, str] = {}
    unknown: list[str] = []
    for i, h in enumerate(headers):
        key = next((k for k, names in ALIASES.items() if h in names), None)
        if key:
            mapping.setdefault(i, key)
        elif h:
            unknown.append(h)
    rows = []
    for r in raw[1:]:
        rows.append({mapping[i]: r[i] for i in mapping if i < len(r)})
    return rows, unknown


@dataclass
class ImportContext:
    departments: dict[str, list[int]]  # nom(lower) -> idlar
    positions: list[tuple[int, str, int | None]]  # (id, nom lower, bo'lim id)
    branches_by_number: dict[int, int]
    branches_by_name: dict[str, list[int]]


def parse_row(raw: dict[str, Any], ctx: ImportContext) -> tuple[dict[str, Any], list[dict[str, str]]]:
    """Qatorni normallashtiradi. (ma'lumot, xatolar[{field, code}])."""
    errors: list[dict[str, str]] = []
    out: dict[str, Any] = {}

    def fail(field: str, code: str) -> None:
        errors.append({"field": field, "code": code})

    full = _stringify(raw.get("full_name"))
    last, first, middle = _stringify(raw.get("last_name")), _stringify(raw.get("first_name")), _stringify(raw.get("middle_name"))
    if full and not (last and first):
        parts = full.split()
        last = last or (parts[0] if parts else None)
        first = first or (parts[1] if len(parts) > 1 else None)
        middle = middle or (" ".join(parts[2:]) or None)
    if not last or not first:
        fail("name", "name_required")
    out.update(last_name=last, first_name=first, middle_name=middle)

    def norm(field: str, fn, key: str | None = None) -> None:
        value = _stringify(raw.get(field))
        try:
            out[key or field] = fn(value) if value is not None else None
        except ValueError as exc:
            fail(field, str(exc))

    norm("phone", svc.normalize_phone)
    norm("email", svc.normalize_email)
    norm("passport_no", svc.normalize_passport)
    norm("pinfl", svc.normalize_pinfl)
    out["address"] = _stringify(raw.get("address"))

    for field in ("birth_date", "hire_date"):
        try:
            out[field] = parse_date(raw.get(field))
        except ValueError as exc:
            fail(field, str(exc))
    if out.get("birth_date") and (out["birth_date"].year < 1920 or out["birth_date"] > date.today()):
        fail("birth_date", "birth_date_invalid")

    g = _stringify(raw.get("gender"))
    if g:
        out["gender"] = GENDER_MAP.get(g.lower())
        if out["gender"] is None:
            fail("gender", "gender_invalid")
    else:
        out["gender"] = None

    # Bo'lim -> Lavozim -> Filial
    dept_id = None
    dname = _stringify(raw.get("department"))
    if dname:
        ids = ctx.departments.get(dname.lower(), [])
        if len(ids) == 1:
            dept_id = ids[0]
        else:
            fail("department", "department_unknown" if not ids else "department_ambiguous")
    pos_id = None
    pname = _stringify(raw.get("position"))
    if pname:
        cands = [(pid, d) for pid, n, d in ctx.positions if n == pname.lower()]
        if dept_id is not None:
            cands = [c for c in cands if c[1] in (dept_id, None)]
        if len(cands) == 1:
            pos_id = cands[0][0]
            if dept_id is None and cands[0][1] is not None:
                dept_id = cands[0][1]
        else:
            fail("position", "position_unknown" if not cands else "position_ambiguous")
    bid = None
    bval = _stringify(raw.get("branch"))
    if bval:
        num = re.sub(r"^[№#\s]+", "", bval)
        if num.isdigit() and int(num) in ctx.branches_by_number:
            bid = ctx.branches_by_number[int(num)]
        else:
            ids = ctx.branches_by_name.get(bval.lower(), [])
            if len(ids) == 1:
                bid = ids[0]
            else:
                fail("branch", "branch_unknown" if not ids else "branch_ambiguous")
    out.update(department_id=dept_id, position_id=pos_id, branch_id=bid)
    return out, errors


# --- Yozish ------------------------------------------------------------------------------


def _style_header(ws, widths: list[int]) -> None:
    fill = PatternFill("solid", fgColor="1E2A78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28


def template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Xodimlar"
    ws.append(HEADERS_UZ)
    ws.append([
        "Aliyev", "Vali", "Karim o'g'li", "15.03.1990", "Erkak", "+998901234567", "vali@example.uz",
        "Toshkent, Chilonzor", "12345678901234", "AA1234567", "Moliya", "Hisobchi", "16", "01.09.2026",
    ])  # fmt: skip
    _style_header(ws, [16, 14, 16, 16, 10, 16, 22, 26, 18, 14, 18, 18, 10, 18])
    info = wb.create_sheet("Yo'riqnoma")
    for line in (
        "1. Birinchi qator sarlavha, uni o'zgartirmang (ustun tartibi muhim emas, nomi muhim).",
        "2. Majburiy: Familiya va Ism (yoki bitta 'F.I.SH' ustuni).",
        "3. Sana: KK.OO.YYYY yoki YYYY-OO-KK. Telefon: +998 901234567 yoki 901234567.",
        "4. Bo'lim, Lavozim va Filial № tizimda mavjud nomlar bilan bir xil bo'lishi kerak.",
        "5. Tizimda bor xodim (PINFL, pasport, telefon yoki ism+tug'ilgan sana bo'yicha) o'tkazib yuboriladi.",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    harden(wb)
    wb.save(buf)
    return buf.getvalue()


def export_bytes(rows: list[dict[str, Any]], sensitive: bool) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Xodimlar"
    headers = ["№", "F.I.SH", "Jinsi", "Telefon", "Email", "Bo'lim", "Lavozim", "Filial", "Holat",
               "Ishga olingan", "Bo'shatilgan", "Login"]  # fmt: skip
    if sensitive:
        headers += ["Tug'ilgan sana", "PINFL", "Pasport", "Manzil"]
    ws.append(headers)
    for r in rows:
        line = [
            r["employee_no"], r["full_name"], {"m": "Erkak", "f": "Ayol"}.get(r["gender"] or "", ""),
            r["phone"], r["email"], (r["department"] or {}).get("name"), (r["position"] or {}).get("name"),
            (r["branch"] or {}).get("name"), "Ishlayapti" if r["status"] == "active" else "Bo'shatilgan",
            r["hire_date"], r["dismissal_date"], (r["user"] or {}).get("login"),
        ]  # fmt: skip
        if sensitive:
            line += [r["birth_date"], r["pinfl"], r["passport_no"], r["address"]]
        ws.append(line)
    _style_header(ws, [7, 30, 9, 16, 24, 20, 20, 24, 14, 14, 14, 20] + ([14, 18, 12, 30] if sensitive else []))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = "DD.MM.YYYY"
    buf = io.BytesIO()
    harden(wb)
    wb.save(buf)
    return buf.getvalue()
