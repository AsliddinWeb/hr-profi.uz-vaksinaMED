"""Eski tizimdan migratsiya: fayl tahlili, ustun moslashtirish, quruq ishga tushirish (dry-run), import, orqaga qaytarish, solishtirish.

Dry-run va haqiqiy import BIR kod yo'li: dry-run tranzaksiyada hammasini bajaradi (cheklovlar ham tekshiriladi) va ROLLBACK qiladi,
shuning uchun "tekshiruvdan o'tdi" degani "import o'tadi" degani.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.deps import error
from app.models.attendance import AttendanceRecord
from app.models.employees import Employee, EmployeeEvent
from app.models.migration import MigrationBatch, MigrationCheck, MigrationLegacyId
from app.models.org import Branch, Department, Position
from app.services import employee_io as eio
from app.services import employees as emp_svc
from app.timeutil import tz

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = {"employees": 20_000, "attendance": 200_000}
ISSUE_LIMIT = 500
CHUNK = 1000
MAX_SHIFT = timedelta(hours=24)

# dataset -> maydonlar: (kalit, majburiy)
FIELDS: dict[str, list[tuple[str, bool]]] = {
    "employees": [
        ("legacy_id", False), ("full_name", False), ("last_name", False), ("first_name", False), ("middle_name", False), ("birth_date", False), ("gender", False),
        ("phone", False), ("email", False), ("address", False), ("pinfl", False), ("passport_no", False), ("department", False), ("position", False),
        ("branch", False), ("hire_date", False), ("dismissal_date", False), ("dismissal_reason", False),
    ],
    "attendance": [("employee_ref", True), ("work_date", True), ("check_in", True), ("check_out", False), ("late_minutes", False)],
}  # fmt: skip

ALIASES: dict[str, dict[str, set[str]]] = {
    "employees": {
        **eio.ALIASES,
        "legacy_id": {"id", "eski id", "legacy id", "legacy_id", "код", "kod", "tabel raqami", "tabel №", "tabel", "табельный номер", "табельный №", "tab. №"},
        "dismissal_date": {"bo'shatilgan sana", "ishdan ketgan sana", "ishdan bo'shagan", "дата увольнения", "dismissal date", "dismissal_date", "fired"},
        "dismissal_reason": {"bo'shatish sababi", "sabab", "причина увольнения", "dismissal reason", "dismissal_reason"},
    },
    "attendance": {
        "employee_ref": {"xodim id", "xodim", "employee id", "employee", "id", "tabel", "tabel raqami", "табельный номер", "pinfl", "пинфл", "telefon", "phone", "employee_ref"},
        "work_date": {"sana", "kun", "дата", "date", "work_date", "work date"},
        "check_in": {"keldi", "kelgan vaqt", "kirish", "приход", "check in", "check_in", "in", "time in"},
        "check_out": {"ketdi", "ketgan vaqt", "chiqish", "уход", "check out", "check_out", "out", "time out"},
        "late_minutes": {"kechikish", "kechikish (daq)", "опоздание", "late", "late_minutes"},
    },
}  # fmt: skip

REASONS = {
    "own_will": {"o'z xohishi", "ariza", "own_will", "по собственному", "resignation", "resigned"},
    "contract_end": {"shartnoma tugadi", "contract_end", "истёк контракт", "истек контракт"},
    "discipline": {"intizom", "discipline", "дисциплина"},
    "reduction": {"qisqartirish", "reduction", "сокращение"},
}


# --- Fayl o'qish -------------------------------------------------------------------------


def read_table(filename: str, data: bytes) -> tuple[list[str], list[list[Any]]]:
    """(sarlavhalar, qatorlar). Birinchi bo'sh bo'lmagan qator sarlavha; bo'sh qatorlar tashlanadi."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        raw = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    elif ext in ("csv", "txt"):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t") if text.strip() else csv.excel
        except csv.Error:
            dialect = csv.excel
        raw = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValueError("file_format")
    raw = [r for r in raw if any(c not in (None, "") for c in r)]
    if len(raw) < 2:
        return [], []
    headers = [str(h or "").strip() for h in raw[0]]
    return headers, raw[1:]


def suggest_mapping(dataset: str, headers: list[str]) -> dict[str, str]:
    """Sarlavha nomlari bo'yicha avtomatik moslashtirish ({maydon: sarlavha})."""
    out: dict[str, str] = {}
    for key, _ in FIELDS[dataset]:
        names = ALIASES[dataset].get(key, set())
        for h in headers:
            if h and h.lower() in names and h not in out.values():
                out[key] = h
                break
    return out


def check_mapping(dataset: str, mapping: dict[str, str], headers: list[str]) -> None:
    allowed = {k for k, _ in FIELDS[dataset]}
    for k, h in mapping.items():
        if k not in allowed:
            raise error(400, "unknown_field", field=k)
        if h not in headers:
            raise error(400, "unknown_column", column=h)
    for k, req in FIELDS[dataset]:
        if req and k not in mapping:
            raise error(400, "mapping_required", field=k)
    if dataset == "employees" and "full_name" not in mapping and not ("last_name" in mapping and "first_name" in mapping):
        raise error(400, "mapping_required", field="full_name")


def to_dicts(dataset: str, headers: list[str], rows: list[list[Any]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    idx = {k: headers.index(h) for k, h in mapping.items() if h in headers}
    return [{k: (r[i] if i < len(r) else None) for k, i in idx.items()} for r in rows]


def sample(headers: list[str], rows: list[list[Any]], n: int = 5) -> list[list[str]]:
    return [[("" if c is None else str(c)) for c in r[: len(headers)]] for r in rows[:n]]


# --- Natija yig'ish ----------------------------------------------------------------------


@dataclass
class Result:
    rows: int = 0
    created: int = 0
    linked: int = 0
    skipped: int = 0
    errors: int = 0
    new_refs: dict[str, int] = field(default_factory=lambda: {"departments": 0, "positions": 0, "branches": 0})
    ref_ids: dict[str, list[int]] = field(default_factory=lambda: {"departments": [], "positions": [], "branches": []})
    issues: list[dict[str, Any]] = field(default_factory=list)
    stats: dict[str, Any] = field(default_factory=dict)

    def issue(self, row: int, status: str, code: str, field_: str | None = None, name: str | None = None) -> None:
        if status == "error":
            self.errors += 1
        else:
            self.skipped += 1
        if len(self.issues) < ISSUE_LIMIT:
            self.issues.append({"row": row, "status": status, "code": code, "field": field_, "name": name})

    def summary(self) -> dict[str, Any]:
        return {"rows": self.rows, "created": self.created, "linked": self.linked, "skipped": self.skipped, "errors": self.errors, "new_refs": self.new_refs, "ref_ids": self.ref_ids, "stats": self.stats}


def _s(v: Any) -> str | None:
    return eio._stringify(v)  # noqa: SLF001


# --- Xodimlar ----------------------------------------------------------------------------


async def _employee_context(db: AsyncSession) -> eio.ImportContext:
    depts: dict[str, list[int]] = {}
    for d in (await db.scalars(select(Department).where(Department.is_active.is_(True)))).all():
        depts.setdefault(d.name.lower(), []).append(d.id)
    positions = [(p.id, p.name.lower(), p.department_id) for p in (await db.scalars(select(Position).where(Position.is_active.is_(True)))).all()]
    by_num: dict[int, int] = {}
    by_name: dict[str, list[int]] = {}
    for b in (await db.scalars(select(Branch).where(Branch.is_active.is_(True)))).all():
        if b.number is not None:
            by_num[b.number] = b.id
        by_name.setdefault(b.name.lower(), []).append(b.id)
    return eio.ImportContext(depts, positions, by_num, by_name)


async def _create_missing_refs(db: AsyncSession, rows: list[dict[str, Any]], res: Result) -> None:
    """Fayldagi, tizimda yo'q bo'lim/lavozim/filiallarni yaratadi (dry-run da ham: keyin rollback)."""
    ctx = await _employee_context(db)
    for raw in rows:
        dname = _s(raw.get("department"))
        if dname and dname.lower() not in ctx.departments:
            d = Department(name=dname[:120], kind="office")
            db.add(d)
            await db.flush()
            ctx.departments[dname.lower()] = [d.id]
            res.new_refs["departments"] += 1
            res.ref_ids["departments"].append(d.id)
        pname = _s(raw.get("position"))
        if pname and not any(n == pname.lower() for _, n, _d in ctx.positions):
            dept_id = (ctx.departments.get(dname.lower()) or [None])[0] if dname else None
            p = Position(name=pname[:120], department_id=dept_id)
            db.add(p)
            await db.flush()
            ctx.positions.append((p.id, pname.lower(), dept_id))
            res.new_refs["positions"] += 1
            res.ref_ids["positions"].append(p.id)
        bval = _s(raw.get("branch"))
        if bval:
            num = re.sub(r"^[№#\s]+", "", bval)
            known = (num.isdigit() and int(num) in ctx.branches_by_number) or bval.lower() in ctx.branches_by_name
            if not known:
                b = Branch(name=bval[:160], number=int(num) if num.isdigit() and int(num) < 10**6 else None)
                if num.isdigit():
                    b.name = f"№{num}"
                db.add(b)
                await db.flush()
                if b.number is not None:
                    ctx.branches_by_number[b.number] = b.id
                ctx.branches_by_name.setdefault(b.name.lower(), []).append(b.id)
                res.new_refs["branches"] += 1
                res.ref_ids["branches"].append(b.id)


def _reason(v: Any) -> str:
    t = (_s(v) or "").lower()
    return next((k for k, names in REASONS.items() if t in names), "other")


async def process_employees(db: AsyncSession, batch: MigrationBatch, rows: list[dict[str, Any]], options: dict[str, Any], actor_id: int | None) -> Result:
    res = Result(rows=len(rows))
    if options.get("create_missing"):
        await _create_missing_refs(db, rows, res)
    ctx = await _employee_context(db)
    legacy = {r.legacy_id: r.entity_id for r in (await db.scalars(select(MigrationLegacyId).where(MigrationLegacyId.entity == "employee"))).all()}
    seen: dict[tuple, int] = {}
    seen_legacy: set[str] = set()
    for i, raw in enumerate(rows, start=2):  # fayldagi qator raqami (1 = sarlavha)
        data, errors = eio.parse_row(raw, ctx)
        name = emp_svc.build_full_name(data.get("last_name") or "", data.get("first_name") or "", None) if data.get("last_name") else None
        lid = _s(raw.get("legacy_id"))
        dismissed = None
        if raw.get("dismissal_date") not in (None, ""):
            try:
                dismissed = eio.parse_date(raw.get("dismissal_date"))
            except ValueError:
                errors.append({"field": "dismissal_date", "code": "date_invalid"})
        if dismissed and data.get("hire_date") and dismissed < data["hire_date"]:
            errors.append({"field": "dismissal_date", "code": "date_before_hire"})
        if errors:
            res.issue(i, "error", errors[0]["code"], errors[0]["field"], name)
            continue
        if lid and (lid in legacy or lid in seen_legacy):
            res.issue(i, "skipped", "legacy_exists", "legacy_id", name)
            continue
        full = emp_svc.build_full_name(data["last_name"], data["first_name"], data.get("middle_name"))
        keys = [("pinfl", data["pinfl"]), ("passport", data["passport_no"]), ("phone", data["phone"])]
        if data.get("birth_date"):
            keys.append(("nb", (emp_svc.name_key(full), data["birth_date"])))
        in_file = next((seen[k] for k in keys if k[1] and k in seen), None)
        for k in keys:
            if k[1]:
                seen.setdefault(k, i)
        if in_file is not None:
            res.issue(i, "skipped", "duplicate_in_file", None, name)
            continue
        matches = await emp_svc.find_duplicates(db, pinfl=data["pinfl"], passport_no=data["passport_no"], phone=data["phone"], full_name=full, birth_date=data.get("birth_date"))
        if matches:
            if lid:  # mavjud xodimga bog'lab qo'yamiz: davomat importi shu orqali topadi
                db.add(MigrationLegacyId(entity="employee", legacy_id=lid[:60], entity_id=matches[0].employee.id, batch_id=batch.id))
                seen_legacy.add(lid)
                res.linked += 1
            res.issue(i, "skipped", "duplicate_existing", None, name)
            continue
        try:
            await emp_svc.validate_assignment(db, data["department_id"], data["position_id"], data["branch_id"])
        except Exception as exc:  # noqa: BLE001
            code = getattr(exc, "detail", {}).get("code", "assignment_invalid") if hasattr(exc, "detail") and isinstance(exc.detail, dict) else "assignment_invalid"
            res.issue(i, "error", code, "department", name)
            continue
        hire = data.get("hire_date")
        emp = Employee(
            last_name=data["last_name"], first_name=data["first_name"], middle_name=data.get("middle_name"), full_name=full, birth_date=data.get("birth_date"),
            gender=data.get("gender"), phone=data["phone"], email=data["email"], address=data["address"], passport_no=data["passport_no"], pinfl=data["pinfl"],
            department_id=data["department_id"], position_id=data["position_id"], branch_id=data["branch_id"], hire_date=hire, status="active", import_batch_id=batch.id,
        )  # fmt: skip
        if dismissed:
            emp.status, emp.dismissal_date, emp.dismissal_reason = "dismissed", dismissed, _reason(raw.get("dismissal_reason"))
        db.add(emp)
        await db.flush()
        emp_svc.add_event(db, emp.id, "hired", hire or dismissed or date.today(), data={"source": "migration"}, actor_id=actor_id)
        if dismissed:
            emp_svc.add_event(db, emp.id, "dismissed", dismissed, data={"source": "migration", "reason": emp.dismissal_reason}, actor_id=actor_id)
        if lid:
            db.add(MigrationLegacyId(entity="employee", legacy_id=lid[:60], entity_id=emp.id, batch_id=batch.id))
            seen_legacy.add(lid)
        res.created += 1
    res.stats = {"dismissed": sum(1 for r in rows if r.get("dismissal_date") not in (None, ""))}
    return res


# --- Davomat -----------------------------------------------------------------------------


def parse_moment(v: Any, base: date) -> tuple[datetime | None, bool]:
    """(tz-aware vaqt, faqat_soat). Datetime, time, 'HH:MM[:SS]', 'DD.MM.YYYY HH:MM' va 'YYYY-MM-DD HH:MM' qabul qilinadi."""
    if v in (None, ""):
        return None, False
    if isinstance(v, datetime):
        return (v if v.tzinfo else v.replace(tzinfo=tz())), False
    if isinstance(v, time):
        return datetime.combine(base, v, tzinfo=tz()), True
    if isinstance(v, (int, float)) and 0 <= float(v) < 1:  # Excel kasr sutka
        secs = round(float(v) * 86400)
        return datetime.combine(base, time(secs // 3600 % 24, secs % 3600 // 60), tzinfo=tz()), True
    t = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})[:.](\d{2})(?::(\d{2}))?", t)
    if m:
        h, mi = int(m.group(1)), int(m.group(2))
        if h > 23 or mi > 59:
            raise ValueError("time_invalid")
        return datetime.combine(base, time(h, mi), tzinfo=tz()), True
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%d/%m/%Y %H:%M"):
        try:
            return datetime.strptime(t, fmt).replace(tzinfo=tz()), False
        except ValueError:
            continue
    raise ValueError("time_invalid")


async def _employee_lookup(db: AsyncSession, match_by: str) -> dict[str, int]:
    if match_by == "legacy_id":
        return {r.legacy_id: r.entity_id for r in (await db.scalars(select(MigrationLegacyId).where(MigrationLegacyId.entity == "employee"))).all()}
    col = {"employee_no": Employee.employee_no, "pinfl": Employee.pinfl, "phone": Employee.phone}[match_by]
    out: dict[str, int] = {}
    for eid, val in (await db.execute(select(Employee.id, col).where(col.is_not(None)))).all():
        out[str(val)] = eid
    return out


def _norm_ref(match_by: str, v: Any) -> str | None:
    s = _s(v)
    if s is None:
        return None
    if match_by == "phone":
        try:
            return emp_svc.normalize_phone(s)
        except ValueError:
            return None
    if match_by == "pinfl":
        digits = re.sub(r"\D", "", s)
        return digits or None
    return s


async def process_attendance(db: AsyncSession, batch: MigrationBatch, rows: list[dict[str, Any]], options: dict[str, Any], actor_id: int | None) -> Result:
    match_by = options.get("match_by", "legacy_id")
    if match_by not in ("legacy_id", "employee_no", "pinfl", "phone"):
        raise error(400, "bad_match_by")
    res = Result(rows=len(rows))
    lookup = await _employee_lookup(db, match_by)
    parsed: list[tuple[int, int, date, datetime, datetime | None, int]] = []
    dates: set[date] = set()
    emp_ids: set[int] = set()
    for i, raw in enumerate(rows, start=2):
        ref = _norm_ref(match_by, raw.get("employee_ref"))
        eid = lookup.get(ref) if ref else None
        if eid is None:
            res.issue(i, "error", "employee_unknown", "employee_ref", ref)
            continue
        try:
            wd = eio.parse_date(raw.get("work_date"))
            if wd is None:
                raise ValueError("date_invalid")
            cin, in_time_only = parse_moment(raw.get("check_in"), wd)
            cout, _ = parse_moment(raw.get("check_out"), wd)
            if cin is None:
                raise ValueError("time_invalid")
            if cout is not None and cout <= cin and in_time_only:
                cout += timedelta(days=1)  # tungi smena
            if cout is not None and (cout <= cin or cout - cin > MAX_SHIFT):
                raise ValueError("shift_invalid")
            late = int(float(_s(raw.get("late_minutes")) or 0))
            if late < 0:
                raise ValueError("late_invalid")
        except ValueError as exc:
            res.issue(i, "error", str(exc) if str(exc) in ("date_invalid", "time_invalid", "shift_invalid", "late_invalid") else "date_invalid", None, ref)
            continue
        parsed.append((i, eid, wd, cin, cout, late))
        dates.add(wd)
        emp_ids.add(eid)
    existing: set[tuple[int, date]] = set()
    if parsed:
        existing = set((await db.execute(select(AttendanceRecord.employee_id, AttendanceRecord.work_date).where(AttendanceRecord.employee_id.in_(emp_ids), AttendanceRecord.work_date >= min(dates), AttendanceRecord.work_date <= max(dates)))).all())
    pending: list[AttendanceRecord] = []
    minutes = 0
    for i, eid, wd, cin, cout, late in parsed:
        if (eid, wd) in existing:
            res.issue(i, "skipped", "attendance_exists", None, None)
            continue
        existing.add((eid, wd))
        worked = int((cout - cin).total_seconds() // 60) if cout else 0
        minutes += worked
        pending.append(AttendanceRecord(
            employee_id=eid, work_date=wd, check_in_at=cin, check_out_at=cout, in_method="import", out_method="import" if cout else None, track_lateness=False,
            late_minutes=late, worked_minutes=worked, is_manual=True, note="migration", created_by=actor_id, import_batch_id=batch.id,
        ))  # fmt: skip
        if len(pending) >= CHUNK:
            db.add_all(pending)
            await db.flush()
            pending = []
    if pending:
        db.add_all(pending)
        await db.flush()
    in_db = (await db.execute(select(func.count(), func.coalesce(func.sum(AttendanceRecord.worked_minutes), 0)).where(AttendanceRecord.import_batch_id == batch.id))).one()
    res.created = in_db[0]
    # Solishtirish: faylga ko'ra yoziladigan soat == bazaga yozilgan soat
    res.stats = {"expected_minutes": minutes, "db_minutes": int(in_db[1]), "reconciled": minutes == int(in_db[1]), "employees": len(emp_ids), "date_from": min(dates).isoformat() if dates else None, "date_to": max(dates).isoformat() if dates else None}
    return res


PROCESSORS = {"employees": process_employees, "attendance": process_attendance}


async def run(db: AsyncSession, batch: MigrationBatch, headers: list[str], raw_rows: list[list[Any]], *, commit: bool, actor_id: int | None) -> Result:
    """Qatorlarni qayta ishlaydi. commit=False: dry-run (hammasi bajariladi va rollback qilinadi)."""
    rows = to_dicts(batch.dataset, headers, raw_rows, batch.mapping)
    try:
        res = await PROCESSORS[batch.dataset](db, batch, rows, batch.options or {}, actor_id)
    except IntegrityError as exc:
        await db.rollback()
        raise error(409, "import_conflict", detail=str(exc.orig)[:200]) from None
    if not commit:
        await db.rollback()
        await db.refresh(batch)
    return res


async def rollback_batch(db: AsyncSession, batch: MigrationBatch) -> dict[str, int]:
    """Paket yaratgan yozuvlarni o'chiradi. Keyin boshqa joyda ishlatilgan bo'lsa (hisob, boshqa davomat) 409."""
    att_other = await db.scalar(
        select(func.count()).select_from(AttendanceRecord).join(Employee, Employee.id == AttendanceRecord.employee_id)
        .where(Employee.import_batch_id == batch.id, AttendanceRecord.import_batch_id.is_distinct_from(batch.id))
    )  # fmt: skip
    users = await db.scalar(select(func.count()).select_from(Employee).where(Employee.import_batch_id == batch.id, Employee.user_id.is_not(None)))
    if att_other or users:
        raise error(409, "batch_in_use", attendance=att_other or 0, accounts=users or 0)
    removed_att = (await db.execute(delete(AttendanceRecord).where(AttendanceRecord.import_batch_id == batch.id))).rowcount
    emp_ids = list((await db.scalars(select(Employee.id).where(Employee.import_batch_id == batch.id))).all())
    try:
        if emp_ids:
            await db.execute(delete(EmployeeEvent).where(EmployeeEvent.employee_id.in_(emp_ids)))
        removed_emp = (await db.execute(delete(Employee).where(Employee.import_batch_id == batch.id))).rowcount
        await db.execute(delete(MigrationLegacyId).where(MigrationLegacyId.batch_id == batch.id))
        await db.flush()
    except IntegrityError:
        await db.rollback()
        raise error(409, "batch_in_use") from None
    refs = 0
    ids = (batch.summary or {}).get("ref_ids") or {}
    for model, key in ((Position, "positions"), (Branch, "branches"), (Department, "departments")):
        for rid in ids.get(key, []):
            try:  # boshqa joyda ishlatilayotgan ma'lumotnomani qoldiramiz
                async with db.begin_nested():
                    await db.execute(delete(model).where(model.id == rid))
                refs += 1
            except IntegrityError:
                continue
    return {"attendance": removed_att or 0, "employees": removed_emp or 0, "references": refs}


# --- Solishtirish (reconcile) --------------------------------------------------------------

METRICS = ("employees_total", "employees_active", "employees_dismissed", "attendance_records", "attendance_hours")
PERIOD_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")


def month_bounds(period: str) -> tuple[date, date]:
    y, m = int(period[:4]), int(period[5:])
    first = date(y, m, 1)
    nxt = date(y + (m == 12), m % 12 + 1, 1)
    return first, nxt - timedelta(days=1)


async def actual_value(db: AsyncSession, c: MigrationCheck) -> Decimal:
    if c.metric.startswith("employees"):
        q = select(func.count()).select_from(Employee)
        if c.metric == "employees_active":
            q = q.where(Employee.status == "active")
        elif c.metric == "employees_dismissed":
            q = q.where(Employee.status == "dismissed")
        if c.branch_id:
            q = q.where(Employee.branch_id == c.branch_id)
        return Decimal(await db.scalar(q) or 0)
    if not c.period:
        raise error(400, "period_required")
    start, end = month_bounds(c.period)
    expr = func.count() if c.metric == "attendance_records" else func.coalesce(func.sum(AttendanceRecord.worked_minutes), 0)
    q = select(expr).select_from(AttendanceRecord).where(AttendanceRecord.work_date >= start, AttendanceRecord.work_date <= end)
    if c.branch_id:
        q = q.join(Employee, Employee.id == AttendanceRecord.employee_id).where(Employee.branch_id == c.branch_id)
    v = Decimal(await db.scalar(q) or 0)
    return (v / 60).quantize(Decimal("0.01")) if c.metric == "attendance_hours" else v


async def reconcile(db: AsyncSession) -> dict[str, Any]:
    checks = (await db.scalars(select(MigrationCheck).order_by(MigrationCheck.id))).all()
    branches = {b.id: b.name for b in (await db.scalars(select(Branch))).all()}
    items = []
    for c in checks:
        actual = await actual_value(db, c)
        diff = (actual - c.expected).quantize(Decimal("0.01"))
        items.append({
            "id": c.id, "label": c.label, "metric": c.metric, "period": c.period, "branch": {"id": c.branch_id, "name": branches.get(c.branch_id or 0)} if c.branch_id else None,
            "expected": float(c.expected), "actual": float(actual), "diff": float(diff), "tolerance": float(c.tolerance), "status": "match" if abs(diff) <= c.tolerance else "mismatch",
        })  # fmt: skip
    return {"total": len(items), "matched": sum(1 for x in items if x["status"] == "match"), "mismatched": sum(1 for x in items if x["status"] == "mismatch"), "items": items}


def parse_checks_csv(data: bytes) -> list[dict[str, Any]]:
    """CSV/Excel: label, metric, period, branch (raqam), expected, tolerance."""
    headers, rows = read_table("c.csv", data)
    h = [x.lower() for x in headers]
    need = {"metric", "expected"}
    if not need <= set(h):
        raise error(400, "columns_missing", columns=sorted(need - set(h)))
    out = []
    for i, r in enumerate(rows, start=2):
        d = {h[j]: (r[j] if j < len(r) else None) for j in range(len(h))}
        metric = (_s(d.get("metric")) or "").lower()
        if metric not in METRICS:
            raise error(400, "metric_invalid", row=i)
        period = _s(d.get("period"))
        if metric.startswith("attendance") and not (period and PERIOD_RE.match(period)):
            raise error(400, "period_invalid", row=i)
        try:
            expected = Decimal(str(_s(d.get("expected")) or "").replace(",", ".").replace(" ", ""))
            tol = Decimal(str(_s(d.get("tolerance")) or "0").replace(",", "."))
        except Exception:  # noqa: BLE001
            raise error(400, "number_invalid", row=i) from None
        out.append({"label": (_s(d.get("label")) or metric)[:160], "metric": metric, "period": period if metric.startswith("attendance") else None, "branch": _s(d.get("branch")), "expected": expected, "tolerance": tol})
    return out

