"""Jarimalar: qoidalar (versiyalangan), avtomatik hisob, tasdiq/bekor, hisobot."""

import asyncio
import io
from datetime import date
from decimal import Decimal
from typing import Any

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import case, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.attendance import XLSX
from app.api.employees import get_employee
from app.db import get_session
from app.deps import AuthContext, error, get_auth
from app.models.auth import User
from app.models.employees import Employee
from app.models.letters import ExplanationLetter
from app.models.org import Branch, Department
from app.models.penalties import Penalty, PenaltyRuleVersion
from app.services import penalties as svc
from app.services.audit import audit
from app.services.letters import has
from app.services.notifications import deliver
from app.services.xlsx_safe import harden
from app.timeutil import local_date, now

router = APIRouter(prefix="/v1/penalties", tags=["penalties"])


def require_any(*perms: str):
    async def dep(auth: AuthContext = Depends(get_auth)) -> AuthContext:
        if not any(has(auth, p) for p in perms):
            raise error(403, "forbidden")
        return auth

    return dep


def can_read_all(auth: AuthContext) -> bool:
    return has(auth, "penalties.read") or has(auth, "penalties.approve") or has(auth, "penalties.manage")


# --- Qoidalar ---------------------------------------------------------------------------


def rules_out(v: PenaltyRuleVersion, author: str | None = None) -> dict[str, Any]:
    return {"id": v.id, "version": v.version, "config": v.config, "note": v.note, "created_at": v.created_at, "author": author}


@router.get("/rules")
async def rules(db: AsyncSession = Depends(get_session), _: AuthContext = Depends(require_any("penalties.read", "penalties.approve", "penalties.manage"))):
    versions = (await db.scalars(select(PenaltyRuleVersion).order_by(PenaltyRuleVersion.version.desc()))).all()
    names = {u.id: u.full_name for u in (await db.scalars(select(User).where(User.id.in_({v.created_by for v in versions if v.created_by})))).all()} if versions else {}
    used = dict((await db.execute(select(Penalty.rule_version_id, func.count()).group_by(Penalty.rule_version_id))).all())
    return {
        "current": rules_out(versions[0], names.get(versions[0].created_by or 0)) if versions else None,
        "versions": [{**rules_out(v, names.get(v.created_by or 0)), "used": used.get(v.id, 0)} for v in versions],
    }


@router.get("/rules/{rule_id}")
async def rule_one(rule_id: int, db: AsyncSession = Depends(get_session), _: AuthContext = Depends(require_any("penalties.read", "penalties.approve", "penalties.manage"))):
    v = await db.get(PenaltyRuleVersion, rule_id)
    if not v:
        raise error(404, "not_found")
    author = await db.get(User, v.created_by) if v.created_by else None
    return rules_out(v, author.full_name if author else None)


class RulesIn(BaseModel):
    config: svc.RulesConfig
    note: str | None = Field(default=None, max_length=500)


@router.post("/rules", status_code=201)
async def new_rules(body: RulesIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(require_any("penalties.manage"))):
    """Qoidani o'zgartirish = yangi versiya. Avval hisoblangan jarimalar o'z versiyasi bo'yicha qoladi."""
    last = await db.scalar(select(func.max(PenaltyRuleVersion.version))) or 0
    v = PenaltyRuleVersion(version=last + 1, config=svc.config_json(body.config), note=(body.note or "").strip() or None, created_by=auth.user.id)
    db.add(v)
    await db.flush()
    await audit(db, request, "penalties.rules", actor=auth.user, entity="penalty_rules", entity_id=v.id, data={"version": v.version}, note=v.note)
    await db.commit()
    return rules_out(v, auth.user.full_name)


# --- Avtomatik hisob --------------------------------------------------------------------


class ScanIn(BaseModel):
    start: date
    end: date
    dry_run: bool = True


@router.post("/scan")
async def scan(body: ScanIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(require_any("penalties.manage"))):
    """O'tgan kunlar uchun davomatdan jarimalarni hisoblaydi (tasdiq kutadi). `dry_run` faqat sanaydi."""
    result = await svc.scan_range(db, body.start, body.end, dry_run=body.dry_run, moment=now())
    if not body.dry_run:
        await audit(db, request, "penalties.scan", actor=auth.user, entity="penalty", data={"start": str(body.start), "end": str(body.end), **{k: v for k, v in result.items() if k != "dry_run"}})
        await db.commit()
    return result


# --- Ro'yxat va qaror -------------------------------------------------------------------


@router.get("")
async def list_penalties(
    scope: str = Query(default="mine", pattern="^(mine|all)$"),
    status: str | None = Query(default=None, pattern="^(pending|approved|cancelled)$"),
    incident_type: str | None = Query(default=None, pattern="^(late|absent|early_leave|no_checkout|discipline|other)$"),
    source: str | None = Query(default=None, pattern="^(auto|manual)$"),
    q: str | None = None,
    employee_id: int | None = None,
    start: date | None = None,
    end: date | None = None,
    page: int = Query(1, ge=1),
    size: int = Query(30, ge=1, le=100),
    db: AsyncSession = Depends(get_session),
    auth: AuthContext = Depends(get_auth),
):
    conds: list[Any] = []
    if scope == "mine":
        conds.append(Employee.user_id == auth.user.id)
    elif not can_read_all(auth):
        raise error(403, "forbidden")
    for col, val in ((Penalty.status, status), (Penalty.incident_type, incident_type), (Penalty.source, source), (Penalty.employee_id, employee_id)):
        if val:
            conds.append(col == val)
    if q and q.strip():
        conds.append(Employee.full_name.ilike(f"%{q.strip()}%"))
    if start:
        conds.append(Penalty.incident_date >= start)
    if end:
        conds.append(Penalty.incident_date <= end)
    base = select(Penalty, Employee).join(Employee, Employee.id == Penalty.employee_id).where(*conds)
    total = await db.scalar(select(func.count()).select_from(Penalty).join(Employee, Employee.id == Penalty.employee_id).where(*conds))
    sums = float(await db.scalar(select(func.coalesce(func.sum(Penalty.amount), 0)).select_from(Penalty).join(Employee, Employee.id == Penalty.employee_id).where(*conds, Penalty.status == "approved")) or 0)
    stmt = base
    rows = (await db.execute(stmt.order_by(Penalty.incident_date.desc(), Penalty.id.desc()).offset((page - 1) * size).limit(size))).all()
    return {"total": total or 0, "page": page, "size": size, "items": [svc.penalty_out(p, e) for p, e in rows], "approved_amount": sums}


@router.get("/summary")
async def summary(db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(get_auth)) -> dict[str, int]:
    """Menyu belgisi: tasdiq kutayotgan jarimalar soni (tasdiqlovchilar uchun)."""
    if not has(auth, "penalties.approve"):
        return {"pending": 0}
    return {"pending": await db.scalar(select(func.count()).select_from(Penalty).where(Penalty.status == "pending")) or 0}


class ManualIn(BaseModel):
    employee_id: int
    incident_type: str = Field(pattern="^(late|absent|early_leave|no_checkout|discipline|other)$")
    incident_date: date
    step: int | None = Field(default=None, ge=1, le=svc.STEPS)
    amount: Decimal | None = Field(default=None, ge=0, le=Decimal("1000000000"))
    block_days: int | None = Field(default=None, ge=0, le=90)
    label: str | None = Field(default=None, max_length=120)
    reason: str = Field(min_length=3, max_length=2000)


@router.post("", status_code=201)
async def create_manual(body: ManualIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(require_any("penalties.manage"))):
    emp = await get_employee(db, body.employee_id)
    if emp.status != "active":
        raise error(409, "employee_dismissed")
    if body.incident_date > local_date(now()):
        raise error(400, "date_in_future")
    p = await svc.create_manual(db, auth, emp, incident_type=body.incident_type, incident_date=body.incident_date, step=body.step, amount=body.amount, block_days=body.block_days, label=body.label, reason=body.reason)
    await audit(db, request, "penalties.create", actor=auth.user, entity="penalty", entity_id=p.id, data={"employee_id": emp.id, "amount": float(p.amount)}, note=body.reason)
    await db.commit()
    return svc.penalty_out(p, emp)


async def _load(db: AsyncSession, penalty_id: int, auth: AuthContext, *, lock: bool = False) -> tuple[Penalty, Employee]:
    stmt = select(Penalty).where(Penalty.id == penalty_id)
    p = await db.scalar(stmt.with_for_update() if lock else stmt)
    if not p:
        raise error(404, "not_found")
    emp = await db.get(Employee, p.employee_id)
    if emp is None or not (emp.user_id == auth.user.id or can_read_all(auth)):
        raise error(404, "not_found")
    return p, emp


@router.get("/report")
async def report(
    start: date,
    end: date,
    group_by: str = Query(default="employee", pattern="^(employee|department|branch|type)$"),
    db: AsyncSession = Depends(get_session),
    _: AuthContext = Depends(require_any("penalties.read", "penalties.approve", "penalties.manage")),
):
    """Jarimalar hisoboti. Tasdiqlangan summalar alohida, kutilayotganlar alohida ko'rsatiladi."""
    if end < start or (end - start).days > 366:
        raise error(400, "range_invalid")
    return await _report(db, start, end, group_by)


async def _report(db: AsyncSession, start: date, end: date, group_by: str) -> dict[str, Any]:
    if group_by == "employee":
        key, name = Employee.id, Employee.full_name
    elif group_by == "department":
        key, name = Department.id, func.coalesce(Department.name, "")
    elif group_by == "branch":
        key, name = Branch.id, func.coalesce(Branch.name, "")
    else:
        key, name = Penalty.incident_type, Penalty.incident_type
    approved = Penalty.status == "approved"
    stmt = (
        select(
            key.label("k"), name.label("n"),
            func.count().filter(Penalty.status == "pending").label("pending"),
            func.count().filter(approved).label("approved"),
            func.count().filter(Penalty.status == "cancelled").label("cancelled"),
            func.coalesce(func.sum(case((approved, Penalty.amount), else_=0)), 0).label("amount"),
            func.coalesce(func.sum(case((Penalty.status == "pending", Penalty.amount), else_=0)), 0).label("pending_amount"),
            func.count().filter(approved, Penalty.block_days > 0).label("blocks"),
        )
        .select_from(Penalty)
        .join(Employee, Employee.id == Penalty.employee_id)
        .where(Penalty.incident_date.between(start, end))
        .group_by(key, name)
    )  # fmt: skip
    if group_by == "department":
        stmt = stmt.outerjoin(Department, Department.id == Employee.department_id)
    elif group_by == "branch":
        stmt = stmt.outerjoin(Branch, Branch.id == Employee.branch_id)
    rows = (await db.execute(stmt.order_by(func.sum(case((approved, Penalty.amount), else_=0)).desc(), name))).all()
    items = [
        {"key": r.k, "name": r.n or None, "pending": r.pending, "approved": r.approved, "cancelled": r.cancelled,
         "amount": float(r.amount), "pending_amount": float(r.pending_amount), "blocks": r.blocks}
        for r in rows
    ]  # fmt: skip
    totals = {k: sum(i[k] for i in items) for k in ("pending", "approved", "cancelled", "amount", "pending_amount", "blocks")}
    return {"start": start, "end": end, "group_by": group_by, "items": items, "totals": totals}


@router.get("/report/export")
async def report_export(
    request: Request,
    start: date,
    end: date,
    group_by: str = Query(default="employee", pattern="^(employee|department|branch|type)$"),
    db: AsyncSession = Depends(get_session),
    auth: AuthContext = Depends(require_any("penalties.read", "penalties.approve", "penalties.manage")),
):
    if end < start or (end - start).days > 366:
        raise error(400, "range_invalid")
    data = await _report(db, start, end, group_by)
    content = await asyncio.to_thread(_xlsx, data)
    await audit(db, request, "penalties.export", actor=auth.user, entity="penalty", data={"start": str(start), "end": str(end), "group_by": group_by})
    await db.commit()
    return Response(content, media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="jarimalar-{start}-{end}.xlsx"'})


def _xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Jarimalar"
    ws.append(["Nomi", "Tasdiqlangan (soni)", "Tasdiqlangan summa", "Kutilayotgan (soni)", "Kutilayotgan summa", "Bekor qilingan", "Blok"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i in data["items"]:
        ws.append([i["name"] or "—", i["approved"], i["amount"], i["pending"], i["pending_amount"], i["cancelled"], i["blocks"]])
    t = data["totals"]
    ws.append(["JAMI", t["approved"], t["amount"], t["pending"], t["pending_amount"], t["cancelled"], t["blocks"]])
    ws[ws.max_row][0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 36
    buf = io.BytesIO()
    harden(wb)
    wb.save(buf)
    return buf.getvalue()


@router.get("/{penalty_id}")
async def detail(penalty_id: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(get_auth)):
    p, emp = await _load(db, penalty_id, auth)
    version = await db.get(PenaltyRuleVersion, p.rule_version_id) if p.rule_version_id else None
    out = svc.penalty_out(p, emp)
    out["rule_version"] = version.version if version else None
    out["can_decide"] = has(auth, "penalties.approve") and emp.user_id != auth.user.id and (p.status == "pending" or (p.status == "approved"))
    out["can_approve"] = has(auth, "penalties.approve") and emp.user_id != auth.user.id and p.status == "pending"
    out["can_cancel"] = has(auth, "penalties.approve") and p.status in ("pending", "approved")
    letters = (
        await db.scalars(
            select(ExplanationLetter).where(ExplanationLetter.employee_id == emp.id, ExplanationLetter.incident_date == p.incident_date, ExplanationLetter.status != "cancelled")
        )
    ).all() if can_read_all(auth) or emp.user_id == auth.user.id else []
    out["letters"] = [{"id": letter.id, "number": letter.number, "status": letter.status} for letter in letters]
    return out


class DecisionIn(BaseModel):
    action: str = Field(pattern="^(approve|cancel)$")
    note: str | None = Field(default=None, max_length=1000)


@router.post("/{penalty_id}/decision")
async def decision(penalty_id: int, body: DecisionIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(require_any("penalties.approve"))):
    p, emp = await _load(db, penalty_id, auth, lock=True)
    if body.action == "approve" and emp.user_id == auth.user.id:
        raise error(403, "cannot_decide_own")  # o'z jarimasini tasdiqlay olmaydi
    notes = await svc.decide(db, p, emp, auth, body.action, body.note, now())
    await audit(db, request, f"penalties.{p.status}", actor=auth.user, entity="penalty", entity_id=p.id,
                data={"employee_id": emp.id, "amount": float(p.amount), "block_days": p.block_days}, note=p.decision_note)  # fmt: skip
    await db.commit()
    await deliver(request.app, notes)
    return svc.penalty_out(p, emp)

