"""Moliya: ish haqi sxemalari, oylik kiritmalar (savdo/reja/KPI/bonus/avans/ushlanma), hisob-kitob, tasdiq, eksport."""

import asyncio
import csv
import io
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any

from fastapi import APIRouter, Depends, Query, Request, Response
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import get_session
from app.deps import AuthContext, error, get_auth
from app.models.employees import Employee
from app.models.payroll import PayItem, PayrollLine, PayrollRun, PayScheme
from app.services import payroll as svc
from app.services.audit import audit
from app.services.notifications import deliver, notify
from app.services.xlsx_safe import harden
from app.timeutil import local_date, now

router = APIRouter(prefix="/v1/payroll", tags=["payroll"])
MONEY_MAX = Decimal("1000000000000")


def require(perm: str):
    async def dep(auth: AuthContext = Depends(get_auth)) -> AuthContext:
        if not svc.has(auth, perm):
            raise error(403, "forbidden")
        return auth

    return dep


Read = require("payroll.read")
Input = require("payroll.input")
Manage = require("payroll.manage")


def _money_field(v: Decimal, *, allow_zero=True) -> Decimal:
    if v < 0 or v > MONEY_MAX or (not allow_zero and v == 0):
        raise error(400, "amount_invalid")
    return v


async def _emp(db: AsyncSession, eid: int) -> Employee:
    e = await db.get(Employee, eid)
    if not e:
        raise error(400, "unknown_employee")
    return e


async def _locked(db: AsyncSession, period: str) -> bool:
    st = await db.scalar(select(PayrollRun.status).where(PayrollRun.period == period))
    return st in ("approved", "paid")


# --- Sxemalar -----------------------------------------------------------------------------


class SchemeIn(BaseModel):
    employee_id: int
    effective_from: date
    fixed_amount: Decimal = Field(default=Decimal(0), ge=0)
    commission_percent: Decimal = Field(default=Decimal(0), ge=0, le=100)
    plan_bonus: Decimal = Field(default=Decimal(0), ge=0)
    kpi_amount: Decimal = Field(default=Decimal(0), ge=0)
    note: str | None = Field(default=None, max_length=300)


def _scheme_out(s: PayScheme, names: dict[int, Employee] | None = None) -> dict[str, Any]:
    e = (names or {}).get(s.employee_id)
    return {
        "id": s.id, "employee": {"id": s.employee_id, "name": e.full_name if e else None, "employee_no": e.employee_no if e else None}, "effective_from": s.effective_from,
        "fixed_amount": int(s.fixed_amount), "commission_percent": float(s.commission_percent), "plan_bonus": int(s.plan_bonus), "kpi_amount": int(s.kpi_amount), "note": s.note,
    }  # fmt: skip


@router.get("/schemes")
async def schemes(
    q: str | None = None,
    current: bool = True,
    employee_id: int | None = None,
    page: int = Query(default=1, ge=1),
    size: int = Query(default=30, ge=1, le=100),
    db: AsyncSession = Depends(get_session),
    auth: AuthContext = Depends(Read),
):
    """Joriy sxemalar (har xodimning eng yangisi) yoki hammasi."""
    S = PayScheme
    stmt = select(S).join(Employee, Employee.id == S.employee_id)
    if q and q.strip():
        stmt = stmt.where(Employee.full_name.ilike(f"%{q.strip()}%"))
    if employee_id:
        stmt = stmt.where(S.employee_id == employee_id)
    if current:
        latest = select(func.max(S.effective_from)).where(S.employee_id == Employee.id, S.effective_from <= local_date(now())).correlate(Employee).scalar_subquery()
        stmt = stmt.where(S.effective_from == latest)
    total = await db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    rows = (await db.scalars(stmt.order_by(Employee.full_name, S.effective_from.desc()).offset((page - 1) * size).limit(size))).all()
    emps = {e.id: e for e in (await db.scalars(select(Employee).where(Employee.id.in_({r.employee_id for r in rows})))).all()} if rows else {}
    return {"total": total, "page": page, "size": size, "items": [_scheme_out(r, emps) for r in rows]}


@router.post("/schemes", status_code=201)
async def create_scheme(body: SchemeIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    e = await _emp(db, body.employee_id)
    for v in (body.fixed_amount, body.plan_bonus, body.kpi_amount):
        _money_field(v)
    if await db.scalar(select(PayScheme.id).where(PayScheme.employee_id == e.id, PayScheme.effective_from == body.effective_from)):
        raise error(409, "scheme_exists")
    s = PayScheme(**body.model_dump(), created_by=auth.user.id)
    db.add(s)
    await db.flush()
    await audit(db, request, "payroll.scheme_create", actor=auth.user, entity="pay_scheme", entity_id=s.id, data={"employee_id": e.id, "fixed": str(s.fixed_amount), "from": str(s.effective_from)})
    await db.commit()
    return _scheme_out(s, {e.id: e})


async def _scheme(db: AsyncSession, sid: int) -> PayScheme:
    s = await db.get(PayScheme, sid)
    if not s:
        raise error(404, "not_found")
    return s


@router.get("/schemes/{sid}")
async def scheme(sid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Read)):
    s = await _scheme(db, sid)
    return _scheme_out(s, {s.employee_id: await db.get(Employee, s.employee_id)})  # type: ignore[dict-item]


@router.put("/schemes/{sid}")
async def update_scheme(sid: int, body: SchemeIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    s = await _scheme(db, sid)
    if body.employee_id != s.employee_id:
        raise error(400, "employee_immutable")
    for v in (body.fixed_amount, body.plan_bonus, body.kpi_amount):
        _money_field(v)
    if await db.scalar(select(PayScheme.id).where(PayScheme.employee_id == s.employee_id, PayScheme.effective_from == body.effective_from, PayScheme.id != s.id)):
        raise error(409, "scheme_exists")
    for k, v in body.model_dump().items():
        setattr(s, k, v)
    await audit(db, request, "payroll.scheme_update", actor=auth.user, entity="pay_scheme", entity_id=s.id, data={"employee_id": s.employee_id, "fixed": str(s.fixed_amount), "from": str(s.effective_from)})
    await db.commit()
    return _scheme_out(s, {s.employee_id: await db.get(Employee, s.employee_id)})  # type: ignore[dict-item]


@router.delete("/schemes/{sid}", status_code=204)
async def delete_scheme(sid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    s = await _scheme(db, sid)
    await audit(db, request, "payroll.scheme_delete", actor=auth.user, entity="pay_scheme", entity_id=s.id, data={"employee_id": s.employee_id, "from": str(s.effective_from)})
    await db.delete(s)
    await db.commit()


# --- Kiritmalar ---------------------------------------------------------------------------


class ItemIn(BaseModel):
    employee_id: int
    period: str
    kind: str = Field(pattern="^(sales|plan|kpi|bonus|advance|deduction)$")
    amount: Decimal
    note: str | None = Field(default=None, max_length=300)


def _check_item(b: ItemIn) -> None:
    svc.check_period(b.period)
    _money_field(b.amount, allow_zero=b.kind == "kpi")
    if b.kind == "kpi" and b.amount > 100:
        raise error(400, "kpi_range")


def _item_out(i: PayItem, emps: dict[int, Employee]) -> dict[str, Any]:
    e = emps.get(i.employee_id)
    return {"id": i.id, "employee": {"id": i.employee_id, "name": e.full_name if e else None, "employee_no": e.employee_no if e else None}, "period": i.period, "kind": i.kind, "amount": float(i.amount), "note": i.note, "created_at": i.created_at}


@router.get("/items")
async def items(
    period: str,
    kind: str | None = Query(default=None, pattern="^(sales|plan|kpi|bonus|advance|deduction)$"),
    q: str | None = None,
    employee_id: int | None = None,
    page: int = Query(default=1, ge=1),
    size: int = Query(default=30, ge=1, le=100),
    db: AsyncSession = Depends(get_session),
    auth: AuthContext = Depends(Read),
):
    svc.check_period(period)
    I = PayItem  # noqa: E741
    stmt = select(I).join(Employee, Employee.id == I.employee_id).where(I.period == period)
    if kind:
        stmt = stmt.where(I.kind == kind)
    if employee_id:
        stmt = stmt.where(I.employee_id == employee_id)
    if q and q.strip():
        stmt = stmt.where(Employee.full_name.ilike(f"%{q.strip()}%"))
    total = await db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    rows = (await db.scalars(stmt.order_by(Employee.full_name, I.kind, I.id).offset((page - 1) * size).limit(size))).all()
    emps = {e.id: e for e in (await db.scalars(select(Employee).where(Employee.id.in_({r.employee_id for r in rows})))).all()} if rows else {}
    return {"total": total, "page": page, "size": size, "locked": await _locked(db, period), "items": [_item_out(r, emps) for r in rows]}


@router.post("/items", status_code=201)
async def create_item(body: ItemIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Input)):
    _check_item(body)
    e = await _emp(db, body.employee_id)
    if await _locked(db, body.period):
        raise error(409, "period_locked")
    it = PayItem(**body.model_dump(), created_by=auth.user.id)
    db.add(it)
    await db.flush()
    await audit(db, request, "payroll.item_create", actor=auth.user, entity="pay_item", entity_id=it.id, data={"employee_id": e.id, "period": it.period, "kind": it.kind, "amount": str(it.amount)})
    await db.commit()
    return _item_out(it, {e.id: e})


async def _item(db: AsyncSession, iid: int) -> PayItem:
    it = await db.get(PayItem, iid)
    if not it:
        raise error(404, "not_found")
    return it


@router.get("/items/{iid}")
async def item(iid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Read)):
    it = await _item(db, iid)
    return _item_out(it, {it.employee_id: await db.get(Employee, it.employee_id)})  # type: ignore[dict-item]


@router.put("/items/{iid}")
async def update_item(iid: int, body: ItemIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Input)):
    it = await _item(db, iid)
    _check_item(body)
    if body.employee_id != it.employee_id:
        raise error(400, "employee_immutable")
    if await _locked(db, it.period) or await _locked(db, body.period):
        raise error(409, "period_locked")
    before = {"period": it.period, "kind": it.kind, "amount": str(it.amount)}
    for k, v in body.model_dump().items():
        setattr(it, k, v)
    await audit(db, request, "payroll.item_update", actor=auth.user, entity="pay_item", entity_id=it.id, data={"employee_id": it.employee_id, "before": before, "after": {"period": it.period, "kind": it.kind, "amount": str(it.amount)}})
    await db.commit()
    return _item_out(it, {it.employee_id: await db.get(Employee, it.employee_id)})  # type: ignore[dict-item]


@router.delete("/items/{iid}", status_code=204)
async def delete_item(iid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Input)):
    it = await _item(db, iid)
    if await _locked(db, it.period):
        raise error(409, "period_locked")
    await audit(db, request, "payroll.item_delete", actor=auth.user, entity="pay_item", entity_id=it.id, data={"employee_id": it.employee_id, "period": it.period, "kind": it.kind, "amount": str(it.amount)})
    await db.delete(it)
    await db.commit()


class ImportIn(BaseModel):
    period: str
    csv: str = Field(min_length=3, max_length=500_000)
    dry_run: bool = True


@router.post("/items/import")
async def import_items(body: ImportIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Input)):
    """CSV (`xodim_no;tur;summa;izoh`, sarlavhali yoki sarlavhasiz). `dry_run` da faqat tekshiradi."""
    svc.check_period(body.period)
    if await _locked(db, body.period):
        raise error(409, "period_locked")
    text = body.csv.strip()
    delim = ";" if text.splitlines()[0].count(";") >= text.splitlines()[0].count(",") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    if rows and not rows[0][0].strip().lstrip("№#").isdigit():
        rows = rows[1:]  # sarlavha
    nos = {int(r[0].strip().lstrip("№#")) for r in rows if r and r[0].strip().lstrip("№#").isdigit()}
    emps = {e.employee_no: e for e in (await db.scalars(select(Employee).where(Employee.employee_no.in_(nos)))).all()} if nos else {}
    good: list[PayItem] = []
    errors: list[dict[str, Any]] = []
    for n, r in enumerate(rows, 1):
        if not r or not any(c.strip() for c in r):
            continue
        try:
            no, kind, amount = int(r[0].strip().lstrip("№#")), r[1].strip().lower(), Decimal(r[2].strip().replace(" ", "").replace(",", "."))
            note = r[3].strip()[:300] if len(r) > 3 and r[3].strip() else None
        except (ValueError, IndexError, InvalidOperation):
            errors.append({"row": n, "code": "row_invalid"})
            continue
        if kind not in svc.KINDS:
            errors.append({"row": n, "code": "kind_invalid"})
        elif no not in emps:
            errors.append({"row": n, "code": "unknown_employee"})
        elif amount < 0 or amount > MONEY_MAX or (kind == "kpi" and amount > 100):
            errors.append({"row": n, "code": "amount_invalid"})
        else:
            good.append(PayItem(employee_id=emps[no].id, period=body.period, kind=kind, amount=amount, note=note, created_by=auth.user.id))
    if errors and not body.dry_run:
        raise error(400, "import_invalid", errors=errors[:50])
    if not body.dry_run and good:
        db.add_all(good)
        await audit(db, request, "payroll.item_import", actor=auth.user, entity="pay_item", data={"period": body.period, "count": len(good)})
        await db.commit()
    return {"ok": not errors, "valid": len(good), "errors": errors[:50], "imported": len(good) if not body.dry_run and not errors else 0}


# --- Hisob-kitoblar -----------------------------------------------------------------------


class RunIn(BaseModel):
    period: str
    tax_percent: Decimal = Field(default=Decimal("12"), ge=0, le=100)
    inps_percent: Decimal = Field(default=Decimal("0.1"), ge=0, le=100)
    paid_leave: bool = True
    note: str | None = Field(default=None, max_length=1000)


def _run_out(r: PayrollRun) -> dict[str, Any]:
    return {"id": r.id, "period": r.period, "status": r.status, "settings": {k: v for k, v in r.settings.items() if not k.startswith("_")}, "note": r.note, "calculated_at": r.calculated_at, "approved_at": r.approved_at, "paid_at": r.paid_at}


async def _run(db: AsyncSession, rid: int, *, lock: bool = False) -> PayrollRun:
    stmt = select(PayrollRun).where(PayrollRun.id == rid)
    r = await db.scalar(stmt.with_for_update() if lock else stmt)
    if not r:
        raise error(404, "not_found")
    return r


@router.get("/runs")
async def runs(db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Read)):
    rows = (await db.scalars(select(PayrollRun).order_by(PayrollRun.period.desc()))).all()
    sums = {r[0]: r for r in (await db.execute(select(PayrollLine.run_id, func.count(), func.coalesce(func.sum(PayrollLine.gross), 0), func.coalesce(func.sum(PayrollLine.to_pay), 0)).group_by(PayrollLine.run_id))).all()}
    return [{**_run_out(r), "employees": sums[r.id][1] if r.id in sums else 0, "gross": int(sums[r.id][2]) if r.id in sums else 0, "to_pay": int(sums[r.id][3]) if r.id in sums else 0} for r in rows]


@router.post("/runs", status_code=201)
async def create_run(body: RunIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    svc.check_period(body.period)
    if body.period > local_date(now()).strftime("%Y-%m"):
        raise error(400, "period_future")
    if await db.scalar(select(PayrollRun.id).where(PayrollRun.period == body.period)):
        raise error(409, "run_exists")
    r = PayrollRun(period=body.period, status="draft", settings={"tax_percent": str(body.tax_percent), "inps_percent": str(body.inps_percent), "paid_leave": body.paid_leave}, note=(body.note or "").strip() or None, created_by=auth.user.id)
    db.add(r)
    await db.flush()
    await audit(db, request, "payroll.run_create", actor=auth.user, entity="payroll_run", entity_id=r.id, data={"period": r.period, "settings": r.settings})
    await db.commit()
    return _run_out(r)


@router.get("/runs/{rid}")
async def run(rid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Read)):
    r = await _run(db, rid)
    start, end = svc.month_bounds(r.period)
    missing = 0
    if r.status == "draft":
        have = set((await db.scalars(select(PayrollLine.employee_id).where(PayrollLine.run_id == r.id))).all())
        covered = set((await db.scalars(select(PayScheme.employee_id).where(PayScheme.effective_from <= end))).all())
        missing = len([e for e in await svc.employees_for(db, start, end) if e.status == "active" and e.id not in covered and e.id not in have])
    flagged = await db.scalar(select(func.count()).select_from(PayrollLine).where(PayrollLine.run_id == r.id, PayrollLine.details["warnings"].astext != "[]")) or 0
    return {**_run_out(r), "totals": await svc.totals(db, r.id), "missing_scheme": missing, "flagged": flagged}


@router.get("/runs/{rid}/lines")
async def run_lines(
    rid: int,
    q: str | None = None,
    flagged: bool = False,
    page: int = Query(default=1, ge=1),
    size: int = Query(default=50, ge=1, le=200),
    db: AsyncSession = Depends(get_session),
    auth: AuthContext = Depends(Read),
):
    r = await _run(db, rid)
    L = PayrollLine
    conds: list = [L.run_id == r.id]
    if q and q.strip():
        conds.append(L.employee_name.ilike(f"%{q.strip()}%"))
    if flagged:
        conds.append(L.details["warnings"].astext != "[]")
    total = await db.scalar(select(func.count()).select_from(L).where(*conds)) or 0
    rows = (await db.scalars(select(L).where(*conds).order_by(L.employee_name, L.id).offset((page - 1) * size).limit(size))).all()
    return {"total": total, "page": page, "size": size, "items": [svc.line_out(x) for x in rows]}


@router.get("/runs/{rid}/lines/{lid}")
async def run_line(rid: int, lid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Read)):
    l = await db.get(PayrollLine, lid)  # noqa: E741
    if not l or l.run_id != rid:
        raise error(404, "not_found")
    r = await _run(db, rid)
    return {**svc.line_out(l, detail=True), "period": r.period, "status": r.status}


@router.post("/runs/{rid}/calculate")
async def calculate(rid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    r = await _run(db, rid, lock=True)
    if r.status != "draft":
        raise error(409, "run_locked")
    res = await svc.calculate(db, r, now())
    await audit(db, request, "payroll.calculate", actor=auth.user, entity="payroll_run", entity_id=r.id, data={"period": r.period, "lines": res["lines"], "missing_scheme": len(res["missing_scheme"])})
    await db.commit()
    return {**res, "run": _run_out(r), "totals": await svc.totals(db, r.id)}


@router.post("/runs/{rid}/approve")
async def approve(rid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    r = await _run(db, rid, lock=True)
    if r.status != "draft":
        raise error(409, "run_locked")
    if not r.calculated_at:
        raise error(409, "not_calculated")
    # kiritma, jarima yoki sxema hisobdan keyin o'zgargan bo'lsa, eskirgan hisobni tasdiqlab bo'lmaydi
    if r.settings.get("_sig") != await svc.inputs_signature(db, r.period):
        raise error(409, "stale_calculation")
    lines = (await db.scalars(select(PayrollLine).where(PayrollLine.run_id == r.id))).all()
    if not lines:
        raise error(409, "no_lines")
    r.status, r.approved_at, r.approved_by = "approved", now(), auth.user.id
    totals = await svc.totals(db, r.id)
    await audit(db, request, "payroll.approve", actor=auth.user, entity="payroll_run", entity_id=r.id, data={"period": r.period, "totals": totals})
    users = {e.id: e.user_id for e in (await db.scalars(select(Employee).where(Employee.id.in_([l.employee_id for l in lines])))).all()}  # noqa: E741
    notes = []
    for l in lines:  # noqa: E741
        if users.get(l.employee_id):
            notes += await notify(db, [users[l.employee_id]], "payroll.payslip", link=f"/payroll/my/{l.id}", params={"period": r.period, "amount": f"{int(l.to_pay):,}".replace(",", " ")})
    await db.commit()
    await deliver(request.app, notes)
    return _run_out(r)


@router.post("/runs/{rid}/reopen")
async def reopen(rid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    r = await _run(db, rid, lock=True)
    if r.status != "approved":
        raise error(409, "run_locked")
    r.status, r.approved_at, r.approved_by = "draft", None, None
    await audit(db, request, "payroll.reopen", actor=auth.user, entity="payroll_run", entity_id=r.id, data={"period": r.period})
    await db.commit()
    return _run_out(r)


@router.post("/runs/{rid}/paid")
async def mark_paid(rid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    r = await _run(db, rid, lock=True)
    if r.status != "approved":
        raise error(409, "run_locked")
    r.status, r.paid_at = "paid", now()
    await audit(db, request, "payroll.paid", actor=auth.user, entity="payroll_run", entity_id=r.id, data={"period": r.period, "totals": await svc.totals(db, r.id)})
    await db.commit()
    return _run_out(r)


@router.delete("/runs/{rid}", status_code=204)
async def delete_run(rid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manage)):
    r = await _run(db, rid, lock=True)
    if r.status != "draft":
        raise error(409, "run_locked")
    await audit(db, request, "payroll.run_delete", actor=auth.user, entity="payroll_run", entity_id=r.id, data={"period": r.period})
    await db.delete(r)
    await db.commit()


def _xlsx(r: PayrollRun, lines: list[PayrollLine]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = r.period
    head = ["№", "F.I.SH", "Filial", "Bo'lim", "Reja kun", "To'langan kun", "Oklad", "Komissiya", "Reja bonus", "KPI bonus", "Bonus", "Yalpi", "Jarima", "Ushlanma", "Soliq bazasi", "JShDS", "INPS", "Sof", "Avans", "To'lanadi"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    keys = ("fixed", "commission", "plan_bonus", "kpi_bonus", "bonus", "gross", "penalties", "deductions", "taxable", "tax", "inps", "net", "advance", "to_pay")
    for i, l in enumerate(lines, 1):  # noqa: E741
        ws.append([l.employee_no or i, l.employee_name, l.branch_name or "", l.department_name or "", l.plan_days, l.paid_days] + [int(getattr(l, k)) for k in keys])
    n = len(lines) + 1
    ws.append(["", "JAMI"] + [""] * 4 + [f"=SUM({chr(71 + j)}2:{chr(71 + j)}{n})" for j in range(len(keys))])
    for c in ws[n + 1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=len(head)):
        for c in row:
            c.number_format = "#,##0"
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.freeze_panes = "C2"
    buf = io.BytesIO()
    harden(wb)
    wb.save(buf)
    return buf.getvalue()


@router.get("/runs/{rid}/export")
async def export(rid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Read)):
    r = await _run(db, rid)
    lines = list((await db.scalars(select(PayrollLine).where(PayrollLine.run_id == r.id).order_by(PayrollLine.employee_name, PayrollLine.id))).all())
    content = await asyncio.to_thread(_xlsx, r, lines)
    await audit(db, request, "payroll.export", actor=auth.user, entity="payroll_run", entity_id=r.id, data={"period": r.period, "lines": len(lines)})
    await db.commit()
    return Response(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="oylik-{r.period}.xlsx"'})


# --- Xodimning o'z varaqasi ---------------------------------------------------------------


async def _my_employees(db: AsyncSession, auth: AuthContext) -> list[int]:
    return list((await db.scalars(select(Employee.id).where(Employee.user_id == auth.user.id))).all())


@router.get("/my")
async def my_payslips(db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(require("payroll.own"))):
    ids = await _my_employees(db, auth)
    if not ids:
        return []
    rows = (await db.execute(select(PayrollLine, PayrollRun).join(PayrollRun, PayrollRun.id == PayrollLine.run_id).where(PayrollLine.employee_id.in_(ids), PayrollRun.status.in_(("approved", "paid"))).order_by(PayrollRun.period.desc()))).all()
    return [{"id": l.id, "period": r.period, "status": r.status, "gross": int(l.gross), "net": int(l.net), "to_pay": int(l.to_pay)} for l, r in rows]  # noqa: E741


@router.get("/my/{lid}")
async def my_payslip(lid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(require("payroll.own"))):
    ids = await _my_employees(db, auth)
    l = await db.get(PayrollLine, lid)  # noqa: E741
    r = await db.get(PayrollRun, l.run_id) if l else None
    if not l or not r or l.employee_id not in ids or r.status not in ("approved", "paid"):
        raise error(404, "not_found")
    return {**svc.line_out(l, detail=True), "period": r.period, "status": r.status}
