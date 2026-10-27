"""Logistika: mashinalar, hujjat muddatlari, haydovchilar, reyslar, jonli GPS, yoqilg'i va hisobotlar."""

import asyncio
import io
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any

from fastapi import APIRouter, Depends, Query, Request, Response
from pydantic import BaseModel, Field
from sqlalchemy import func, literal_column, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import get_session
from app.deps import AuthContext, error, get_auth
from app.models.employees import Employee
from app.models.logistics import DriverProfile, FuelLog, GpsPoint, Trip, Vehicle, VehicleDocument
from app.models.org import Branch
from app.services import logistics as svc
from app.services.audit import audit
from app.services.xlsx_safe import harden
from app.timeutil import local_date, now, tz

router = APIRouter(prefix="/v1/logistics", tags=["logistics"])


def require_any(*perms: str):
    async def dep(auth: AuthContext = Depends(get_auth)) -> AuthContext:
        if not any(svc.has(auth, p) for p in perms):
            raise error(403, "forbidden")
        return auth

    return dep


Access = require_any("logistics.own", "logistics.read", "logistics.manage")
Reader = require_any("logistics.read", "logistics.manage")
Manager = require_any("logistics.manage")
Driver = require_any("logistics.own")


async def _names(db: AsyncSession, ids: set[int | None]) -> dict[int, str]:
    ids2 = {i for i in ids if i}
    return {e.id: e.full_name for e in (await db.scalars(select(Employee).where(Employee.id.in_(ids2)))).all()} if ids2 else {}


async def _my_employee(db: AsyncSession, auth: AuthContext) -> Employee | None:
    return await db.scalar(select(Employee).where(Employee.user_id == auth.user.id, Employee.status == "active"))


# --- Mashinalar ---------------------------------------------------------------------------


class VehicleIn(BaseModel):
    plate: str = Field(min_length=3, max_length=16)
    name: str = Field(min_length=1, max_length=120)
    kind: str = Field(default="car", pattern="^(car|van|truck|moto|other)$")
    fuel_type: str = Field(default="petrol", pattern="^(petrol|diesel|gas|propane|electric)$")
    fuel_norm: Decimal = Field(default=Decimal(0), ge=0, le=200)
    tank_l: int | None = Field(default=None, ge=1, le=2000)
    odometer_km: int = Field(default=0, ge=0, le=5_000_000)
    branch_id: int | None = None
    status: str = Field(default="active", pattern="^(active|service|retired)$")
    note: str | None = Field(default=None, max_length=300)


def _plate(p: str) -> str:
    return " ".join(p.upper().split())


def _vehicle_out(v: Vehicle, branches: dict[int, str], extra: dict[str, Any] | None = None) -> dict[str, Any]:
    return {
        "id": v.id, "plate": v.plate, "name": v.name, "kind": v.kind, "fuel_type": v.fuel_type, "fuel_norm": float(v.fuel_norm), "tank_l": v.tank_l, "odometer_km": v.odometer_km,
        "branch": {"id": v.branch_id, "name": branches.get(v.branch_id or 0)} if v.branch_id else None, "status": v.status, "note": v.note, **(extra or {}),
    }  # fmt: skip


async def _vehicle(db: AsyncSession, vid: int) -> Vehicle:
    v = await db.get(Vehicle, vid)
    if not v:
        raise error(404, "not_found")
    return v


async def _doc_flags(db: AsyncSession, ids: list[int]) -> dict[int, dict[str, int]]:
    today = local_date(now())
    out: dict[int, dict[str, int]] = {i: {"expired": 0, "soon": 0} for i in ids}
    if ids:
        for d in (await db.scalars(select(VehicleDocument).where(VehicleDocument.vehicle_id.in_(ids)))).all():
            st = svc.doc_stage(d.expires_at, today)
            if st == 3:
                out[d.vehicle_id]["expired"] += 1
            elif st >= 1:
                out[d.vehicle_id]["soon"] += 1
    return out


@router.get("/vehicles")
async def vehicles(q: str | None = None, status: str | None = Query(default=None, pattern="^(active|service|retired)$"), branch_id: int | None = None, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    stmt = select(Vehicle).order_by(Vehicle.plate)
    if q and q.strip():
        like = f"%{q.strip()}%"
        stmt = stmt.where(Vehicle.plate.ilike(like) | Vehicle.name.ilike(like))
    if status:
        stmt = stmt.where(Vehicle.status == status)
    if branch_id:
        stmt = stmt.where(Vehicle.branch_id == branch_id)
    rows = list((await db.scalars(stmt)).all())
    flags = await _doc_flags(db, [v.id for v in rows])
    active = {t.vehicle_id: t.id for t in (await db.scalars(select(Trip).where(Trip.status == "active"))).all()}
    branches = {b.id: b.name for b in (await db.scalars(select(Branch))).all()}
    return [_vehicle_out(v, branches, {"docs": flags[v.id], "active_trip_id": active.get(v.id)}) for v in rows]


@router.post("/vehicles", status_code=201)
async def create_vehicle(body: VehicleIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    if body.branch_id and not await db.get(Branch, body.branch_id):
        raise error(400, "unknown_branch")
    v = Vehicle(**{**body.model_dump(), "plate": _plate(body.plate), "name": body.name.strip()})
    db.add(v)
    try:
        await db.flush()
    except IntegrityError:
        raise error(409, "plate_exists") from None
    await audit(db, request, "logistics.vehicle_create", actor=auth.user, entity="vehicle", entity_id=v.id, data={"plate": v.plate})
    await db.commit()
    return _vehicle_out(v, {})


@router.get("/vehicles/{vid}")
async def vehicle(vid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    v = await _vehicle(db, vid)
    branches = {b.id: b.name for b in (await db.scalars(select(Branch))).all()}
    return _vehicle_out(v, branches, {"docs": (await _doc_flags(db, [v.id]))[v.id]})


@router.put("/vehicles/{vid}")
async def update_vehicle(vid: int, body: VehicleIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    v = await _vehicle(db, vid)
    if body.branch_id and not await db.get(Branch, body.branch_id):
        raise error(400, "unknown_branch")
    for k, val in {**body.model_dump(), "plate": _plate(body.plate), "name": body.name.strip()}.items():
        setattr(v, k, val)
    try:
        await db.flush()
    except IntegrityError:
        raise error(409, "plate_exists") from None
    await audit(db, request, "logistics.vehicle_update", actor=auth.user, entity="vehicle", entity_id=v.id, data={"plate": v.plate})
    await db.commit()
    return _vehicle_out(v, {})


@router.delete("/vehicles/{vid}", status_code=204)
async def delete_vehicle(vid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    v = await _vehicle(db, vid)
    if await db.scalar(select(func.count()).select_from(Trip).where(Trip.vehicle_id == v.id)):
        raise error(409, "vehicle_in_use")
    await audit(db, request, "logistics.vehicle_delete", actor=auth.user, entity="vehicle", entity_id=v.id, data={"plate": v.plate})
    await db.delete(v)
    await db.commit()


# --- Hujjatlar ----------------------------------------------------------------------------


class DocIn(BaseModel):
    kind: str = Field(pattern="^(tech_passport|insurance|osago|inspection|permit|other)$")
    number: str | None = Field(default=None, max_length=60)
    issued_at: date | None = None
    expires_at: date
    note: str | None = Field(default=None, max_length=300)


def _doc_out(d: VehicleDocument, today: date) -> dict[str, Any]:
    return {"id": d.id, "vehicle_id": d.vehicle_id, "kind": d.kind, "number": d.number, "issued_at": d.issued_at, "expires_at": d.expires_at, "note": d.note, "days_left": (d.expires_at - today).days, "stage": svc.doc_stage(d.expires_at, today)}


@router.get("/vehicles/{vid}/documents")
async def documents(vid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    await _vehicle(db, vid)
    today = local_date(now())
    rows = (await db.scalars(select(VehicleDocument).where(VehicleDocument.vehicle_id == vid).order_by(VehicleDocument.expires_at))).all()
    return [_doc_out(d, today) for d in rows]


@router.post("/vehicles/{vid}/documents", status_code=201)
async def create_document(vid: int, body: DocIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    await _vehicle(db, vid)
    if body.issued_at and body.expires_at < body.issued_at:
        raise error(400, "dates_invalid")
    d = VehicleDocument(vehicle_id=vid, **body.model_dump())
    d.notified_stage = svc.doc_stage(d.expires_at, local_date(now()))  # allaqachon o'tib ketganini qayta xabar qilmaymiz
    db.add(d)
    await db.flush()
    await audit(db, request, "logistics.document_create", actor=auth.user, entity="vehicle_document", entity_id=d.id, data={"vehicle_id": vid, "kind": d.kind, "expires_at": str(d.expires_at)})
    await db.commit()
    return _doc_out(d, local_date(now()))


async def _document(db: AsyncSession, did: int) -> VehicleDocument:
    d = await db.get(VehicleDocument, did)
    if not d:
        raise error(404, "not_found")
    return d


@router.get("/documents/{did}")
async def document(did: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    return _doc_out(await _document(db, did), local_date(now()))


@router.put("/documents/{did}")
async def update_document(did: int, body: DocIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    d = await _document(db, did)
    if body.issued_at and body.expires_at < body.issued_at:
        raise error(400, "dates_invalid")
    if body.expires_at != d.expires_at:
        d.notified_stage = 0  # yangi muddat: eslatmalar qaytadan
    for k, v in body.model_dump().items():
        setattr(d, k, v)
    await audit(db, request, "logistics.document_update", actor=auth.user, entity="vehicle_document", entity_id=d.id, data={"vehicle_id": d.vehicle_id, "kind": d.kind, "expires_at": str(d.expires_at)})
    await db.commit()
    return _doc_out(d, local_date(now()))


@router.delete("/documents/{did}", status_code=204)
async def delete_document(did: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    d = await _document(db, did)
    await audit(db, request, "logistics.document_delete", actor=auth.user, entity="vehicle_document", entity_id=d.id, data={"vehicle_id": d.vehicle_id, "kind": d.kind})
    await db.delete(d)
    await db.commit()


# --- Haydovchilar -------------------------------------------------------------------------


class DriverIn(BaseModel):
    license_no: str | None = Field(default=None, max_length=30)
    categories: str | None = Field(default=None, max_length=30)
    license_expires: date | None = None
    note: str | None = Field(default=None, max_length=300)


def _driver_out(p: DriverProfile | None, e: Employee, today: date) -> dict[str, Any]:
    left = (p.license_expires - today).days if p and p.license_expires else None
    return {
        "employee": {"id": e.id, "name": e.full_name, "phone": e.phone, "has_account": e.user_id is not None, "status": e.status}, "has_profile": p is not None,
        "license_no": p.license_no if p else None, "categories": p.categories if p else None, "license_expires": p.license_expires if p else None, "note": p.note if p else None,
        "license_days_left": left, "license_stage": svc.doc_stage(p.license_expires, today) if p and p.license_expires else 0,
    }  # fmt: skip


@router.get("/drivers")
async def drivers(db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    today = local_date(now())
    rows = (await db.execute(select(DriverProfile, Employee).join(Employee, Employee.id == DriverProfile.employee_id).order_by(Employee.full_name))).all()
    active = {t.driver_id: t.id for t in (await db.scalars(select(Trip).where(Trip.status == "active"))).all()}
    return [{**_driver_out(p, e, today), "active_trip_id": active.get(e.id)} for p, e in rows]


@router.get("/drivers/{eid}")
async def driver(eid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    e = await db.get(Employee, eid)
    if not e:
        raise error(404, "not_found")
    p = await db.scalar(select(DriverProfile).where(DriverProfile.employee_id == eid))
    return _driver_out(p, e, local_date(now()))


@router.put("/drivers/{eid}")
async def put_driver(eid: int, body: DriverIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    e = await db.get(Employee, eid)
    if not e or e.status != "active":
        raise error(400, "unknown_employee")
    p = await db.scalar(select(DriverProfile).where(DriverProfile.employee_id == eid))
    if p is None:
        p = DriverProfile(employee_id=eid)
        db.add(p)
    if body.license_expires != p.license_expires:
        p.notified_stage = 0
    p.license_no, p.categories, p.license_expires, p.note = (body.license_no or "").strip() or None, (body.categories or "").strip() or None, body.license_expires, (body.note or "").strip() or None
    await db.flush()
    await audit(db, request, "logistics.driver_save", actor=auth.user, entity="driver", entity_id=eid, data={"license_expires": str(p.license_expires) if p.license_expires else None})
    await db.commit()
    return _driver_out(p, e, local_date(now()))


@router.delete("/drivers/{eid}", status_code=204)
async def delete_driver(eid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    p = await db.scalar(select(DriverProfile).where(DriverProfile.employee_id == eid))
    if not p:
        raise error(404, "not_found")
    if await db.scalar(select(func.count()).select_from(Trip).where(Trip.driver_id == eid, Trip.status.in_(("planned", "active")))):
        raise error(409, "driver_has_trips")
    await audit(db, request, "logistics.driver_delete", actor=auth.user, entity="driver", entity_id=eid)
    await db.delete(p)
    await db.commit()


# --- Reyslar ------------------------------------------------------------------------------


class TripIn(BaseModel):
    vehicle_id: int
    driver_id: int
    purpose: str = Field(min_length=2, max_length=200)
    origin: str | None = Field(default=None, max_length=200)
    destination: str | None = Field(default=None, max_length=200)
    planned_start: datetime | None = None


async def _trip_out_many(db: AsyncSession, rows: list[Trip]) -> list[dict[str, Any]]:
    if not rows:
        return []
    names = await _names(db, {t.driver_id for t in rows})
    veh = {v.id: v for v in (await db.scalars(select(Vehicle).where(Vehicle.id.in_({t.vehicle_id for t in rows})))).all()}
    return [
        {
            "id": t.id, "status": t.status, "purpose": t.purpose, "origin": t.origin, "destination": t.destination, "planned_start": t.planned_start,
            "vehicle": {"id": t.vehicle_id, "plate": veh[t.vehicle_id].plate, "name": veh[t.vehicle_id].name}, "driver": {"id": t.driver_id, "name": names.get(t.driver_id)},
            "started_at": t.started_at, "finished_at": t.finished_at, "start_odometer": t.start_odometer, "end_odometer": t.end_odometer,
            "distance_km": float(t.distance_km) if t.distance_km is not None else None, "gps_km": float(t.gps_km) if t.gps_km is not None else None,
            "fuel_norm_l": float(t.fuel_norm_l) if t.fuel_norm_l is not None else None, "note": t.note,
        }
        for t in rows
    ]  # fmt: skip


async def _trip(db: AsyncSession, tid: int, auth: AuthContext, *, lock: bool = False) -> Trip:
    stmt = select(Trip).where(Trip.id == tid)
    t = await db.scalar(stmt.with_for_update() if lock else stmt)
    if not t:
        raise error(404, "not_found")
    if not svc.sees_all(auth):
        e = await _my_employee(db, auth)
        if not e or e.id != t.driver_id:
            raise error(404, "not_found")
    return t


async def _check_trip_refs(db: AsyncSession, body: TripIn) -> None:
    v = await db.get(Vehicle, body.vehicle_id)
    if not v or v.status != "active":
        raise error(400, "vehicle_unavailable")
    e = await db.get(Employee, body.driver_id)
    if not e or e.status != "active":
        raise error(400, "unknown_employee")
    if not await db.scalar(select(DriverProfile.id).where(DriverProfile.employee_id == e.id)):
        raise error(400, "not_a_driver")


@router.get("/trips")
async def trips(
    scope: str = Query(default="all", pattern="^(mine|all)$"),
    status: str | None = Query(default=None, pattern="^(planned|active|completed|cancelled|open)$"),
    vehicle_id: int | None = None,
    driver_id: int | None = None,
    start: date | None = None,
    end: date | None = None,
    page: int = Query(default=1, ge=1),
    size: int = Query(default=20, ge=1, le=100),
    db: AsyncSession = Depends(get_session),
    auth: AuthContext = Depends(Access),
):
    T = Trip
    conds: list = []
    if scope == "mine" or not svc.sees_all(auth):
        e = await _my_employee(db, auth)
        conds.append(T.driver_id == (e.id if e else -1))
    if status == "open":
        conds.append(T.status.in_(("planned", "active")))
    elif status:
        conds.append(T.status == status)
    if vehicle_id:
        conds.append(T.vehicle_id == vehicle_id)
    if driver_id:
        conds.append(T.driver_id == driver_id)
    if start:
        conds.append(func.coalesce(T.started_at, T.planned_start, T.created_at) >= datetime.combine(start, datetime.min.time(), tzinfo=tz()))
    if end:
        conds.append(func.coalesce(T.started_at, T.planned_start, T.created_at) < datetime.combine(end + timedelta(days=1), datetime.min.time(), tzinfo=tz()))
    total = await db.scalar(select(func.count()).select_from(T).where(*conds)) or 0
    rows = (await db.scalars(select(T).where(*conds).order_by(func.coalesce(T.started_at, T.planned_start, T.created_at).desc(), T.id.desc()).offset((page - 1) * size).limit(size))).all()
    return {"total": total, "page": page, "size": size, "items": await _trip_out_many(db, list(rows))}


@router.post("/trips", status_code=201)
async def create_trip(body: TripIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    await _check_trip_refs(db, body)
    t = Trip(**{**body.model_dump(), "purpose": body.purpose.strip()}, status="planned", created_by=auth.user.id)
    db.add(t)
    await db.flush()
    await audit(db, request, "logistics.trip_create", actor=auth.user, entity="trip", entity_id=t.id, data={"vehicle_id": t.vehicle_id, "driver_id": t.driver_id})
    await db.commit()
    return (await _trip_out_many(db, [t]))[0]


@router.get("/trips/{tid}")
async def trip(tid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Access)):
    t = await _trip(db, tid, auth)
    out = (await _trip_out_many(db, [t]))[0]
    e = await _my_employee(db, auth)
    mine = bool(e and e.id == t.driver_id)
    n_points = await db.scalar(select(func.count()).select_from(GpsPoint).where(GpsPoint.trip_id == t.id)) or 0
    return {**out, "points": n_points, "can": {"edit": svc.can_manage(auth) and t.status == "planned", "start": t.status == "planned" and (mine or svc.can_manage(auth)), "finish": t.status == "active" and (mine or svc.can_manage(auth)), "cancel": t.status in ("planned", "active") and svc.can_manage(auth), "mine": mine}}


@router.put("/trips/{tid}")
async def update_trip(tid: int, body: TripIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    t = await _trip(db, tid, auth, lock=True)
    if t.status != "planned":
        raise error(409, "invalid_transition")
    await _check_trip_refs(db, body)
    for k, v in {**body.model_dump(), "purpose": body.purpose.strip()}.items():
        setattr(t, k, v)
    await audit(db, request, "logistics.trip_update", actor=auth.user, entity="trip", entity_id=t.id, data={"vehicle_id": t.vehicle_id, "driver_id": t.driver_id})
    await db.commit()
    return (await _trip_out_many(db, [t]))[0]


class StartIn(BaseModel):
    odometer: int = Field(ge=0, le=5_000_000)


@router.post("/trips/{tid}/start")
async def start_trip(tid: int, body: StartIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Access)):
    t = await _trip(db, tid, auth, lock=True)
    e = await _my_employee(db, auth)
    if not ((e and e.id == t.driver_id and svc.has(auth, "logistics.own")) or svc.can_manage(auth)):
        raise error(403, "forbidden")
    if t.status != "planned":
        raise error(409, "invalid_transition")
    v = await db.get(Vehicle, t.vehicle_id)
    if not v or v.status != "active":
        raise error(400, "vehicle_unavailable")
    if await db.scalar(select(Trip.id).where(Trip.status == "active", (Trip.vehicle_id == t.vehicle_id) | (Trip.driver_id == t.driver_id))):
        raise error(409, "trip_conflict")
    if body.odometer < v.odometer_km:
        raise error(400, "odometer_low", current=v.odometer_km)
    t.status, t.started_at, t.start_odometer = "active", now(), body.odometer
    v.odometer_km = body.odometer
    await audit(db, request, "logistics.trip_start", actor=auth.user, entity="trip", entity_id=t.id, data={"odometer": body.odometer})
    await db.commit()
    return (await _trip_out_many(db, [t]))[0]


class FinishIn(BaseModel):
    odometer: int = Field(ge=0, le=5_000_000)
    note: str | None = Field(default=None, max_length=2000)


@router.post("/trips/{tid}/finish")
async def finish_trip(tid: int, body: FinishIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Access)):
    t = await _trip(db, tid, auth, lock=True)
    e = await _my_employee(db, auth)
    if not ((e and e.id == t.driver_id and svc.has(auth, "logistics.own")) or svc.can_manage(auth)):
        raise error(403, "forbidden")
    if t.status != "active":
        raise error(409, "invalid_transition")
    if body.odometer < (t.start_odometer or 0):
        raise error(400, "odometer_low", current=t.start_odometer)
    v = await db.get(Vehicle, t.vehicle_id)
    pts = await svc.vehicle_points(db, t.id)
    t.gps_km = svc.path_km([(p.lat, p.lng, p.recorded_at) for p in pts]) if pts else None
    t.status, t.finished_at, t.end_odometer = "completed", now(), body.odometer
    t.distance_km = svc.dec(body.odometer - (t.start_odometer or 0))
    t.fuel_norm_l = svc.norm_liters(t.distance_km, v.fuel_norm) if v else None
    t.note = (body.note or "").strip() or t.note
    if v:
        v.odometer_km = max(v.odometer_km, body.odometer)
    await audit(db, request, "logistics.trip_finish", actor=auth.user, entity="trip", entity_id=t.id, data={"distance_km": float(t.distance_km), "gps_km": float(t.gps_km) if t.gps_km is not None else None})
    await db.commit()
    return (await _trip_out_many(db, [t]))[0]


class CancelIn(BaseModel):
    note: str = Field(min_length=3, max_length=500)


@router.post("/trips/{tid}/cancel")
async def cancel_trip(tid: int, body: CancelIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    t = await _trip(db, tid, auth, lock=True)
    if t.status not in ("planned", "active"):
        raise error(409, "invalid_transition")
    t.status, t.note = "cancelled", body.note.strip()
    await audit(db, request, "logistics.trip_cancel", actor=auth.user, entity="trip", entity_id=t.id, data={"note": t.note})
    await db.commit()
    return (await _trip_out_many(db, [t]))[0]


@router.get("/trips/{tid}/track")
async def trip_track(tid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Access)):
    t = await _trip(db, tid, auth)
    pts = await svc.vehicle_points(db, t.id)
    step = max(1, len(pts) // 800)
    return {"trip_id": t.id, "total": len(pts), "points": [{"lat": p.lat, "lng": p.lng, "speed": p.speed_kmh, "at": p.recorded_at} for p in pts[::step]]}


# --- GPS ----------------------------------------------------------------------------------


class PointIn(BaseModel):
    lat: float = Field(ge=-90, le=90)
    lng: float = Field(ge=-180, le=180)
    speed: float | None = Field(default=None, ge=0, le=500)  # km/soat
    accuracy: float | None = Field(default=None, ge=0)
    heading: float | None = Field(default=None, ge=0, le=360)
    at: datetime


class TrackIn(BaseModel):
    points: list[PointIn] = Field(min_length=1, max_length=svc.MAX_BATCH)


@router.post("/track")
async def track(body: TrackIn, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Driver)):
    """Haydovchi telefonidan GPS nuqtalar. Faqat faol reys davomida qabul qilinadi (reysdan tashqari kuzatuv yo'q)."""
    e = await _my_employee(db, auth)
    t = await db.scalar(select(Trip).where(Trip.driver_id == e.id, Trip.status == "active")) if e else None
    if not t:
        raise error(409, "no_active_trip")
    moment = now()
    existing = {r for (r,) in (await db.execute(select(GpsPoint.recorded_at).where(GpsPoint.trip_id == t.id, GpsPoint.recorded_at.in_([p.at for p in body.points])))).all()}
    saved = skipped = 0
    seen: set[datetime] = set()
    for p in sorted(body.points, key=lambda x: x.at):
        if p.at in existing or p.at in seen or p.at > moment + timedelta(minutes=2) or not svc.clean_point({"lat": p.lat, "lng": p.lng, "accuracy": p.accuracy}) or (t.started_at and p.at < t.started_at - timedelta(minutes=5)):
            skipped += 1
            continue
        seen.add(p.at)
        db.add(GpsPoint(vehicle_id=t.vehicle_id, trip_id=t.id, driver_id=t.driver_id, lat=p.lat, lng=p.lng, speed_kmh=p.speed, accuracy_m=p.accuracy, heading=p.heading, recorded_at=p.at))
        saved += 1
    await db.commit()
    return {"saved": saved, "skipped": skipped, "trip_id": t.id}


@router.get("/live")
async def live(db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    """Faol reysdagi mashinalarning oxirgi joylashuvi."""
    active = list((await db.scalars(select(Trip).where(Trip.status == "active"))).all())
    if not active:
        return []
    last = (
        select(GpsPoint.trip_id, func.max(GpsPoint.recorded_at).label("t")).where(GpsPoint.trip_id.in_([t.id for t in active])).group_by(GpsPoint.trip_id).subquery()
    )
    pts = {p.trip_id: p for p in (await db.scalars(select(GpsPoint).join(last, (GpsPoint.trip_id == last.c.trip_id) & (GpsPoint.recorded_at == last.c.t)))).all()}
    names = await _names(db, {t.driver_id for t in active})
    veh = {v.id: v for v in (await db.scalars(select(Vehicle).where(Vehicle.id.in_({t.vehicle_id for t in active})))).all()}
    moment = now()
    return [
        {
            "trip_id": t.id, "vehicle": {"id": t.vehicle_id, "plate": veh[t.vehicle_id].plate, "name": veh[t.vehicle_id].name}, "driver": {"id": t.driver_id, "name": names.get(t.driver_id)}, "purpose": t.purpose,
            "lat": pts[t.id].lat if t.id in pts else None, "lng": pts[t.id].lng if t.id in pts else None, "speed": pts[t.id].speed_kmh if t.id in pts else None,
            "at": pts[t.id].recorded_at if t.id in pts else None, "stale": t.id not in pts or (moment - pts[t.id].recorded_at) > timedelta(minutes=svc.STALE_MIN),
        }
        for t in active
    ]  # fmt: skip


@router.get("/my")
async def my(db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Driver)):
    """Haydovchi sahifasi: faol reys va rejalashtirilgan reyslar."""
    e = await _my_employee(db, auth)
    if not e:
        return {"employee": None, "active": None, "planned": []}
    rows = list((await db.scalars(select(Trip).where(Trip.driver_id == e.id, Trip.status.in_(("planned", "active"))).order_by(Trip.planned_start.nulls_last(), Trip.id))).all())
    out = await _trip_out_many(db, rows)
    veh = {v.id: v for v in (await db.scalars(select(Vehicle).where(Vehicle.id.in_({t.vehicle_id for t in rows})))).all()} if rows else {}
    for o, t in zip(out, rows, strict=True):
        o["vehicle"]["odometer_km"] = veh[t.vehicle_id].odometer_km
    return {"employee": {"id": e.id, "name": e.full_name}, "active": next((o for o in out if o["status"] == "active"), None), "planned": [o for o in out if o["status"] == "planned"]}


# --- Yoqilg'i -----------------------------------------------------------------------------


class FuelIn(BaseModel):
    vehicle_id: int | None = None
    trip_id: int | None = None
    fueled_at: datetime | None = None
    liters: Decimal = Field(gt=0, le=2000)
    price_per_liter: Decimal | None = Field(default=None, gt=0, le=1_000_000)
    odometer_km: int | None = Field(default=None, ge=0, le=5_000_000)
    station: str | None = Field(default=None, max_length=120)
    full_tank: bool = True


def _fuel_out(f: FuelLog, plates: dict[int, str], names: dict[int, str]) -> dict[str, Any]:
    return {
        "id": f.id, "vehicle": {"id": f.vehicle_id, "plate": plates.get(f.vehicle_id)}, "trip_id": f.trip_id, "driver": {"id": f.driver_id, "name": names.get(f.driver_id or 0)} if f.driver_id else None,
        "fueled_at": f.fueled_at, "liters": float(f.liters), "price_per_liter": float(f.price_per_liter) if f.price_per_liter is not None else None,
        "total": float(f.total) if f.total is not None else None, "odometer_km": f.odometer_km, "station": f.station, "full_tank": f.full_tank,
    }  # fmt: skip


@router.get("/fuel")
async def fuel_list(
    vehicle_id: int | None = None,
    start: date | None = None,
    end: date | None = None,
    page: int = Query(default=1, ge=1),
    size: int = Query(default=30, ge=1, le=100),
    db: AsyncSession = Depends(get_session),
    auth: AuthContext = Depends(Reader),
):
    F = FuelLog
    conds: list = []
    if vehicle_id:
        conds.append(F.vehicle_id == vehicle_id)
    if start:
        conds.append(F.fueled_at >= datetime.combine(start, datetime.min.time(), tzinfo=tz()))
    if end:
        conds.append(F.fueled_at < datetime.combine(end + timedelta(days=1), datetime.min.time(), tzinfo=tz()))
    total = await db.scalar(select(func.count()).select_from(F).where(*conds)) or 0
    rows = (await db.scalars(select(F).where(*conds).order_by(F.fueled_at.desc(), F.id.desc()).offset((page - 1) * size).limit(size))).all()
    plates = {v.id: v.plate for v in (await db.scalars(select(Vehicle))).all()}
    names = await _names(db, {f.driver_id for f in rows})
    return {"total": total, "page": page, "size": size, "items": [_fuel_out(f, plates, names) for f in rows]}


async def _fuel_vehicle(db: AsyncSession, body: FuelIn, auth: AuthContext) -> tuple[Vehicle, Trip | None, int | None]:
    """Boshqaruvchi: istalgan mashina. Haydovchi: faqat o'zining faol reysi mashinasi."""
    e = await _my_employee(db, auth)
    trip: Trip | None = None
    if svc.can_manage(auth):
        if body.trip_id:
            trip = await db.get(Trip, body.trip_id)
            if not trip:
                raise error(400, "unknown_trip")
        vid = body.vehicle_id or (trip.vehicle_id if trip else None)
        if not vid:
            raise error(400, "vehicle_required")
        driver = trip.driver_id if trip else None
    else:
        if not svc.has(auth, "logistics.own") or not e:
            raise error(403, "forbidden")
        trip = await db.scalar(select(Trip).where(Trip.driver_id == e.id, Trip.status == "active"))
        if not trip:
            raise error(409, "no_active_trip")
        vid, driver = trip.vehicle_id, e.id
    v = await db.get(Vehicle, vid)
    if not v:
        raise error(400, "unknown_vehicle")
    return v, trip, driver


@router.post("/fuel", status_code=201)
async def fuel_create(body: FuelIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Access)):
    v, trip, driver = await _fuel_vehicle(db, body, auth)
    moment = now()
    f = FuelLog(
        vehicle_id=v.id, trip_id=trip.id if trip else None, driver_id=driver, fueled_at=body.fueled_at or moment, liters=body.liters, price_per_liter=body.price_per_liter,
        total=svc.dec(body.liters * body.price_per_liter, "0.01") if body.price_per_liter else None, odometer_km=body.odometer_km, station=(body.station or "").strip() or None, full_tank=body.full_tank, created_by=auth.user.id,
    )  # fmt: skip
    if f.odometer_km:
        v.odometer_km = max(v.odometer_km, f.odometer_km)
    db.add(f)
    await db.flush()
    await audit(db, request, "logistics.fuel_create", actor=auth.user, entity="fuel_log", entity_id=f.id, data={"vehicle_id": v.id, "liters": str(f.liters), "total": str(f.total)})
    await db.commit()
    return _fuel_out(f, {v.id: v.plate}, {})


async def _fuel(db: AsyncSession, fid: int) -> FuelLog:
    f = await db.get(FuelLog, fid)
    if not f:
        raise error(404, "not_found")
    return f


@router.get("/fuel/{fid}")
async def fuel_one(fid: int, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    f = await _fuel(db, fid)
    v = await db.get(Vehicle, f.vehicle_id)
    return _fuel_out(f, {f.vehicle_id: v.plate if v else ""}, await _names(db, {f.driver_id}))


@router.put("/fuel/{fid}")
async def fuel_update(fid: int, body: FuelIn, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    f = await _fuel(db, fid)
    if body.vehicle_id and body.vehicle_id != f.vehicle_id:
        if not await db.get(Vehicle, body.vehicle_id):
            raise error(400, "unknown_vehicle")
        f.vehicle_id = body.vehicle_id
    f.fueled_at = body.fueled_at or f.fueled_at
    f.liters, f.price_per_liter, f.odometer_km, f.station, f.full_tank = body.liters, body.price_per_liter, body.odometer_km, (body.station or "").strip() or None, body.full_tank
    f.total = svc.dec(body.liters * body.price_per_liter, "0.01") if body.price_per_liter else None
    await audit(db, request, "logistics.fuel_update", actor=auth.user, entity="fuel_log", entity_id=f.id, data={"vehicle_id": f.vehicle_id, "liters": str(f.liters)})
    await db.commit()
    v = await db.get(Vehicle, f.vehicle_id)
    return _fuel_out(f, {f.vehicle_id: v.plate if v else ""}, await _names(db, {f.driver_id}))


@router.delete("/fuel/{fid}", status_code=204)
async def fuel_delete(fid: int, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Manager)):
    f = await _fuel(db, fid)
    await audit(db, request, "logistics.fuel_delete", actor=auth.user, entity="fuel_log", entity_id=f.id, data={"vehicle_id": f.vehicle_id, "liters": str(f.liters)})
    await db.delete(f)
    await db.commit()


# --- Hisobotlar ---------------------------------------------------------------------------


def _range(start: date, end: date) -> tuple[datetime, datetime]:
    if end < start or (end - start).days > 366:
        raise error(400, "range_invalid", max_days=366)
    return datetime.combine(start, datetime.min.time(), tzinfo=tz()), datetime.combine(end + timedelta(days=1), datetime.min.time(), tzinfo=tz())


async def _fuel_report(db: AsyncSession, start: date, end: date) -> dict[str, Any]:
    s_dt, e_dt = _range(start, end)
    trips_ = (await db.execute(select(Trip.vehicle_id, func.count(), func.coalesce(func.sum(Trip.distance_km), 0), func.coalesce(func.sum(Trip.gps_km), 0), func.coalesce(func.sum(Trip.fuel_norm_l), 0)).where(Trip.status == "completed", Trip.finished_at >= s_dt, Trip.finished_at < e_dt).group_by(Trip.vehicle_id))).all()
    fuel = (await db.execute(select(FuelLog.vehicle_id, func.count(), func.coalesce(func.sum(FuelLog.liters), 0), func.coalesce(func.sum(FuelLog.total), 0)).where(FuelLog.fueled_at >= s_dt, FuelLog.fueled_at < e_dt).group_by(FuelLog.vehicle_id))).all()
    veh = {v.id: v for v in (await db.scalars(select(Vehicle))).all()}
    by: dict[int, dict[str, Any]] = {}
    for vid, n, km, gps, norm in trips_:
        by.setdefault(vid, {})["trips"] = (n, float(km), float(gps), float(norm))
    for vid, n, liters, cost in fuel:
        by.setdefault(vid, {})["fuel"] = (n, float(liters), float(cost))
    rows = []
    for vid, d in by.items():
        v = veh.get(vid)
        n, km, gps, norm = d.get("trips", (0, 0.0, 0.0, 0.0))
        fn, liters, cost = d.get("fuel", (0, 0.0, 0.0))
        rows.append({
            "vehicle": {"id": vid, "plate": v.plate if v else "—", "name": v.name if v else ""}, "norm_per_100": float(v.fuel_norm) if v else 0, "trips": n, "km": round(km, 1), "gps_km": round(gps, 1),
            "norm_liters": round(norm, 2), "fills": fn, "liters": round(liters, 2), "cost": round(cost), "variance": round(liters - norm, 2), "actual_per_100": round(liters / km * 100, 2) if km > 0 and liters else None,
        })  # fmt: skip
    rows.sort(key=lambda r: r["vehicle"]["plate"])
    tot = {k: round(sum(r[k] for r in rows), 2) for k in ("trips", "km", "norm_liters", "fills", "liters", "cost", "variance")}
    return {"start": start, "end": end, "items": rows, "totals": tot}


@router.get("/reports/fuel")
async def report_fuel(start: date, end: date, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    return await _fuel_report(db, start, end)


def _fuel_xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Yoqilg'i"
    ws.append(["Davlat raqami", "Mashina", "Me'yor (l/100km)", "Reyslar", "Masofa (km)", "GPS (km)", "Me'yor bo'yicha (l)", "Quyilgan (l)", "Summa", "Farq (l)", "Amalda (l/100km)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in data["items"]:
        ws.append([r["vehicle"]["plate"], r["vehicle"]["name"], r["norm_per_100"], r["trips"], r["km"], r["gps_km"], r["norm_liters"], r["liters"], r["cost"], r["variance"], r["actual_per_100"]])
    n = len(data["items"]) + 1
    ws.append(["JAMI", "", "", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})", f"=SUM(G2:G{n})", f"=SUM(H2:H{n})", f"=SUM(I2:I{n})", f"=SUM(J2:J{n})", ""])
    for c in ws[n + 1]:
        c.font = Font(bold=True)
    ws.column_dimensions["B"].width = 28
    buf = io.BytesIO()
    harden(wb)
    wb.save(buf)
    return buf.getvalue()


@router.get("/reports/fuel/export")
async def report_fuel_export(start: date, end: date, request: Request, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    data = await _fuel_report(db, start, end)
    content = await asyncio.to_thread(_fuel_xlsx, data)
    await audit(db, request, "logistics.export", actor=auth.user, entity="logistics", data={"start": str(start), "end": str(end)})
    await db.commit()
    return Response(content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="yoqilgi-{start}-{end}.xlsx"'})


@router.get("/reports/trips")
async def report_trips(start: date, end: date, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    s_dt, e_dt = _range(start, end)
    T = Trip
    rows = (await db.execute(select(T.driver_id, T.status, func.count(), func.coalesce(func.sum(T.distance_km), 0)).where(func.coalesce(T.started_at, T.planned_start, T.created_at) >= s_dt, func.coalesce(T.started_at, T.planned_start, T.created_at) < e_dt).group_by(T.driver_id, T.status))).all()
    names = await _names(db, {r[0] for r in rows})
    agg: dict[int, dict[str, Any]] = {}
    for did, st, n, km in rows:
        a = agg.setdefault(did, {"driver": {"id": did, "name": names.get(did)}, "completed": 0, "cancelled": 0, "open": 0, "km": 0.0})
        if st == "completed":
            a["completed"] += n
            a["km"] += float(km)
        elif st == "cancelled":
            a["cancelled"] += n
        else:
            a["open"] += n
    return {"start": start, "end": end, "items": sorted(agg.values(), key=lambda a: -a["km"])}


@router.get("/reports/driver-days")
async def report_driver_days(start: date, end: date, db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    """GPS ish kuni: haydovchining kun bo'yicha birinchi/oxirgi signali, yo'ldagi vaqt va reyslar."""
    s_dt, e_dt = _range(start, end)
    day = func.date(func.timezone(literal_column("'Asia/Tashkent'"), GpsPoint.recorded_at))
    rows = (await db.execute(select(GpsPoint.driver_id, day.label("d"), func.min(GpsPoint.recorded_at), func.max(GpsPoint.recorded_at), func.count(), func.count(func.distinct(GpsPoint.trip_id))).where(GpsPoint.recorded_at >= s_dt, GpsPoint.recorded_at < e_dt).group_by(GpsPoint.driver_id, day).order_by(day.desc()))).all()
    names = await _names(db, {r[0] for r in rows})
    km = {}
    for did, finished, dist in (await db.execute(select(Trip.driver_id, func.date(func.timezone(literal_column("'Asia/Tashkent'"), Trip.finished_at)), Trip.distance_km).where(Trip.status == "completed", Trip.finished_at >= s_dt, Trip.finished_at < e_dt))).all():
        km[(did, finished)] = km.get((did, finished), 0.0) + float(dist or 0)
    return {"start": start, "end": end, "items": [
        {"driver": {"id": r[0], "name": names.get(r[0])}, "date": r[1], "first": r[2], "last": r[3], "hours": round((r[3] - r[2]).total_seconds() / 3600, 1), "points": r[4], "trips": r[5], "km": round(km.get((r[0], r[1]), 0.0), 1)} for r in rows
    ]}  # fmt: skip


@router.get("/reports/expiry")
async def report_expiry(days: int = Query(default=30, ge=1, le=365), db: AsyncSession = Depends(get_session), auth: AuthContext = Depends(Reader)):
    """Muddati o'tgan yoki `days` kun ichida tugaydigan hujjatlar va guvohnomalar."""
    today = local_date(now())
    limit = today + timedelta(days=days)
    out: list[dict[str, Any]] = []
    plates = {v.id: v.plate for v in (await db.scalars(select(Vehicle))).all()}
    for d in (await db.scalars(select(VehicleDocument).where(VehicleDocument.expires_at <= limit).order_by(VehicleDocument.expires_at))).all():
        out.append({"type": "document", "kind": d.kind, "subject": plates.get(d.vehicle_id), "link": f"/logistics/vehicles/{d.vehicle_id}", "expires_at": d.expires_at, "days_left": (d.expires_at - today).days})
    for p, e in (await db.execute(select(DriverProfile, Employee).join(Employee, Employee.id == DriverProfile.employee_id).where(DriverProfile.license_expires.is_not(None), DriverProfile.license_expires <= limit))).all():
        out.append({"type": "license", "kind": "license", "subject": e.full_name, "link": f"/logistics/drivers/{e.id}/edit", "expires_at": p.license_expires, "days_left": (p.license_expires - today).days})  # type: ignore[operator]
    out.sort(key=lambda r: r["expires_at"])
    return {"days": days, "items": out}
