import io
from datetime import UTC, date, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from app import worker
from app.db import SessionLocal
from app.models.auth import User
from app.models.reports import ReportRun, ReportSchedule
from app.services import reports as svc
from app.timeutil import tz
from tests.conftest import FakeStorage, login, make_user

pytestmark = pytest.mark.skipif(
    "not __import__('os').environ.get('DATABASE_URL')", reason="DATABASE_URL kerak"
)

R = "/api/v1/reports"
NOW = datetime(2026, 11, 10, 12, 0, tzinfo=tz())


@pytest.fixture
def clock(monkeypatch):
    fn = lambda: NOW  # noqa: E731
    monkeypatch.setattr("app.api.reports.now", fn)
    return fn


async def as_user(new_client, name, role):
    uid = await make_user(name, role)
    c = new_client()
    assert (await login(c, name)).status_code == 200
    return c, uid


async def setup(new_client):
    hr, hr_id = await as_user(new_client, "hr", "hr_direktor")
    ali, ali_id = await as_user(new_client, "ali", "farmasevt")
    dr, dr_id = await as_user(new_client, "dir", "direktor")
    return {"hr": hr, "hr_id": hr_id, "ali": ali, "ali_id": ali_id, "dir": dr, "dir_id": dr_id}


def sched_body(ctx, **kw):
    return {"name": "Haftalik davomat", "report_type": "attendance", "period": "last_week", "format": "xlsx", "frequency": "weekly", "weekday": 0, "hour": 8, "minute": 30, "recipient_ids": [ctx["dir_id"]], "channels": ["app"], **kw}


def test_resolve_period_and_next_run():
    today = date(2026, 11, 11)  # chorshanba
    assert svc.resolve_period("yesterday", today) == (date(2026, 11, 10), date(2026, 11, 10))
    assert svc.resolve_period("last_week", today) == (date(2026, 11, 2), date(2026, 11, 8))
    assert svc.resolve_period("last_7d", today) == (date(2026, 11, 4), date(2026, 11, 10))
    assert svc.resolve_period("last_month", today) == (date(2026, 10, 1), date(2026, 10, 31))
    assert svc.resolve_period("this_month", today) == (date(2026, 11, 1), date(2026, 11, 10))
    assert svc.resolve_period("last_month", date(2026, 1, 5)) == (date(2025, 12, 1), date(2025, 12, 31))
    s = ReportSchedule(frequency="daily", hour=8, minute=0)
    after = datetime(2026, 11, 10, 9, 0, tzinfo=tz())
    assert svc.next_run(s, after).astimezone(tz()) == datetime(2026, 11, 11, 8, 0, tzinfo=tz())
    assert svc.next_run(s, datetime(2026, 11, 10, 7, 0, tzinfo=tz())).astimezone(tz()) == datetime(2026, 11, 10, 8, 0, tzinfo=tz())
    s = ReportSchedule(frequency="weekly", weekday=0, hour=8, minute=0)  # dushanba
    assert svc.next_run(s, after).astimezone(tz()) == datetime(2026, 11, 16, 8, 0, tzinfo=tz())
    s = ReportSchedule(frequency="monthly", month_day=1, hour=6, minute=0)
    assert svc.next_run(s, after).astimezone(tz()) == datetime(2026, 12, 1, 6, 0, tzinfo=tz())
    assert svc.next_run(s, datetime(2026, 12, 20, tzinfo=tz())).astimezone(tz()) == datetime(2027, 1, 1, 6, 0, tzinfo=tz())


async def test_catalog_is_permission_filtered(new_client):
    ctx = await setup(new_client)
    assert (await ctx["ali"].get(f"{R}/types")).status_code == 403
    hr = (await ctx["hr"].get(f"{R}/types")).json()
    assert {t["type"] for t in hr["types"]} == set(svc.CATALOG)
    dr = {t["type"] for t in (await ctx["dir"].get(f"{R}/types")).json()["types"]}
    assert "attendance" in dr and "headcount" in dr and dr <= set(svc.CATALOG)
    assert svc.allowed(frozenset({"reports.read"}), "tasks") is False and svc.allowed(frozenset({"tasks.read"}), "tasks") is True
    assert (await ctx["hr"].post(f"{R}/run", json={"type": "nope"})).json()["detail"]["code"] == "unknown_report"


async def test_run_exports_xlsx_pdf_and_download_rules(new_client, clock):
    ctx = await setup(new_client)
    for fmt, marker in (("xlsx", b"PK"), ("pdf", b"%PDF")):
        r = await ctx["hr"].post(f"{R}/run", json={"type": "tasks", "start": "2026-11-01", "end": "2026-11-07", "format": fmt, "lang": "ru"})
        assert r.status_code == 201, r.text
        run = r.json()
        assert run["status"] == "ok" and run["filename"].endswith(fmt) and run["size"] > 0
        f = await ctx["hr"].get(f"{R}/runs/{run['id']}/file")
        assert f.status_code == 200 and f.content.startswith(marker)
        assert f.headers["content-type"].startswith("application/")
    # xlsx mazmuni
    run = (await ctx["hr"].post(f"{R}/run", json={"type": "headcount", "format": "xlsx"})).json()
    wb = load_workbook(io.BytesIO((await ctx["hr"].get(f"{R}/runs/{run['id']}/file")).content))
    texts = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value is not None]
    assert any("Faol xodimlar" in t for t in texts) and any("Filial" in t for t in texts)
    # begona yuklay olmaydi
    assert (await ctx["dir"].get(f"{R}/runs/{run['id']}/file")).status_code == 404
    # diapazon va talablar
    assert (await ctx["hr"].post(f"{R}/run", json={"type": "tasks"})).json()["detail"]["code"] == "range_required"
    r = await ctx["hr"].post(f"{R}/run", json={"type": "tasks", "start": "2026-01-01", "end": "2026-11-01"})
    assert r.status_code == 400 and r.json()["detail"]["code"] == "range_invalid"
    # tarix
    mine = (await ctx["hr"].get(f"{R}/runs")).json()
    assert mine["total"] == 3 and (await ctx["dir"].get(f"{R}/runs")).json()["total"] == 0
    assert (await ctx["dir"].get(f"{R}/runs", params={"scope": "all"})).status_code == 403
    assert (await ctx["hr"].get(f"{R}/runs", params={"scope": "all", "type": "headcount"})).json()["total"] == 1


async def test_every_report_type_builds(new_client, clock):
    ctx = await setup(new_client)
    for rtype in svc.CATALOG:
        r = await ctx["hr"].post(f"{R}/run", json={"type": rtype, "start": "2026-11-01", "end": "2026-11-07", "format": "xlsx"})
        assert r.status_code == 201 and r.json()["status"] == "ok", (rtype, r.text)


async def test_schedule_crud_validation_and_permissions(new_client, clock):
    ctx = await setup(new_client)
    assert (await ctx["dir"].get(f"{R}/schedules")).status_code == 403
    assert (await ctx["dir"].post(f"{R}/schedules", json=sched_body(ctx))).status_code == 403
    r = await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx))
    assert r.status_code == 201, r.text
    s = r.json()
    assert s["next_run_at"] and datetime.fromisoformat(s["next_run_at"]).astimezone(tz()).weekday() == 0
    assert s["recipients"][0]["name"] == "Dir" and s["recipients"][0]["telegram"] is False
    # tekshiruvlar
    assert (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, weekday=None))).status_code == 422
    assert (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, frequency="monthly", month_day=None))).status_code == 422
    assert (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, channels=["sms"]))).status_code == 422
    assert (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, recipient_ids=[]))).status_code == 422
    assert (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, recipient_ids=[999999]))).json()["detail"]["code"] == "unknown_user"
    assert (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, report_type="nope"))).json()["detail"]["code"] == "unknown_report"

    r = await ctx["hr"].put(f"{R}/schedules/{s['id']}", json=sched_body(ctx, name="Kunlik", frequency="daily", weekday=None, is_active=False))
    assert r.json()["name"] == "Kunlik" and r.json()["next_run_at"] is None and r.json()["weekday"] is None
    assert (await ctx["hr"].get(f"{R}/schedules/{s['id']}")).json()["is_active"] is False
    assert len((await ctx["hr"].get(f"{R}/schedules")).json()) == 1
    assert (await ctx["hr"].delete(f"{R}/schedules/{s['id']}")).status_code == 204
    assert (await ctx["hr"].get(f"{R}/schedules/{s['id']}")).status_code == 404


async def test_run_now_delivers_notification_and_telegram(new_client, clock, app):
    ctx = await setup(new_client)
    async with SessionLocal() as db:
        await db.execute(update(User).where(User.id == ctx["dir_id"]).values(telegram_chat_id=777001))
        await db.commit()
    s = (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, report_type="headcount", period="yesterday", channels=["app", "telegram"]))).json()
    assert s["recipients"][0]["telegram"] is True
    r = await ctx["hr"].post(f"{R}/schedules/{s['id']}/run")
    assert r.status_code == 200 and r.json()["status"] == "ok" and r.json()["schedule"]["name"] == "Haftalik davomat"
    run_id = r.json()["id"]
    kinds = [n["kind"] for n in (await ctx["dir"].get("/api/v1/notifications")).json()["items"]]
    assert "reports.ready" in kinds
    assert ("deliver_report_telegram", run_id, ctx["dir_id"]) in app.state.queue.calls
    # qabul qiluvchi yuklay oladi va tarixida ko'radi; boshqa foydalanuvchi yo'q
    assert (await ctx["dir"].get(f"{R}/runs/{run_id}/file")).status_code == 200
    assert (await ctx["dir"].get(f"{R}/runs")).json()["total"] == 1
    # app-only kanal: telegram job yo'q
    s2 = (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, report_type="headcount", period="yesterday"))).json()
    before = len(app.state.queue.calls)
    await ctx["hr"].post(f"{R}/schedules/{s2['id']}/run")
    assert not [c for c in app.state.queue.calls[before:] if c[0] == "deliver_report_telegram"]


class Redis:
    def __init__(self):
        self.jobs: list[tuple] = []

    async def enqueue_job(self, fn, *args):
        self.jobs.append((fn, *args))


async def test_worker_scheduler_runs_due_failures_and_purge(new_client, clock, monkeypatch):
    ctx = await setup(new_client)
    s_ok = (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, report_type="headcount", period="yesterday", frequency="daily", weekday=None))).json()
    s_off = (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, name="Faol emas", report_type="headcount", frequency="daily", weekday=None, is_active=False))).json()
    s_bad = (await ctx["hr"].post(f"{R}/schedules", json=sched_body(ctx, name="Buzuq", report_type="visits", period="yesterday", frequency="daily", weekday=None))).json()
    past = datetime.now(UTC) - timedelta(minutes=5)
    async with SessionLocal() as db:
        await db.execute(update(ReportSchedule).where(ReportSchedule.id.in_([s_ok["id"], s_bad["id"]])).values(next_run_at=past))
        await db.execute(update(ReportSchedule).where(ReportSchedule.id == s_off["id"]).values(next_run_at=past))
        await db.commit()

    real = svc.build_table

    async def flaky(db, rtype, start, end):
        if rtype == "visits":
            raise RuntimeError("manba ishlamadi")
        return await real(db, rtype, start, end)

    monkeypatch.setattr(svc, "build_table", flaky)
    storage, redis = FakeStorage(), Redis()
    out = await worker.report_scheduler({"storage": storage, "redis": redis})
    assert out["ran"] == 2  # faol emas jadval ishlamaydi
    async with SessionLocal() as db:
        runs = {r.report_type: r for r in (await db.scalars(select(ReportRun))).all()}
        scheds = {s.id: s for s in (await db.scalars(select(ReportSchedule))).all()}
    assert runs["headcount"].status == "ok" and runs["headcount"].file_key in storage.objects
    assert runs["visits"].status == "failed" and "manba ishlamadi" in runs["visits"].error and runs["visits"].file_key is None
    assert scheds[s_ok["id"]].last_status == "ok" and scheds[s_bad["id"]].last_status == "failed"
    assert scheds[s_ok["id"]].next_run_at > datetime.now(UTC) and scheds[s_off["id"]].last_run_at is None
    # qayta chaqirilsa takror ishlamaydi
    assert (await worker.report_scheduler({"storage": storage, "redis": redis}))["ran"] == 0
    # muvaffaqiyatsiz jadval egasiga xabar, qabul qiluvchiga faqat muvaffaqiyatli hisobot
    assert "reports.failed" in [n["kind"] for n in (await ctx["hr"].get("/api/v1/notifications")).json()["items"]]
    assert "reports.ready" in [n["kind"] for n in (await ctx["dir"].get("/api/v1/notifications")).json()["items"]]
    assert (await ctx["hr"].get(f"{R}/runs/{runs['visits'].id}/file")).json()["detail"]["code"] == "report_failed"

    # 90 kundan eski fayl o'chadi
    async with SessionLocal() as db:
        await db.execute(update(ReportRun).where(ReportRun.id == runs["headcount"].id).values(created_at=datetime.now(UTC) - timedelta(days=91)))
        await db.commit()
    async with SessionLocal() as db:
        assert await svc.purge_runs(db, storage, datetime.now(UTC)) == 1
        await db.commit()
    assert runs["headcount"].file_key not in storage.objects
