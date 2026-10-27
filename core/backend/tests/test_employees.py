import io
from datetime import date, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import login, make_user

pytestmark = pytest.mark.skipif(
    "not __import__('os').environ.get('DATABASE_URL')", reason="DATABASE_URL kerak"
)


async def as_user(new_client, login_name, role, **kw):
    uid = await make_user(login_name, role, **kw)
    c = new_client()
    assert (await login(c, login_name)).status_code == 200
    return c, uid


async def org(c):
    """Bo'lim, lavozim va filial yaratadi."""
    d = (await c.post("/api/v1/org/departments", json={"name": "Moliya"})).json()["id"]
    d2 = (await c.post("/api/v1/org/departments", json={"name": "Logistika"})).json()["id"]
    pos = (await c.get("/api/v1/org/positions", params={"department_id": d})).json()
    pos2 = (await c.get("/api/v1/org/positions", params={"department_id": d2})).json()
    b = (await c.post("/api/v1/branches", json={"number": 16, "name": "Filial 16"})).json()["id"]
    return {"dept": d, "dept2": d2, "pos": pos[1]["id"], "pos2": pos2[1]["id"], "branch": b}


def person(**kw):
    base = {"last_name": "Aliyev", "first_name": "Vali", "phone": "90 123 45 67"}
    return {**base, **kw}


async def hr(new_client):
    return (await as_user(new_client, "hr", "hr_direktor"))[0]


async def create(c, **kw):
    r = await c.post("/api/v1/employees", json=person(**kw))
    assert r.status_code == 201, r.text
    return r.json()["employee"]


# --- Ruxsatlar -----------------------------------------------------------------------


async def test_permissions(new_client):
    kor, _ = await as_user(new_client, "k", "koordinator")
    assert (await kor.get("/api/v1/employees")).status_code == 200
    assert (await kor.post("/api/v1/employees", json=person())).status_code == 403
    mudir, _ = await as_user(new_client, "m", "mudir")
    assert (await mudir.get("/api/v1/employees")).status_code == 403
    assert (await mudir.get("/api/v1/employees/duplicates")).status_code == 403


# --- Yaratish va normallashtirish ----------------------------------------------------


async def test_create_normalizes_and_numbers_employees(new_client):
    c = await hr(new_client)
    o = await org(c)
    e = await create(
        c, middle_name="Karim o'g'li", phone="+998 (90) 123-45-67", pinfl="12345678901234",
        passport_no="aa 1234567", email="  Vali@Example.UZ ", birth_date="1990-03-15", gender="m",
        department_id=o["dept"], position_id=o["pos"], branch_id=o["branch"],
    )  # fmt: skip
    assert e["full_name"] == "Aliyev Vali Karim o'g'li"
    assert e["phone"] == "+998901234567" and e["passport_no"] == "AA1234567"
    assert e["email"] == "vali@example.uz" and e["status"] == "active"
    assert e["employee_no"] >= 1 and e["hire_date"] == str(date.today())
    assert e["department"]["name"] == "Moliya" and e["branch"]["number"] == 16
    second = await create(c, last_name="Karimov", phone="+998 91 000 00 01")
    assert second["employee_no"] > e["employee_no"]


@pytest.mark.parametrize(
    "field,value,code",
    [
        ("phone", "123", "phone_invalid"),
        ("pinfl", "123", "pinfl_invalid"),
        ("passport_no", "1234567", "passport_invalid"),
        ("email", "not-an-email", "email_invalid"),
        ("birth_date", "2020-01-01", "birth_date_invalid"),
        ("birth_date", "1800-01-01", "birth_date_invalid"),
    ],
)
async def test_field_validation_returns_field_codes(new_client, field, value, code):
    c = await hr(new_client)
    r = await c.post("/api/v1/employees", json=person(**{field: value}))
    assert r.status_code == 422
    detail = r.json()["detail"]
    assert detail["code"] == "validation"
    assert {"field": field, "code": code} in detail["fields"]


async def test_assignment_validation(new_client):
    c = await hr(new_client)
    o = await org(c)
    # Lavozim boshqa bo'limniki
    r = await c.post("/api/v1/employees", json=person(department_id=o["dept"], position_id=o["pos2"]))
    assert r.status_code == 400 and r.json()["detail"]["code"] == "position_department_mismatch"
    for body, code in (
        ({"department_id": 9999}, "unknown_department"),
        ({"position_id": 9999}, "unknown_position"),
        ({"branch_id": 9999}, "unknown_branch"),
    ):
        r = await c.post("/api/v1/employees", json=person(**body))
        assert r.status_code == 400 and r.json()["detail"]["code"] == code


# --- Dublikat nazorati ---------------------------------------------------------------


async def test_hard_duplicates_by_pinfl_and_passport_cannot_be_forced(new_client):
    c = await hr(new_client)
    await create(c, pinfl="12345678901234", passport_no="AA1234567", phone="+998900000001")
    for extra in ({"pinfl": "12345678901234"}, {"passport_no": "AA1234567"}):
        r = await c.post("/api/v1/employees", json=person(last_name="Boshqa", phone="+998900000002", confirm_duplicate=True, **extra))
        assert r.status_code == 409 and r.json()["detail"]["code"] == "duplicate_identifier"
        assert r.json()["detail"]["matches"][0]["reasons"]


async def test_soft_duplicates_need_confirmation(new_client):
    c = await hr(new_client)
    first = await create(c, phone="+998901112233", birth_date="1990-01-01")
    # Telefon bir xil
    r = await c.post("/api/v1/employees", json=person(last_name="Boshqa", phone="+998901112233"))
    assert r.status_code == 409 and r.json()["detail"]["code"] == "possible_duplicate"
    assert r.json()["detail"]["matches"][0]["id"] == first["id"]
    assert r.json()["detail"]["matches"][0]["reasons"] == ["phone"]
    ok = await c.post("/api/v1/employees", json=person(last_name="Boshqa", phone="+998901112233", confirm_duplicate=True))
    assert ok.status_code == 201
    # Ism + tug'ilgan sana bir xil (telefon boshqa)
    r = await c.post("/api/v1/employees", json=person(phone="+998907776655", birth_date="1990-01-01"))
    assert r.json()["detail"]["matches"][0]["reasons"] == ["name_birth"]


async def test_update_duplicate_check_ignores_self_and_preexisting(new_client):
    c = await hr(new_client)
    a = await create(c, phone="+998901000001")
    b = await create(c, last_name="Boshqa", phone="+998901000002")
    # O'zining telefonini qayta saqlash xato emas
    assert (await c.patch(f"/api/v1/employees/{a['id']}", json={"phone": "+998901000001"})).status_code == 200
    # Boshqaning telefoniga o'zgartirish
    r = await c.patch(f"/api/v1/employees/{b['id']}", json={"phone": "+998901000001"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "possible_duplicate"
    ok = await c.patch(f"/api/v1/employees/{b['id']}", json={"phone": "+998901000001", "confirm_duplicate": True})
    assert ok.status_code == 200


# --- Sezgir ma'lumotlar --------------------------------------------------------------


async def test_sensitive_fields_are_masked_without_permission(new_client):
    boss = await hr(new_client)
    e = await create(boss, pinfl="12345678901234", passport_no="AA1234567", birth_date="1990-03-15", address="Toshkent")
    # Direktor: employees.read bor, sensitive yo'q
    d, _ = await as_user(new_client, "dir", "direktor")
    got = (await d.get(f"/api/v1/employees/{e['id']}")).json()
    assert got["sensitive"] is False and got["birth_date"] is None and got["address"] is None
    assert got["pinfl"] == "123••••••••234" and got["passport_no"] == "AA•••567"
    full = (await boss.get(f"/api/v1/employees/{e['id']}")).json()
    assert full["pinfl"] == "12345678901234" and full["birth_date"] == "1990-03-15"


async def test_sensitive_fields_cannot_be_written_without_permission(new_client):
    boss = await hr(new_client)
    e = await create(boss)
    # employees.update bor, lekin employees.sensitive yo'q rol yarataylik
    roles = {r["key"]: r for r in (await boss.get("/api/v1/roles")).json()}
    await boss.put(f"/api/v1/roles/{roles['koordinator']['id']}/permissions", json={"permissions": ["employees.read", "employees.update", "employees.create"]})
    kor, _ = await as_user(new_client, "kor", "koordinator")
    assert (await kor.patch(f"/api/v1/employees/{e['id']}", json={"pinfl": "98765432109876"})).status_code == 403
    assert (await kor.post("/api/v1/employees", json=person(last_name="X", phone="+998900001111", passport_no="AB7654321"))).status_code == 403
    assert (await kor.patch(f"/api/v1/employees/{e['id']}", json={"email": "ok@example.uz"})).status_code == 200


# --- Yangilash -----------------------------------------------------------------------


async def test_update_recomputes_full_name_and_syncs_account(new_client):
    c = await hr(new_client)
    e = await create(c)
    acc = (await c.post(f"/api/v1/employees/{e['id']}/account", json={"role_key": "farmasevt"})).json()
    r = await c.patch(f"/api/v1/employees/{e['id']}", json={"last_name": "Yangiyev", "phone": "+998903334455"})
    assert r.json()["full_name"] == "Yangiyev Vali" and r.json()["phone"] == "+998903334455"
    users = (await c.get("/api/v1/users", params={"q": acc["login"]})).json()["items"]
    assert users[0]["full_name"] == "Yangiyev Vali" and users[0]["phone"] == "+998903334455"
    assert (await c.patch(f"/api/v1/employees/{e['id']}", json={"last_name": ""})).status_code in (400, 422)


# --- Tarix, o'tkazish, bo'shatish ----------------------------------------------------


async def test_transfer_records_history_and_notifies(new_client):
    c = await hr(new_client)
    o = await org(c)
    e = await create(c, department_id=o["dept"], position_id=o["pos"], hire_date=str(date.today() - timedelta(days=30)))
    acc = (await c.post(f"/api/v1/employees/{e['id']}/account", json={"role_key": "farmasevt"})).json()

    r = await c.post(f"/api/v1/employees/{e['id']}/transfer", json={
        "department_id": o["dept2"], "position_id": o["pos2"], "branch_id": o["branch"],
        "effective_date": str(date.today()), "note": "Ehtiyoj",
    })  # fmt: skip
    assert r.status_code == 200
    assert r.json()["department"]["name"] == "Logistika" and r.json()["branch"]["number"] == 16

    events = (await c.get(f"/api/v1/employees/{e['id']}/events")).json()
    assert [ev["kind"] for ev in events] == ["transfer", "hired"]
    ch = events[0]["data"]["changes"]
    assert ch["department"]["from"]["name"] == "Moliya" and ch["department"]["to"]["name"] == "Logistika"
    assert events[0]["note"] == "Ehtiyoj" and events[0]["actor"] == "Hr"

    # Hisobga bildirishnoma ketgan
    user = await new_client_login(new_client, acc["login"], acc["temp_password"])
    assert user is not None


async def new_client_login(new_client, login_name, password):
    c = new_client()
    r = await c.post("/api/auth/login", json={"login": login_name, "password": password})
    return r.json() if r.status_code == 200 else None


async def test_transfer_validation(new_client):
    c = await hr(new_client)
    o = await org(c)
    e = await create(c, department_id=o["dept"], position_id=o["pos"], hire_date=str(date.today()))
    url = f"/api/v1/employees/{e['id']}/transfer"
    today = str(date.today())
    r = await c.post(url, json={"department_id": o["dept"], "effective_date": today})
    assert r.status_code == 400 and r.json()["detail"]["code"] == "no_changes"
    # Bo'lim o'zgardi, lekin lavozim eski bo'limniki qoldi
    r = await c.post(url, json={"department_id": o["dept2"], "effective_date": today})
    assert r.status_code == 400 and r.json()["detail"]["code"] == "position_department_mismatch"
    r = await c.post(url, json={"branch_id": o["branch"], "effective_date": str(date.today() - timedelta(days=5))})
    assert r.status_code == 400 and r.json()["detail"]["code"] == "date_before_hire"
    # null yuborib filialni olib tashlash mumkin
    await c.post(url, json={"branch_id": o["branch"], "effective_date": today})
    r = await c.post(url, json={"branch_id": None, "effective_date": today})
    assert r.status_code == 200 and r.json()["branch"] is None


async def test_dismiss_blocks_account_and_reinstate(new_client):
    c = await hr(new_client)
    e = await create(c, hire_date=str(date.today() - timedelta(days=10)))
    acc = (await c.post(f"/api/v1/employees/{e['id']}/account", json={"role_key": "farmasevt"})).json()
    emp_client = new_client()
    first = await emp_client.post("/api/auth/login", json={"login": acc["login"], "password": acc["temp_password"]})
    assert first.status_code == 200

    url = f"/api/v1/employees/{e['id']}"
    bad = await c.post(f"{url}/dismiss", json={"date": str(date.today()), "reason": "other"})
    assert bad.json()["detail"]["code"] == "note_required"
    future = await c.post(f"{url}/dismiss", json={"date": str(date.today() + timedelta(days=1)), "reason": "own_will"})
    assert future.json()["detail"]["code"] == "date_in_future"
    early = await c.post(f"{url}/dismiss", json={"date": str(date.today() - timedelta(days=11)), "reason": "own_will"})
    assert early.json()["detail"]["code"] == "date_before_hire"

    r = await c.post(f"{url}/dismiss", json={"date": str(date.today()), "reason": "own_will", "note": "Arizasi bo'yicha"})
    assert r.status_code == 200 and r.json()["status"] == "dismissed" and r.json()["user"]["is_active"] is False
    assert (await emp_client.get("/api/auth/me")).status_code == 401  # sessiyalari yopilgan
    again = await c.post(f"{url}/dismiss", json={"date": str(date.today()), "reason": "own_will"})
    assert again.json()["detail"]["code"] == "already_dismissed"
    assert (await c.post(f"{url}/transfer", json={"branch_id": None, "effective_date": str(date.today())})).status_code == 409

    r = await c.post(f"{url}/reinstate", json={"date": str(date.today()), "reactivate_account": True})
    assert r.json()["status"] == "active" and r.json()["user"]["is_active"] is True
    assert r.json()["dismissal_date"] is None
    kinds = [ev["kind"] for ev in (await c.get(f"{url}/events")).json()]
    assert kinds == ["reinstated", "dismissed", "hired"]


async def test_dismiss_can_keep_account_and_protects_last_superuser(new_client):
    root, root_id = await as_user(new_client, "root", "tasischi")
    e = await create(root)
    await root.post(f"/api/v1/employees/{e['id']}/account", json={"user_id": root_id})
    r = await root.post(f"/api/v1/employees/{e['id']}/dismiss", json={"date": str(date.today()), "reason": "own_will"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "cannot_deactivate_self"
    keep = await root.post(f"/api/v1/employees/{e['id']}/dismiss", json={"date": str(date.today()), "reason": "own_will", "block_account": False})
    assert keep.status_code == 200 and keep.json()["user"]["is_active"] is True


# --- Hisob yaratish/ulash ------------------------------------------------------------


async def test_create_with_account_and_account_rules(new_client):
    c = await hr(new_client)
    r = await c.post("/api/v1/employees", json=person(create_account=True, role_key="farmasevt"))
    body = r.json()
    assert r.status_code == 201 and body["login"] == "aliyev.vali" and body["temp_password"]
    assert body["employee"]["user"]["login"] == "aliyev.vali"
    emp_id = body["employee"]["id"]
    dup = await c.post(f"/api/v1/employees/{emp_id}/account", json={"role_key": "farmasevt"})
    assert dup.json()["detail"]["code"] == "account_exists"
    # Berish huquqi yo'q rol
    other = await create(c, last_name="Boshqa", phone="+998900009999")
    deny = await c.post(f"/api/v1/employees/{other['id']}/account", json={"role_key": "tasischi"})
    assert deny.status_code == 403 and deny.json()["detail"]["code"] == "cannot_assign_role"
    neither = await c.post(f"/api/v1/employees/{other['id']}/account", json={})
    assert neither.json()["detail"]["code"] == "role_or_user"
    # Mavjud foydalanuvchini ulash; bitta foydalanuvchi bitta xodimga
    uid = await make_user("solo", "mudir")
    assert (await c.post(f"/api/v1/employees/{other['id']}/account", json={"user_id": uid})).status_code == 200
    third = await create(c, last_name="Uchinchi", phone="+998900008888")
    r = await c.post(f"/api/v1/employees/{third['id']}/account", json={"user_id": uid})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "user_already_linked"
    assert (await c.delete(f"/api/v1/employees/{other['id']}/account")).status_code == 204
    assert (await c.get(f"/api/v1/employees/{other['id']}")).json()["user"] is None


# --- Ro'yxat -------------------------------------------------------------------------


async def test_list_filters_search_sort_and_summary(new_client):
    c = await hr(new_client)
    o = await org(c)
    a = await create(c, last_name="Aliyev", phone="+998901000001", department_id=o["dept"])
    b = await create(c, last_name="Valiyev", phone="+998901000002", department_id=o["dept2"], branch_id=o["branch"])
    await create(c, last_name="Zokirov", phone="+998901000003")
    await c.post(f"/api/v1/employees/{b['id']}/dismiss", json={"date": str(date.today()), "reason": "own_will"})

    q = lambda **kw: c.get("/api/v1/employees", params=kw)  # noqa: E731
    r = (await q()).json()
    assert r["total"] == 3 and r["summary"] == {"total": 3, "active": 2, "dismissed": 1, "with_account": 0}
    assert [e["last_name"] for e in r["items"]] == ["Aliyev", "Valiyev", "Zokirov"]
    assert [e["last_name"] for e in (await q(sort="full_name", desc="true")).json()["items"]][0] == "Zokirov"
    assert (await q(status="active")).json()["total"] == 2
    assert (await q(department_id=o["dept2"])).json()["total"] == 1
    assert (await q(branch_id=o["branch"])).json()["items"][0]["id"] == b["id"]
    assert (await q(q="valiyev")).json()["total"] == 1
    assert (await q(q="90100000")).json()["total"] == 3  # telefon bo'yicha qidirish
    assert (await q(q=f"№{a['employee_no']}")).json()["items"][0]["id"] == a["id"]
    assert (await q(has_account="true")).json()["total"] == 0
    page = (await q(size=2, page=2)).json()
    assert len(page["items"]) == 1 and page["total"] == 3


async def test_department_filter_includes_children(new_client):
    c = await hr(new_client)
    root = (await c.post("/api/v1/org/departments", json={"name": "Rahbariyat"})).json()["id"]
    child = (await c.post("/api/v1/org/departments", json={"name": "Moliya", "parent_id": root})).json()["id"]
    await create(c, department_id=child)
    assert (await c.get("/api/v1/employees", params={"department_id": root})).json()["total"] == 1
    assert (await c.get("/api/v1/employees", params={"department_id": root, "include_children": "false"})).json()["total"] == 0


async def test_org_objects_with_employees_cannot_be_deleted(new_client):
    c = await hr(new_client)
    o = await org(c)
    await create(c, department_id=o["dept"], position_id=o["pos"], branch_id=o["branch"])
    r = await c.delete(f"/api/v1/org/departments/{o['dept']}")
    assert r.status_code == 409 and r.json()["detail"]["code"] == "has_employees"
    assert (await c.delete(f"/api/v1/org/positions/{o['pos']}")).json()["detail"]["code"] == "has_employees"
    assert (await c.delete(f"/api/v1/branches/{o['branch']}")).json()["detail"]["code"] == "has_employees"


# --- Rasm va hujjatlar ---------------------------------------------------------------

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
PDF = b"%PDF-1.4\n" + b"x" * 64


async def test_photo_upload_download_and_validation(new_client, app):
    c = await hr(new_client)
    e = await create(c)
    ok = await c.put(f"/api/v1/employees/{e['id']}/photo", files={"file": ("me.png", PNG, "image/png")})
    assert ok.status_code == 204
    got = await c.get(f"/api/v1/employees/{e['id']}/photo")
    assert got.status_code == 200 and got.content == PNG and got.headers["content-type"] == "image/png"
    assert (await c.get(f"/api/v1/employees/{e['id']}")).json()["has_photo"] is True
    # PDF rasm sifatida qabul qilinmaydi; haqiqiy tur tekshiriladi (kengaytma/MIME aldamasin)
    bad = await c.put(f"/api/v1/employees/{e['id']}/photo", files={"file": ("x.png", PDF, "image/png")})
    assert bad.status_code == 415
    fake = await c.put(f"/api/v1/employees/{e['id']}/photo", files={"file": ("x.png", b"<html>not an image</html>", "image/png")})
    assert fake.status_code == 415
    # Almashtirilganda eski fayl o'chadi
    before = set(app.state.storage.objects)
    await c.put(f"/api/v1/employees/{e['id']}/photo", files={"file": ("n.png", PNG + b"1", "image/png")})
    assert len(set(app.state.storage.objects) - before) == 1 and len(app.state.storage.objects) == len(before)
    assert (await c.delete(f"/api/v1/employees/{e['id']}/photo")).status_code == 204
    assert (await c.get(f"/api/v1/employees/{e['id']}/photo")).status_code == 404
    assert app.state.storage.objects == {}


async def test_document_lifecycle(new_client, app):
    c = await hr(new_client)
    e = await create(c)
    url = f"/api/v1/employees/{e['id']}/documents"
    soon = str(date.today() + timedelta(days=10))
    r = await c.post(url, files={"file": ("Mening ish shartnomam.pdf", PDF, "application/pdf")},
                     data={"kind": "contract", "title": "Mehnat shartnomasi", "expires_at": soon})  # fmt: skip
    assert r.status_code == 201
    doc = r.json()
    assert doc["filename"] == "Mening ish shartnomam.pdf" and doc["size"] == len(PDF) and doc["expires_soon"] is True
    assert doc["mime"] == "application/pdf"

    listed = (await c.get(url)).json()
    assert [d["id"] for d in listed] == [doc["id"]]
    dl = await c.get(f"{url}/{doc['id']}/download")
    assert dl.content == PDF and "attachment" in dl.headers["content-disposition"]
    assert dl.headers["x-content-type-options"] == "nosniff"

    upd = await c.patch(f"{url}/{doc['id']}", json={"title": "Yangi nom", "expires_at": None})
    assert upd.json()["title"] == "Yangi nom" and upd.json()["expires_at"] is None
    assert (await c.patch(f"{url}/{doc['id']}", json={"kind": "xxx"})).status_code == 400

    assert (await c.delete(f"{url}/{doc['id']}")).status_code == 204
    assert app.state.storage.objects == {}
    assert (await c.get(f"{url}/{doc['id']}")).status_code == 404


async def test_document_upload_validation(new_client):
    c = await hr(new_client)
    e = await create(c)
    url = f"/api/v1/employees/{e['id']}/documents"
    post = lambda name, data, **form: c.post(url, files={"file": (name, data, "application/octet-stream")}, data={"kind": "other", "title": "T", **form})  # noqa: E731
    assert (await post("evil.exe", b"MZ\x90\x00" + b"\x00" * 50)).status_code == 415
    assert (await post("fake.pdf", b"not a pdf at all")).status_code == 415
    assert (await post("a.pdf", b"")).status_code == 400
    assert (await post("a.pdf", PDF, kind="nope")).status_code == 400
    assert (await post("a.pdf", PDF, issued_at="2026-02-01", expires_at="2026-01-01")).status_code == 400
    big = b"%PDF-" + b"0" * (15 * 1024 * 1024 + 10)
    assert (await post("big.pdf", big)).status_code == 413
    # Kengaytma magic ga mos kelmasa rad etiladi (pdf ichida png)
    assert (await post("sneaky.pdf", PNG)).status_code == 415


async def test_documents_require_documents_permission_and_belong_to_employee(new_client):
    c = await hr(new_client)
    a = await create(c)
    b = await create(c, last_name="Boshqa", phone="+998900000009")
    doc = (await c.post(f"/api/v1/employees/{a['id']}/documents", files={"file": ("a.pdf", PDF, "application/pdf")}, data={"kind": "other", "title": "T"})).json()
    # Boshqa xodimning URL i orqali hujjatga kirib bo'lmaydi
    assert (await c.get(f"/api/v1/employees/{b['id']}/documents/{doc['id']}/download")).status_code == 404
    d, _ = await as_user(new_client, "dir", "direktor")
    assert (await d.get(f"/api/v1/employees/{a['id']}/documents")).status_code == 403
    assert (await d.get(f"/api/v1/employees/{a['id']}/documents/{doc['id']}/download")).status_code == 403


# --- Ommaviy amallar -----------------------------------------------------------------


async def test_bulk_transfer_reports_per_employee(new_client):
    c = await hr(new_client)
    o = await org(c)
    a = await create(c, phone="+998901000001", department_id=o["dept"], position_id=o["pos"])
    b = await create(c, last_name="Boshqa", phone="+998901000002", department_id=o["dept2"], position_id=o["pos2"])
    c3 = await create(c, last_name="Uchinchi", phone="+998901000003")
    await c.post(f"/api/v1/employees/{c3['id']}/dismiss", json={"date": str(date.today()), "reason": "own_will"})

    r = await c.post("/api/v1/employees/bulk", json={
        "ids": [a["id"], b["id"], c3["id"], 99999], "action": "transfer",
        "params": {"department_id": o["dept2"], "position_id": o["pos2"], "effective_date": str(date.today())},
    })  # fmt: skip
    body = r.json()
    res = {x["id"]: x for x in body["results"]}
    assert body["succeeded"] == 1 and res[a["id"]]["ok"] is True
    assert res[b["id"]]["error"] == "no_changes"  # allaqachon shu yerda
    assert res[c3["id"]]["error"] == "employee_dismissed" and res[99999]["error"] == "not_found"
    assert (await c.get(f"/api/v1/employees/{a['id']}")).json()["department"]["name"] == "Logistika"


async def test_bulk_failure_does_not_poison_the_rest(new_client):
    c = await hr(new_client)
    ids = [(await create(c, last_name=f"E{i}", phone=f"+99890100000{i}", hire_date=str(date.today() - timedelta(days=3))))["id"] for i in range(3)]
    ids.insert(1, 99999)
    r = (await c.post("/api/v1/employees/bulk", json={"ids": ids, "action": "dismiss", "params": {"date": str(date.today()), "reason": "contract_end"}})).json()
    assert r["succeeded"] == 3
    assert (await c.get("/api/v1/employees", params={"status": "dismissed"})).json()["total"] == 3


async def test_bulk_create_accounts_returns_credentials(new_client):
    c = await hr(new_client)
    ids = [(await create(c, last_name=f"Xodim{i}", phone=f"+99890200000{i}"))["id"] for i in range(2)]
    r = (await c.post("/api/v1/employees/bulk", json={"ids": ids, "action": "create_accounts", "params": {"role_key": "farmasevt"}})).json()
    assert r["succeeded"] == 2 and all(x["login"] and x["temp_password"] for x in r["results"])
    again = (await c.post("/api/v1/employees/bulk", json={"ids": ids, "action": "create_accounts", "params": {"role_key": "farmasevt"}})).json()
    assert again["succeeded"] == 0 and {x["error"] for x in again["results"]} == {"account_exists"}
    no_role = await c.post("/api/v1/employees/bulk", json={"ids": ids, "action": "create_accounts", "params": {}})
    assert no_role.json()["detail"]["code"] == "role_required"


async def test_bulk_requires_permission_for_the_action(new_client):
    boss = await hr(new_client)
    e = await create(boss)
    roles = {r["key"]: r for r in (await boss.get("/api/v1/roles")).json()}
    await boss.put(f"/api/v1/roles/{roles['koordinator']['id']}/permissions", json={"permissions": ["employees.read", "employees.transfer"]})
    kor, _ = await as_user(new_client, "kor", "koordinator")
    r = await kor.post("/api/v1/employees/bulk", json={"ids": [e["id"]], "action": "dismiss", "params": {"date": str(date.today()), "reason": "own_will"}})
    assert r.status_code == 403
    assert (await kor.post("/api/v1/employees/bulk", json={"ids": [], "action": "transfer"})).status_code == 422


# --- Dublikat guruhlari va birlashtirish ---------------------------------------------


async def test_duplicate_groups_and_merge(new_client, app):
    c = await hr(new_client)
    o = await org(c)
    a = await create(c, phone="+998901234567", birth_date="1990-01-01", department_id=o["dept"], position_id=o["pos"])
    b = (await c.post("/api/v1/employees", json=person(last_name="Aliyev", first_name="Vali", phone="+998901234567", confirm_duplicate=True, pinfl="12345678901234", email="a@b.uz"))).json()["employee"]
    await create(c, last_name="Yolg'iz", phone="+998909999999")

    groups = (await c.get("/api/v1/employees/duplicates")).json()
    assert len(groups) == 1 and {m["id"] for m in groups[0]["members"]} == {a["id"], b["id"]}
    assert "phone" in groups[0]["reasons"]

    # b ga hujjat va hisob qo'shamiz
    await c.post(f"/api/v1/employees/{b['id']}/documents", files={"file": ("d.pdf", PDF, "application/pdf")}, data={"kind": "other", "title": "Hujjat"})
    await c.post(f"/api/v1/employees/{b['id']}/account", json={"role_key": "farmasevt"})

    r = await c.post("/api/v1/employees/merge", json={"primary_id": a["id"], "secondary_ids": [b["id"]]})
    assert r.status_code == 200
    m = r.json()
    assert m["id"] == a["id"] and m["pinfl"] == "12345678901234" and m["email"] == "a@b.uz"  # bo'sh maydonlar to'ldirildi
    assert m["user"] is not None  # hisob ko'chdi
    assert m["department"]["name"] == "Moliya"  # primaryniki saqlandi
    assert (await c.get(f"/api/v1/employees/{b['id']}")).status_code == 404
    assert len((await c.get(f"/api/v1/employees/{a['id']}/documents")).json()) == 1
    kinds = [e["kind"] for e in (await c.get(f"/api/v1/employees/{a['id']}/events")).json()]
    assert "merged" in kinds and kinds.count("hired") == 2  # ikkala tarix saqlandi
    assert (await c.get("/api/v1/employees/duplicates")).json() == []
    same = await c.post("/api/v1/employees/merge", json={"primary_id": a["id"], "secondary_ids": [a["id"]]})
    assert same.json()["detail"]["code"] == "merge_same"


async def test_merge_requires_permission(new_client):
    hrm, _ = await as_user(new_client, "hrm", "hr_menejer")  # employees.merge yo'q
    assert (await hrm.get("/api/v1/employees/duplicates")).status_code == 403
    assert (await hrm.post("/api/v1/employees/merge", json={"primary_id": 1, "secondary_ids": [2]})).status_code == 403


# --- Excel -----------------------------------------------------------------------------


def xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADER = ["Familiya", "Ism", "Tug'ilgan sana", "Telefon", "PINFL", "Pasport", "Bo'lim", "Lavozim", "Filial №", "Ishga olingan sana", "Noma'lum ustun"]


async def test_template_and_export(new_client):
    c = await hr(new_client)
    o = await org(c)
    await create(c, pinfl="12345678901234", birth_date="1990-03-15", department_id=o["dept"])
    tpl = await c.get("/api/v1/employees/import/template")
    wb = load_workbook(io.BytesIO(tpl.content))
    assert wb.worksheets[0].cell(1, 1).value == "Familiya" and wb.worksheets[0].max_row == 2

    exp = await c.get("/api/v1/employees/export")
    assert exp.status_code == 200 and "attachment" in exp.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(exp.content)).worksheets[0]
    headers = [cell.value for cell in ws[1]]
    assert "PINFL" in headers and ws.cell(2, 2).value == "Aliyev Vali"

    # Sezgir ustunlarsiz rol (direktor) PINFL ni olmaydi
    d, _ = await as_user(new_client, "dir", "direktor")
    ws2 = load_workbook(io.BytesIO((await d.get("/api/v1/employees/export")).content)).worksheets[0]
    assert "PINFL" not in [cell.value for cell in ws2[1]]
    kor, _ = await as_user(new_client, "k", "koordinator")
    assert (await kor.get("/api/v1/employees/export")).status_code == 403


async def test_import_preview_and_apply_flow(new_client):
    c = await hr(new_client)
    await org(c)  # Moliya (Xodim lavozimi bor), Filial 16
    await create(c, last_name="Mavjud", first_name="Odam", phone="+998905550000")
    rows = [
        HEADER,
        ["Karimov", "Anvar", "15.03.1990", "90 111 22 33", "11111111111111", "AA1111111", "Moliya", "Xodim", "16", "01.09.2026", "x"],
        ["Rahimova", "Dilfuza", "1992-07-01", "+998 90 222 33 44", None, None, "Logistika", "Xodim", "№16", None, None],
        ["", "Yo'q", None, "123", None, None, "Mavjud emas", "Xodim", "999", "31.02.2026", None],  # ko'p xato
        ["Mavjud", "Odam", None, "905550000", None, None, None, None, None, None, None],  # bazada bor (telefon)
        ["Karimov", "Anvar", "15.03.1990", "90 999 99 99", None, None, None, None, None, None, None],  # faylda takror
    ]
    prev = await c.post("/api/v1/employees/import/preview", files={"file": ("x.xlsx", xlsx(rows), "application/octet-stream")})
    assert prev.status_code == 200, prev.text
    p = prev.json()
    assert (p["total"], p["valid"], p["errors"], p["duplicates"]) == (5, 2, 1, 2)
    assert "noma'lum ustun" in p["unknown_columns"]
    err = next(i for i in p["issues"] if i["status"] == "error")
    codes = {(e["field"], e["code"]) for e in err["errors"]}
    assert {("name", "name_required"), ("phone", "phone_invalid"), ("department", "department_unknown"), ("branch", "branch_unknown"), ("hire_date", "date_invalid")} <= codes
    dup_rows = {i["row"]: i for i in p["issues"] if i["status"] == "duplicate"}
    assert dup_rows[4]["matches"][0]["reasons"] == ["phone"] and dup_rows[5]["duplicate_of_row"] == 1

    applied = await c.post("/api/v1/employees/import/apply", json={"token": p["token"]})
    assert applied.json() == {"created": 2, "skipped": 0}
    listed = (await c.get("/api/v1/employees", params={"q": "karimov"})).json()["items"]
    assert listed[0]["phone"] == "+998901112233" and listed[0]["department"]["name"] == "Moliya" and listed[0]["branch"]["number"] == 16
    assert (await c.get("/api/v1/employees")).json()["total"] == 3
    # Token bir martalik
    assert (await c.post("/api/v1/employees/import/apply", json={"token": p["token"]})).status_code == 410


async def test_import_csv_and_single_fio_column(new_client):
    c = await hr(new_client)
    csv_data = b"F.I.SH;Telefon;Jinsi\nTo'rayev Sardor Akmal o'g'li;901010101;Erkak\nSoliyeva Nigora;902020202;ayol\n"
    prev = (await c.post("/api/v1/employees/import/preview", files={"file": ("x.csv", csv_data, "text/csv")})).json()
    assert prev["valid"] == 2
    await c.post("/api/v1/employees/import/apply", json={"token": prev["token"]})
    emps = {e["last_name"]: e for e in (await c.get("/api/v1/employees")).json()["items"]}
    assert emps["To'rayev"]["first_name"] == "Sardor" and emps["To'rayev"]["middle_name"] == "Akmal o'g'li" and emps["To'rayev"]["gender"] == "m"
    assert emps["Soliyeva"]["gender"] == "f"


async def test_import_rejects_bad_files_and_foreign_tokens(new_client):
    c = await hr(new_client)
    post = lambda name, data: c.post("/api/v1/employees/import/preview", files={"file": (name, data, "application/octet-stream")})  # noqa: E731
    assert (await post("x.pdf", PDF)).status_code == 415
    assert (await post("x.xlsx", b"not a zip")).status_code == 400
    assert (await post("x.xlsx", xlsx([["Familiya", "Ism"]]))).status_code == 400  # faqat sarlavha
    prev = (await c.post("/api/v1/employees/import/preview", files={"file": ("x.csv", b"Familiya,Ism\nA,B\n", "text/csv")})).json()
    other, _ = await as_user(new_client, "hr2", "hr_direktor")
    r = await other.post("/api/v1/employees/import/apply", json={"token": prev["token"]})
    assert r.status_code == 403  # token boshqa foydalanuvchiniki
    kor, _ = await as_user(new_client, "k", "koordinator")
    assert (await kor.post("/api/v1/employees/import/preview", files={"file": ("x.csv", b"a\nb", "text/csv")})).status_code == 403


async def test_audit_trail_for_employee_actions(new_client):
    c = await hr(new_client)
    e = await create(c, pinfl="12345678901234")
    await c.patch(f"/api/v1/employees/{e['id']}", json={"pinfl": "98765432109876", "email": "n@e.uz"})
    items = (await c.get("/api/v1/audit", params={"action": "employee."})).json()["items"]
    upd = next(i for i in items if i["action"] == "employee.update")
    assert upd["data"]["pinfl"] == "***" and upd["data"]["email"] == "n@e.uz"  # PINFL jurnalga yozilmaydi
    assert "employee.create" in [i["action"] for i in items]
