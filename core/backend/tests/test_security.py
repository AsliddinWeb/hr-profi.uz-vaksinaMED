"""Xavfsizlik testlari (pentest checklist, docs/SECURITY.md): avtorizatsiya supurish, CSRF, cheklov, sarlavhalar, sozlamalar."""

import re

from app.config import Settings, get_settings
from app.hardening import check_production_config
from tests.conftest import PASSWORD, login, make_user

PUBLIC = {
    ("POST", "/api/auth/login"), ("POST", "/api/auth/refresh"), ("GET", "/api/health"), ("GET", "/api/health/ready"), ("GET", "/api/v1/meta"),
    ("POST", "/api/v1/calls/webhook"), ("GET", "/api/metrics"),
}  # fmt: skip


def routes(app):
    """OpenAPI sxemasidan (barcha ulangan routerlar) (METHOD, yo'l)."""
    for path, item in app.openapi()["paths"].items():
        if path.startswith("/api"):
            for m in item:
                if m.upper() in ("GET", "POST", "PUT", "PATCH", "DELETE"):
                    yield m.upper(), path


def concrete(path: str) -> str:
    return re.sub(r"\{[^}]+\}", "1", path)


async def test_every_non_public_endpoint_rejects_anonymous(app, new_client):
    """Har bir himoyalangan endpoint tokensiz 401/403 qaytarishi kerak (ruxsatsiz kirish yo'q)."""
    anon = new_client()  # cookie yo'q
    open_routes: list[str] = []
    checked = 0
    for method, path in sorted(set(routes(app))):
        if (method, path) in PUBLIC:
            continue
        r = await anon.request(method, concrete(path), headers={"X-Requested-With": "t"})
        checked += 1
        if r.status_code not in (401, 403):
            open_routes.append(f"{method} {path} -> {r.status_code}")
    assert checked > 300
    # ochiq bo'lishi ataylab: shaxsiy havola bilan ishlaydigan jamoat endpointlari (token URL ichida)
    allowed_open = [x for x in open_routes if "/public/letters" in x or "/public/leave" in x or "/public/apply" in x or "/canteen/public" in x or "/canteen/qr" in x]
    assert [x for x in open_routes if x not in allowed_open] == []


async def test_mutating_endpoints_require_csrf_header(app, new_client):
    bare = new_client()
    bare.headers.pop("X-Requested-With", None)
    bad: list[str] = []
    for method, path in sorted(set(routes(app))):
        if method in ("GET",) or (method, path) in PUBLIC and path == "/api/v1/calls/webhook":
            continue
        if path.startswith("/api/ext/"):  # tashqi API: cookie yo'q, X-API-Key bilan himoyalangan (kalitsiz 401)
            assert (await bare.request(method, concrete(path))).status_code == 401
            continue
        r = await bare.request(method, concrete(path))
        if r.status_code != 403 or r.json().get("detail", {}).get("code") != "csrf":
            bad.append(f"{method} {path} -> {r.status_code}")
    assert bad == []


async def test_security_headers_and_docs_disabled(client):
    r = await client.get("/api/v1/meta")
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["cache-control"] == "no-store"
    assert r.headers["referrer-policy"]
    assert (await client.get("/api/docs")).status_code == 404 and (await client.get("/api/openapi.json")).status_code == 404


async def test_auth_cookies_are_httponly_and_scoped(new_client):
    await make_user("ali", "farmasevt")
    c = new_client()
    r = await login(c, "ali")
    cookies = r.headers.get_list("set-cookie")
    assert len(cookies) == 2 and all("httponly" in x.lower() and "samesite=lax" in x.lower() for x in cookies)
    assert any("path=/api/auth" in x.lower() for x in cookies)  # refresh token faqat auth yo'lida


async def test_rate_limit_returns_429_and_exempts_health(client):
    s = get_settings()
    old = s.rate_limit_multiplier
    s.rate_limit_multiplier = 0.001  # api qoidasi: 3000 * 0.001 = 3 ta/daqiqa
    try:
        codes = [(await client.get("/api/v1/meta")).status_code for _ in range(6)]
        assert codes[:3] == [200, 200, 200] and codes[3:] == [429, 429, 429]
        r = await client.get("/api/v1/meta")
        assert r.json()["detail"]["code"] == "rate_limited" and int(r.headers["retry-after"]) >= 1
        assert (await client.get("/api/health")).status_code == 200  # sog'liq tekshiruvi cheklanmaydi
    finally:
        s.rate_limit_multiplier = old


async def test_rate_limit_fails_open_when_redis_down(app, client):
    class Broken:
        async def incr(self, *a, **k):
            raise ConnectionError("redis down")

    s = get_settings()
    old_redis, old = app.state.redis, s.rate_limit_multiplier
    app.state.redis, s.rate_limit_multiplier = Broken(), 0.001
    try:
        for _ in range(5):
            assert (await client.get("/api/v1/meta")).status_code == 200
    finally:
        app.state.redis, s.rate_limit_multiplier = old_redis, old


def test_production_config_guard():
    ok = Settings(env="production", secret_key="x" * 40, minio_root_password="m" * 20, database_url="postgresql+asyncpg://vaksina:S3cret@db:5432/vaksina")
    assert check_production_config(ok) == []
    weak = Settings(env="production", secret_key="change-me", minio_root_password="", database_url="postgresql+asyncpg://vaksina:vaksina@db:5432/vaksina", livekit_api_key="k", livekit_api_secret="short")
    assert len(check_production_config(weak)) == 4
    assert check_production_config(Settings(env="development", secret_key="change-me")) == []


async def test_metrics_internal_only_and_token(client):
    await client.get("/api/v1/meta")
    r = await client.get("/api/metrics")
    assert r.status_code == 200 and "vaksina_db_up 1" in r.text and "vaksina_info{" in r.text
    assert "vaksina_http_request_duration_seconds_bucket" in r.text and 'vaksina_http_requests_total{method="GET",route="/api/v1/meta",status="2xx"} 1' in r.text
    # gateway orqali (X-Forwarded-For bor) tashqaridan ko'rinmaydi
    assert (await client.get("/api/metrics", headers={"X-Forwarded-For": "1.2.3.4"})).status_code == 404
    s = get_settings()
    s.metrics_token = "tok123"
    try:
        assert (await client.get("/api/metrics")).status_code == 404
        assert (await client.get("/api/metrics", headers={"Authorization": "Bearer tok123", "X-Forwarded-For": "1.2.3.4"})).status_code == 200
    finally:
        s.metrics_token = ""


async def test_injection_strings_are_inert(new_client):
    await make_user("hr", "hr_direktor")
    c = new_client()
    assert (await login(c, "hr")).status_code == 200
    for payload in ("'; DROP TABLE users; --", "%' OR '1'='1", "\\\\", "<script>alert(1)</script>", "x" * 99):
        assert (await c.get("/api/v1/search", params={"q": payload})).status_code == 200
        assert (await c.get("/api/v1/employees", params={"q": payload})).status_code == 200
    assert (await c.get("/api/v1/employees")).status_code == 200  # jadvallar joyida
    r = await c.post("/api/auth/login", json={"login": "hr' OR '1'='1", "password": PASSWORD})
    assert r.status_code == 401


def test_spreadsheet_formula_injection_is_neutralised():
    import io

    from openpyxl import Workbook, load_workbook

    from app.services.xlsx_safe import csv_safe, harden

    wb = Workbook()
    ws = wb.active
    ws.append(["=HYPERLINK(\"http://evil\",\"x\")", "+1+1", "-2+3", "@SUM(1)", "Oddiy", 5, "-"])
    ws.append(["JAMI", "=SUM(B2:B9)"])
    buf = io.BytesIO()
    harden(wb)
    wb.save(buf)
    out = load_workbook(io.BytesIO(buf.getvalue())).active
    for c in ("A1", "B1", "C1", "D1", "G1"):
        assert out[c].data_type == "s", c  # matn, formula emas
    assert out["E1"].value == "Oddiy" and out["F1"].value == 5
    assert out["B2"].data_type == "f"  # o'zimizning SUM formulamiz saqlanadi
    assert csv_safe("=cmd|' /C calc'!A0") == "'=cmd|' /C calc'!A0" and csv_safe("Karimov") == "Karimov" and csv_safe(5) == 5


async def test_exports_neutralise_hostile_names(new_client):
    import io

    from openpyxl import load_workbook

    await make_user("hr", "hr_direktor")
    c = new_client()
    assert (await login(c, "hr")).status_code == 200
    r = await c.post("/api/v1/employees", json={"last_name": "=HYPERLINK(\"http://evil\")", "first_name": "+Aziz"})
    assert r.status_code == 201, r.text
    ex = await c.get("/api/v1/employees/export")
    assert ex.status_code == 200
    ws = load_workbook(io.BytesIO(ex.content)).active
    names = [row[1] for row in ws.iter_rows(min_row=2)]
    assert names and all(cell.data_type == "s" for cell in names)
