import io

import pytest
from openpyxl import load_workbook

from tests.test_answers import T0, A, as_user, day_req, setup

pytestmark = pytest.mark.skipif(
    "not __import__('os').environ.get('DATABASE_URL')", reason="DATABASE_URL kerak"
)


@pytest.fixture
def clock(monkeypatch):
    fn = lambda: T0  # noqa: E731
    monkeypatch.setattr("app.api.answers.now", fn)
    monkeypatch.setattr("app.api.attendance.now", fn)
    monkeypatch.setattr("app.services.attendance.now", fn)


async def make_flow(ctx):
    """3 ta so'rov: tasdiqlangan (kunlik), soatlik kutilayotgan, rad etilgan."""
    r1 = (await ctx["emp"].post(f"{A}/requests", json=day_req())).json()
    assert (await ctx["coord"].post(f"{A}/requests/{r1['id']}/decision", json={"action": "approve"})).status_code == 200
    r2 = await ctx["emp"].post(f"{A}/requests", json={"kind": "hourly", "reason_type": "sick", "start_date": "2026-11-10", "start_time": "10:00", "end_time": "12:30"})
    assert r2.status_code == 201, r2.text
    r3 = (await ctx["emp"].post(f"{A}/requests", json=day_req(start_date="2026-11-20", end_date="2026-11-20"))).json()
    assert (await ctx["coord"].post(f"{A}/requests/{r3['id']}/decision", json={"action": "reject", "note": "Kadr yetishmaydi"})).status_code == 200


async def test_report_summary_and_filters(new_client, clock):
    ctx = await setup(new_client)
    await make_flow(ctx)
    r = await ctx["hr"].get(f"{A}/report", params={"month": "2026-11"})
    assert r.status_code == 200, r.text
    d = r.json()
    k = d["kpis"]
    assert (k["total"], k["people"], k["approved"], k["pending_coordinator"], k["rejected"]) == (3, 1, 1, 1, 1)
    assert d["days"]["2026-11-04"] == 1 and d["days"]["2026-11-05"] == 1 and d["days"]["2026-11-10"] == 1 and "2026-11-20" not in d["days"]  # rad etilgan sanaga kirmaydi
    p = d["people"][0]
    assert p["employee"]["full_name"] == "Aliyev Ali" and p["count"] == 1 and p["days"] == 2 and p["hourly"] == 0
    assert d["branches"] == [{"id": ctx["branch"]["id"], "name": "F5"}]
    it = next(i for i in d["items"] if i["status"] == "approved")
    assert it["branch"]["name"] == "F5" and it["coordinator"] == "Coord" and it["days"] == 2
    assert [e["action"] for e in it["events"]] == ["submitted", "approved"] and it["events"][1]["actor"] == "Coord"
    # holat va qidiruv filtri
    assert (await ctx["hr"].get(f"{A}/report", params={"month": "2026-11", "status": "rejected"})).json()["total"] == 1
    assert (await ctx["hr"].get(f"{A}/report", params={"month": "2026-11", "status": "pending"})).json()["total"] == 1
    assert (await ctx["hr"].get(f"{A}/report", params={"month": "2026-11", "q": "Yo'q-xodim"})).json()["total"] == 0
    assert (await ctx["hr"].get(f"{A}/report", params={"month": "2026-12"})).json()["kpis"]["total"] == 0
    # soatlik xodim reytingda faqat tasdiqlanganda hisoblanadi
    pend = next(i for i in d["items"] if i["kind"] == "hourly")
    assert pend["minutes"] == 150


async def test_report_access_scope(new_client, clock):
    ctx = await setup(new_client)
    await make_flow(ctx)
    assert (await ctx["emp"].get(f"{A}/report", params={"month": "2026-11"})).status_code == 403  # oddiy xodim
    own = await ctx["coord"].get(f"{A}/report", params={"month": "2026-11"})  # koordinator o'z so'rovlarini ko'radi
    assert own.status_code == 200 and own.json()["kpis"]["total"] == 3
    other, _ = await as_user(new_client, "coord2", "koordinator")
    assert (await other.get(f"{A}/report", params={"month": "2026-11"})).json()["kpis"]["total"] == 0


async def test_report_exports(new_client, clock):
    ctx = await setup(new_client)
    await make_flow(ctx)
    pdf = await ctx["hr"].get(f"{A}/report/export", params={"month": "2026-11", "format": "pdf", "lang": "uz"})
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    pdf_ru = await ctx["hr"].get(f"{A}/report/export", params={"month": "2026-11", "format": "pdf", "lang": "ru", "status": "approved"})
    assert pdf_ru.status_code == 200 and pdf_ru.content.startswith(b"%PDF")
    xl = await ctx["hr"].get(f"{A}/report/export", params={"month": "2026-11", "format": "xlsx"})
    assert xl.status_code == 200
    wb = load_workbook(io.BytesIO(xl.content))
    ws = wb.worksheets[0]
    assert ws.max_row == 4 and ws["B2"].value == "Aliyev Ali"
    assert {c.value for c in ws[1]} >= {"Xodim", "Holat"}
    assert (await ctx["emp"].get(f"{A}/report/export", params={"month": "2026-11"})).status_code == 403
    assert (await ctx["hr"].get(f"{A}/report/export", params={"month": "2026-13"})).status_code == 422
