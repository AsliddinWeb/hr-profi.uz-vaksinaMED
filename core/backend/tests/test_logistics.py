from datetime import date, datetime, timedelta

import pytest

from app.db import SessionLocal
from app.models.employees import Employee
from app.services import logistics as svc
from app.timeutil import tz
from tests.conftest import login, make_user

pytestmark = pytest.mark.skipif(
    "not __import__('os').environ.get('DATABASE_URL')", reason="DATABASE_URL kerak"
)

L = "/api/v1/logistics"
NOW = datetime(2026, 11, 10, 9, 0, tzinfo=tz())
LAT, LNG = 41.30, 69.28


@pytest.fixture
def clock(monkeypatch):
    state = {"t": NOW}
    fn = lambda: state["t"]  # noqa: E731
    monkeypatch.setattr("app.api.logistics.now", fn)
    return state


async def as_user(new_client, name, role):
    uid = await make_user(name, role)
    c = new_client()
    assert (await login(c, name)).status_code == 200
    return c, uid


async def mk_emp(name, user_id=None):
    async with SessionLocal() as s:
        e = Employee(last_name=name, first_name="X", full_name=f"{name} X", user_id=user_id, status="active")
        s.add(e)
        await s.commit()
        return e.id


async def setup(new_client):
    hr, _ = await as_user(new_client, "hr", "hr_direktor")
    boss, _ = await as_user(new_client, "boss", "direktor")
    far, _ = await as_user(new_client, "far", "farmasevt")
    d1, d1_uid = await as_user(new_client, "driver1", "haydovchi")
    d2, d2_uid = await as_user(new_client, "driver2", "haydovchi")
    e1, e2 = await mk_emp("Haydovchi1", d1_uid), await mk_emp("Haydovchi2", d2_uid)
    for e in (e1, e2):
        assert (await hr.put(f"{L}/drivers/{e}", json={"license_no": f"AB{e}", "categories": "B", "license_expires": "2030-01-01"})).status_code == 200
    v = (await hr.post(f"{L}/vehicles", json={"plate": "01 a 123 bc", "name": "Damas", "kind": "van", "fuel_type": "gas", "fuel_norm": 10, "odometer_km": 50000})).json()
    return {"hr": hr, "boss": boss, "far": far, "d1": d1, "d2": d2, "e1": e1, "e2": e2, "veh": v["id"]}


async def mk_trip(ctx, driver="e1", **kw):
    r = await ctx["hr"].post(f"{L}/trips", json={"vehicle_id": ctx["veh"], "driver_id": ctx[driver], "purpose": "Filiallarga yetkazish", **kw})
    assert r.status_code == 201, r.text
    return r.json()


def pt(clock_t, dlat=0.0, mins=0, **kw):
    return {"lat": LAT + dlat, "lng": LNG, "speed": 40, "accuracy": 10, "at": (clock_t + timedelta(minutes=mins)).isoformat(), **kw}


def test_gps_units():
    from datetime import timezone

    t0 = datetime(2026, 1, 1, tzinfo=timezone.utc)  # noqa: UP017
    pts = [(41.30, 69.28, t0), (41.31, 69.28, t0 + timedelta(minutes=5)), (41.32, 69.28, t0 + timedelta(minutes=10))]
    assert float(svc.path_km(pts)) == 2.2  # 0.02 daraja ~ 2.224 km, 0.1 gacha yaxlitlanadi
    jump = [(41.30, 69.28, t0), (45.0, 69.28, t0 + timedelta(seconds=10)), (41.301, 69.28, t0 + timedelta(minutes=2))]
    assert float(svc.path_km(jump)) < 0.2  # imkonsiz "sakrash" hisobga olinmaydi
    assert svc.path_km([]) == 0
    assert svc.clean_point({"lat": 41.3, "lng": 69.2, "accuracy": 30}) and not svc.clean_point({"lat": 41.3, "lng": 69.2, "accuracy": 500}) and not svc.clean_point({"lat": 99, "lng": 0})
    assert float(svc.norm_liters(120, 10)) == 12.0 and float(svc.norm_liters(333.3, 8.5)) == 28.33
    today = date(2026, 11, 10)
    assert [svc.doc_stage(today + timedelta(days=d), today) for d in (-1, 0, 7, 8, 30, 31)] == [3, 2, 2, 1, 1, 0]


async def test_vehicle_document_driver_crud(new_client, clock):
    ctx = await setup(new_client)
    hr, boss = ctx["hr"], ctx["boss"]
    assert (await ctx["far"].get(f"{L}/vehicles")).status_code == 403
    assert (await ctx["d1"].get(f"{L}/vehicles")).status_code == 403
    assert (await boss.post(f"{L}/vehicles", json={"plate": "01 B 000 AA", "name": "X"})).status_code == 403
    v = (await boss.get(f"{L}/vehicles/{ctx['veh']}")).json()
    assert v["plate"] == "01 A 123 BC" and v["fuel_norm"] == 10.0  # katta harfga keltiriladi
    assert (await hr.post(f"{L}/vehicles", json={"plate": "01a123bc", "name": "Dubl"})).status_code == 201
    assert (await hr.post(f"{L}/vehicles", json={"plate": "01 a 123 bc", "name": "Dubl"})).json()["detail"]["code"] == "plate_exists"
    assert (await hr.post(f"{L}/vehicles", json={"plate": "01 C 111 AA", "name": "X", "fuel_norm": -1})).status_code == 422
    assert (await hr.put(f"{L}/vehicles/{ctx['veh']}", json={"plate": "01 A 123 BC", "name": "Damas-2", "fuel_norm": 11})).json()["name"] == "Damas-2"
    # hujjatlar
    r = await hr.post(f"{L}/vehicles/{ctx['veh']}/documents", json={"kind": "insurance", "number": "S-1", "expires_at": "2026-11-20"})
    assert r.status_code == 201 and r.json()["days_left"] == 10 and r.json()["stage"] == 1
    assert (await hr.post(f"{L}/vehicles/{ctx['veh']}/documents", json={"kind": "osago", "issued_at": "2026-05-01", "expires_at": "2026-04-01"})).json()["detail"]["code"] == "dates_invalid"
    await hr.post(f"{L}/vehicles/{ctx['veh']}/documents", json={"kind": "inspection", "expires_at": "2026-10-01"})
    lst = (await boss.get(f"{L}/vehicles")).json()
    assert lst[0]["docs"] == {"expired": 1, "soon": 1}
    did = r.json()["id"]
    assert (await hr.put(f"{L}/documents/{did}", json={"kind": "insurance", "expires_at": "2027-11-20"})).json()["stage"] == 0
    assert (await hr.delete(f"{L}/documents/{did}")).status_code == 204
    # haydovchi profili
    d = (await boss.get(f"{L}/drivers/{ctx['e1']}")).json()
    assert d["has_profile"] and d["license_no"] == f"AB{ctx['e1']}" and d["license_stage"] == 0
    assert len((await boss.get(f"{L}/drivers")).json()) == 2
    plain = await mk_emp("Oddiy")
    assert (await hr.post(f"{L}/trips", json={"vehicle_id": ctx["veh"], "driver_id": plain, "purpose": "Reys"})).json()["detail"]["code"] == "not_a_driver"
    for a in ("vehicle_create", "document_create", "document_update", "document_delete", "driver_save"):
        assert (await hr.get("/api/v1/audit", params={"action": f"logistics.{a}"})).json()["total"] >= 1, a


async def test_expiry_notifications_once(new_client, clock):
    from app.services.logistics import check_expiry

    ctx = await setup(new_client)
    hr = ctx["hr"]
    await hr.post(f"{L}/vehicles/{ctx['veh']}/documents", json={"kind": "osago", "number": "O-1", "expires_at": "2026-12-05"})
    await hr.put(f"{L}/drivers/{ctx['e1']}", json={"license_expires": "2026-11-12"})
    today = date(2026, 11, 10)
    async with SessionLocal() as s:
        # hujjat yaratilgan paytda muddati 25 kun (bosqich 1) belgilangan, shuning uchun u haqida xabar yo'q; guvohnoma 2 kun (bosqich 2)
        n, _ = await check_expiry(s, today)
        await s.commit()
        assert n == 1
        assert (await check_expiry(s, today))[0] == 0  # qayta xabar yo'q
        n, _ = await check_expiry(s, date(2026, 12, 4))  # hujjat: 1 kun qoldi (bosqich 2), guvohnoma: o'tgan (3)
        await s.commit()
        assert n == 2
        assert (await check_expiry(s, date(2026, 12, 6)))[0] == 1  # hujjat o'tdi
    kinds = [n["kind"] for n in (await hr.get("/api/v1/notifications")).json()["items"]]
    assert kinds.count("logistics.expiry") >= 3
    rep = (await ctx["boss"].get(f"{L}/reports/expiry", params={"days": 60})).json()["items"]
    assert [(r["type"], r["days_left"]) for r in rep] == [("license", 2), ("document", 25)]


async def test_trip_flow_gps_fuel_and_reports(new_client, clock):
    ctx = await setup(new_client)
    hr, boss, d1, d2 = ctx["hr"], ctx["boss"], ctx["d1"], ctx["d2"]
    t = await mk_trip(ctx, planned_start=NOW.isoformat())
    tid = t["id"]
    assert t["status"] == "planned"
    # begona haydovchi ko'rmaydi
    assert (await d2.get(f"{L}/trips/{tid}")).status_code == 404
    assert (await d2.post(f"{L}/trips/{tid}/start", json={"odometer": 50000})).status_code == 404
    assert (await d1.get(f"{L}/my")).json()["planned"][0]["id"] == tid
    # GPS faol reysdan oldin qabul qilinmaydi
    assert (await d1.post(f"{L}/track", json={"points": [pt(NOW)]})).json()["detail"]["code"] == "no_active_trip"
    assert (await d1.post(f"{L}/trips/{tid}/start", json={"odometer": 49000})).json()["detail"]["code"] == "odometer_low"
    r = await d1.post(f"{L}/trips/{tid}/start", json={"odometer": 50000})
    assert r.status_code == 200 and r.json()["status"] == "active"
    # bir mashinaga/haydovchiga ikkinchi faol reys yo'q
    t2 = await mk_trip(ctx, driver="e2")
    assert (await d2.post(f"{L}/trips/{t2['id']}/start", json={"odometer": 50000})).json()["detail"]["code"] == "trip_conflict"

    # GPS: yomon aniqlik, takror va kelajak nuqtalar tashlanadi
    clock["t"] = NOW + timedelta(minutes=30)
    pts = [pt(NOW, 0.0, 1), pt(NOW, 0.01, 10), pt(NOW, 0.02, 20), pt(NOW, 0.02, 20), pt(NOW, 0.5, 25, accuracy=900), pt(NOW, 0.0, 90)]
    r = await d1.post(f"{L}/track", json={"points": pts})
    assert r.json()["saved"] == 3 and r.json()["skipped"] == 3
    assert (await d1.post(f"{L}/track", json={"points": [pts[1]]})).json()["skipped"] == 1  # qayta yuborish ikkilanmaydi
    clock["t"] = NOW + timedelta(minutes=22)
    live = (await boss.get(f"{L}/live")).json()
    assert len(live) == 1 and live[0]["trip_id"] == tid and live[0]["lat"] == pytest.approx(LAT + 0.02) and live[0]["stale"] is False
    assert (await d1.get(f"{L}/live")).status_code == 403
    clock["t"] = NOW + timedelta(minutes=45)
    assert (await boss.get(f"{L}/live")).json()[0]["stale"] is True  # 25 daqiqa signal yo'q
    tr = (await boss.get(f"{L}/trips/{tid}/track")).json()
    assert tr["total"] == 3 and tr["points"][0]["lat"] == pytest.approx(LAT)

    # haydovchi o'z reysida yoqilg'i kiritadi; boshqa haydovchi (faol reysi yo'q) kirita olmaydi
    f = await d1.post(f"{L}/fuel", json={"liters": 15, "price_per_liter": 9000, "odometer_km": 50060, "station": "Mustaqillik AGZS"})
    assert f.status_code == 201 and f.json()["total"] == 135000.0 and f.json()["trip_id"] == tid
    assert (await ctx["d2"].post(f"{L}/fuel", json={"liters": 5})).json()["detail"]["code"] == "no_active_trip"
    assert (await ctx["far"].post(f"{L}/fuel", json={"liters": 5})).status_code == 403

    # yakunlash: odometr 50120 -> 120 km, me'yor 10 l/100 -> 12.00 l, GPS ~2.2 km
    assert (await d1.post(f"{L}/trips/{tid}/finish", json={"odometer": 49990})).json()["detail"]["code"] == "odometer_low"
    assert (await d2.post(f"{L}/trips/{tid}/finish", json={"odometer": 50120})).status_code == 404
    r = (await d1.post(f"{L}/trips/{tid}/finish", json={"odometer": 50120, "note": "Hammasi joyida"})).json()
    assert r["status"] == "completed" and r["distance_km"] == 120.0 and r["fuel_norm_l"] == 12.0 and r["gps_km"] == 2.2
    assert (await hr.get(f"{L}/vehicles/{ctx['veh']}")).json()["odometer_km"] == 50120
    assert (await d1.post(f"{L}/trips/{tid}/finish", json={"odometer": 50130})).status_code == 409
    assert (await boss.get(f"{L}/live")).json() == []

    # yoqilg'i hisoboti qo'lda hisob bilan mos: 15 l quyildi, me'yor 12 l, farq +3, amalda 12.5 l/100km
    clock["t"] = NOW + timedelta(hours=2)
    rep = (await boss.get(f"{L}/reports/fuel", params={"start": "2026-11-10", "end": "2026-11-10"})).json()
    row = rep["items"][0]
    assert (row["km"], row["norm_liters"], row["liters"], row["cost"], row["variance"], row["actual_per_100"]) == (120.0, 12.0, 15.0, 135000, 3.0, 12.5)
    assert rep["totals"]["variance"] == 3.0 and rep["totals"]["trips"] == 1
    import io

    from openpyxl import load_workbook

    resp = await boss.get(f"{L}/reports/fuel/export", params={"start": "2026-11-10", "end": "2026-11-10"})
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws[2][0].value == "01 A 123 BC" and ws[2][7].value == 15.0 and ws[3][0].value == "JAMI"
    # reyslar va GPS ish kuni
    tr_rep = (await boss.get(f"{L}/reports/trips", params={"start": "2026-11-01", "end": "2026-11-30"})).json()["items"]
    assert tr_rep[0]["driver"]["name"] == "Haydovchi1 X" and tr_rep[0]["completed"] == 1 and tr_rep[0]["km"] == 120.0 and tr_rep[0]["open"] == 0
    days = (await boss.get(f"{L}/reports/driver-days", params={"start": "2026-11-10", "end": "2026-11-10"})).json()["items"]
    assert len(days) == 1 and days[0]["points"] == 3 and days[0]["hours"] == 0.3 and days[0]["trips"] == 1 and days[0]["km"] == 120.0
    for a in ("trip_create", "trip_start", "trip_finish", "fuel_create"):
        assert (await hr.get("/api/v1/audit", params={"action": f"logistics.{a}"})).json()["total"] >= 1, a
    # mashina/haydovchi o'chmaydi
    assert (await hr.delete(f"{L}/vehicles/{ctx['veh']}")).json()["detail"]["code"] == "vehicle_in_use"


async def test_trip_edit_cancel_and_manager_fuel(new_client, clock):
    ctx = await setup(new_client)
    hr = ctx["hr"]
    t = await mk_trip(ctx)
    r = await hr.put(f"{L}/trips/{t['id']}", json={"vehicle_id": ctx["veh"], "driver_id": ctx["e2"], "purpose": "Boshqa reys", "destination": "Chilonzor"})
    assert r.json()["driver"]["id"] == ctx["e2"] and r.json()["destination"] == "Chilonzor"
    assert (await ctx["boss"].put(f"{L}/trips/{t['id']}", json={"vehicle_id": ctx["veh"], "driver_id": ctx["e2"], "purpose": "X yy"})).status_code == 403
    assert (await hr.post(f"{L}/trips/{t['id']}/cancel", json={"note": "x"})).status_code == 422
    r = await hr.post(f"{L}/trips/{t['id']}/cancel", json={"note": "Mashina ta'mirda"})
    assert r.json()["status"] == "cancelled"
    assert (await hr.put(f"{L}/trips/{t['id']}", json={"vehicle_id": ctx["veh"], "driver_id": ctx["e2"], "purpose": "Qayta reys"})).status_code == 409
    # boshqaruvchi yoqilg'ini mashinaga qo'lda kiritadi, tahrirlaydi, o'chiradi
    f = (await hr.post(f"{L}/fuel", json={"vehicle_id": ctx["veh"], "liters": 40, "price_per_liter": 5000, "fueled_at": "2026-11-05T10:00:00+05:00"})).json()
    assert f["total"] == 200000.0 and f["driver"] is None
    assert (await hr.post(f"{L}/fuel", json={"liters": 5})).json()["detail"]["code"] == "vehicle_required"
    assert (await hr.put(f"{L}/fuel/{f['id']}", json={"vehicle_id": ctx["veh"], "liters": 42, "price_per_liter": 5000, "fueled_at": "2026-11-05T10:00:00+05:00"})).json()["total"] == 210000.0
    assert (await ctx["boss"].get(f"{L}/fuel", params={"vehicle_id": ctx["veh"]})).json()["total"] == 1
    assert (await hr.delete(f"{L}/fuel/{f['id']}")).status_code == 204
    assert (await hr.delete(f"{L}/drivers/{ctx['e1']}")).status_code == 204
