import io
from datetime import date, datetime
from decimal import Decimal

import pytest

from app.db import SessionLocal
from app.models.employees import Employee
from app.models.payroll import PayScheme
from app.models.penalties import Penalty
from app.services import payroll as svc
from app.timeutil import tz
from tests.conftest import login, make_user

pytestmark = pytest.mark.skipif(
    "not __import__('os').environ.get('DATABASE_URL')", reason="DATABASE_URL kerak"
)

P = "/api/v1/payroll"
NOW = datetime(2026, 12, 5, 10, 0, tzinfo=tz())
D = Decimal


@pytest.fixture
def clock(monkeypatch):
    fn = lambda: NOW  # noqa: E731
    for mod in ("app.api.payroll", "app.api.attendance", "app.services.attendance"):
        monkeypatch.setattr(f"{mod}.now", fn)


async def as_user(new_client, name, role):
    uid = await make_user(name, role)
    c = new_client()
    assert (await login(c, name)).status_code == 200
    return c, uid


def scheme(**kw):
    base = {"id": 1, "effective_from": date(2026, 1, 1), "fixed_amount": D(5_000_000), "commission_percent": D(2), "plan_bonus": D(300_000), "kpi_amount": D(500_000)}
    return PayScheme(**{**base, **kw})


SETTINGS = {"tax_percent": "12", "inps_percent": "0.1", "paid_leave": True}


def days(present=21, absent=4, leave=0, extra=()):
    return [("present", True)] * present + [("absent", True)] * absent + [("leave", True)] * leave + [("off", False)] * 5 + list(extra)


def test_compute_hand_calculation():
    inp = svc.Inputs(sales=D(20_000_000), plan=D(18_000_000), kpi=D(80), bonus=D(100_000), advance=D(1_000_000), deduction=D(50_000))
    r = svc.compute_line(scheme(), days(), inp, D(150_000), SETTINGS)
    assert (r["plan_days"], r["paid_days"]) == (25, 21)
    assert r["fixed"] == 4_200_000 and r["commission"] == 400_000 and r["plan_bonus"] == 300_000 and r["kpi_bonus"] == 400_000 and r["bonus"] == 100_000
    assert r["gross"] == 5_400_000 and r["taxable"] == 5_200_000 and r["tax"] == 624_000 and r["inps"] == 5_200
    assert r["net"] == 4_570_800 and r["to_pay"] == 3_570_800 and r["details"]["warnings"] == []


def test_compute_edges():
    # reja bajarilmadi: reja bonus yo'q; KPI kiritilmadi: ogohlantirish
    r = svc.compute_line(scheme(), days(25, 0), svc.Inputs(sales=D(10_000_000), plan=D(18_000_000)), D(0), SETTINGS)
    assert r["plan_bonus"] == 0 and r["fixed"] == 5_000_000 and r["commission"] == 200_000 and "kpi_missing" in r["details"]["warnings"]
    # ruxsat to'lanmaydigan rejim
    r = svc.compute_line(scheme(kpi_amount=D(0)), days(20, 0, 5), svc.Inputs(), D(0), {**SETTINGS, "paid_leave": False})
    assert r["paid_days"] == 20 and r["fixed"] == 4_000_000
    r = svc.compute_line(scheme(kpi_amount=D(0)), days(20, 0, 5), svc.Inputs(), D(0), SETTINGS)
    assert r["paid_days"] == 25 and r["fixed"] == 5_000_000
    # yaxlitlash: 1_000_000 * 7 / 22 = 318181.818..
    r = svc.compute_line(scheme(fixed_amount=D(1_000_000), kpi_amount=D(0)), [("present", True)] * 7 + [("absent", True)] * 15, svc.Inputs(), D(0), {"tax_percent": "0", "inps_percent": "0", "paid_leave": True})
    assert r["fixed"] == 318_182 and r["to_pay"] == 318_182
    # jarima yalpidan katta: soliq bazasi 0, farq to'lovdan ayriladi
    r = svc.compute_line(scheme(fixed_amount=D(100_000), kpi_amount=D(0)), days(25, 0), svc.Inputs(advance=D(0)), D(250_000), SETTINGS)
    assert r["taxable"] == 0 and r["tax"] == 0 and r["to_pay"] == -150_000 and "negative" in r["details"]["warnings"]
    # reja kunlari yo'q
    r = svc.compute_line(scheme(), [("off", False)] * 30, svc.Inputs(), D(0), SETTINGS)
    assert r["fixed"] == 0 and "no_plan_days" in r["details"]["warnings"]
    # KPI 100 dan oshsa 100 deb olinadi
    r = svc.compute_line(scheme(), days(25, 0), svc.Inputs(kpi=D(150)), D(0), SETTINGS)
    assert r["kpi_bonus"] == 500_000


def test_scheme_resolution_and_period():
    a, b, c = scheme(id=1, effective_from=date(2026, 1, 1)), scheme(id=2, effective_from=date(2026, 11, 20)), scheme(id=3, effective_from=date(2027, 1, 1))
    assert svc.resolve_scheme([a, b, c], date(2026, 10, 31)).id == 1
    assert svc.resolve_scheme([a, b, c], date(2026, 11, 30)).id == 2
    assert svc.resolve_scheme([c], date(2026, 11, 30)) is None
    assert svc.month_bounds("2026-02") == (date(2026, 2, 1), date(2026, 2, 28)) and svc.month_bounds("2026-12")[1] == date(2026, 12, 31)


async def setup(new_client):
    hr, _ = await as_user(new_client, "hr", "hr_direktor")
    mgr, _ = await as_user(new_client, "menejer", "hr_menejer")
    boss, _ = await as_user(new_client, "boss", "direktor")
    emp, emp_uid = await as_user(new_client, "ali", "farmasevt")
    other, other_uid = await as_user(new_client, "vali", "farmasevt")
    shift = (await hr.post("/api/v1/schedule/shifts", json={"name": "K", "start_time": "09:00", "end_time": "18:00", "late_grace_minutes": 5})).json()
    await hr.post("/api/v1/schedule/groups", json={"name": "Ofis", "default_shift_id": shift["id"], "weekly_off_days": [6], "is_default": True})
    async with SessionLocal() as s:
        e = Employee(last_name="Aliyev", first_name="Ali", full_name="Aliyev Ali", user_id=emp_uid, status="active")
        o = Employee(last_name="Valiyev", first_name="Vali", full_name="Valiyev Vali", user_id=other_uid, status="active")
        s.add_all([e, o])
        await s.commit()
        ids = (e.id, e.employee_no, o.id)
    return {"hr": hr, "mgr": mgr, "boss": boss, "emp": emp, "other": other, "eid": ids[0], "eno": ids[1], "oid": ids[2], "shift": shift["id"]}


async def present_days(hr, ctx):
    """Noyabr 2026: Du-Sha (yakshanba dam). 2-7, 9-14, 16-21, 23-25 => 21 kun; 26, 27, 28, 30 kelmagan."""
    for d in [*range(2, 8), *range(9, 15), *range(16, 22), *range(23, 26)]:
        day = f"2026-11-{d:02d}"
        r = await hr.post("/api/v1/attendance/records", json={"employee_id": ctx["eid"], "work_date": day, "check_in_at": f"{day}T09:00:00+05:00", "check_out_at": f"{day}T18:00:00+05:00", "shift_id": ctx["shift"], "reason": "Test"})
        assert r.status_code == 201, r.text


async def mk_scheme(hr, ctx, **kw):
    body = {"employee_id": ctx["eid"], "effective_from": "2026-01-01", "fixed_amount": 5000000, "commission_percent": 2, "plan_bonus": 300000, "kpi_amount": 500000, **kw}
    r = await hr.post(f"{P}/schemes", json=body)
    assert r.status_code == 201, r.text
    return r.json()


async def item(c, ctx, kind, amount, period="2026-11", **kw):
    r = await c.post(f"{P}/items", json={"employee_id": ctx["eid"], "period": period, "kind": kind, "amount": amount, **kw})
    assert r.status_code == 201, r.text
    return r.json()


async def test_scheme_and_item_crud_permissions(new_client, clock):
    ctx = await setup(new_client)
    hr, mgr, boss = ctx["hr"], ctx["mgr"], ctx["boss"]
    assert (await ctx["emp"].get(f"{P}/schemes")).status_code == 403
    assert (await mgr.post(f"{P}/schemes", json={"employee_id": ctx["eid"], "effective_from": "2026-01-01", "fixed_amount": 1})).status_code == 403
    s = await mk_scheme(hr, ctx)
    assert (await hr.post(f"{P}/schemes", json={"employee_id": ctx["eid"], "effective_from": "2026-01-01"})).json()["detail"]["code"] == "scheme_exists"
    assert (await hr.post(f"{P}/schemes", json={"employee_id": 99999, "effective_from": "2026-01-01"})).json()["detail"]["code"] == "unknown_employee"
    assert (await hr.post(f"{P}/schemes", json={"employee_id": ctx["eid"], "effective_from": "2026-06-01", "fixed_amount": -5})).status_code == 422
    s2 = await mk_scheme(hr, ctx, effective_from="2026-11-20", fixed_amount=6000000)
    cur = (await boss.get(f"{P}/schemes")).json()
    assert cur["total"] == 1 and cur["items"][0]["id"] == s2["id"] and cur["items"][0]["fixed_amount"] == 6000000
    assert (await boss.get(f"{P}/schemes", params={"current": "false"})).json()["total"] == 2
    r = await hr.put(f"{P}/schemes/{s['id']}", json={**{k: s[k] for k in ("effective_from", "fixed_amount", "commission_percent", "plan_bonus", "kpi_amount")}, "employee_id": ctx["eid"], "fixed_amount": 5500000})
    assert r.json()["fixed_amount"] == 5500000
    assert (await hr.delete(f"{P}/schemes/{s2['id']}")).status_code == 204

    # kiritmalar: hr_menejer kirita oladi, direktor ko'radi xolos
    it = await item(mgr, ctx, "sales", 1000000)
    assert (await boss.post(f"{P}/items", json={"employee_id": ctx["eid"], "period": "2026-11", "kind": "bonus", "amount": 1})).status_code == 403
    assert (await mgr.post(f"{P}/items", json={"employee_id": ctx["eid"], "period": "2026-13", "kind": "bonus", "amount": 1})).json()["detail"]["code"] == "period_invalid"
    assert (await mgr.post(f"{P}/items", json={"employee_id": ctx["eid"], "period": "2026-11", "kind": "kpi", "amount": 101})).json()["detail"]["code"] == "kpi_range"
    assert (await mgr.post(f"{P}/items", json={"employee_id": ctx["eid"], "period": "2026-11", "kind": "bonus", "amount": 0})).json()["detail"]["code"] == "amount_invalid"
    assert (await mgr.put(f"{P}/items/{it['id']}", json={"employee_id": ctx["eid"], "period": "2026-11", "kind": "sales", "amount": 2000000})).json()["amount"] == 2000000
    assert (await boss.get(f"{P}/items", params={"period": "2026-11"})).json()["total"] == 1
    assert (await mgr.delete(f"{P}/items/{it['id']}")).status_code == 204
    for a in ("scheme_create", "scheme_update", "scheme_delete", "item_create", "item_update", "item_delete"):
        assert (await hr.get("/api/v1/audit", params={"action": f"payroll.{a}"})).json()["total"] >= 1, a


async def test_import_items(new_client, clock):
    ctx = await setup(new_client)
    mgr = ctx["mgr"]
    csv = f"xodim_no;tur;summa;izoh\n{ctx['eno']};sales;12 000 000;Noyabr\n{ctx['eno']};kpi;90\n"
    r = await mgr.post(f"{P}/items/import", json={"period": "2026-11", "csv": csv, "dry_run": True})
    assert r.json() == {"ok": True, "valid": 2, "errors": [], "imported": 0}
    assert (await mgr.get(f"{P}/items", params={"period": "2026-11"})).json()["total"] == 0
    bad = csv + "99999;bonus;5\n" + f"{ctx['eno']};nimadir;5\n{ctx['eno']};bonus;abc\n"
    r = await mgr.post(f"{P}/items/import", json={"period": "2026-11", "csv": bad, "dry_run": True})
    assert r.json()["ok"] is False and [e["code"] for e in r.json()["errors"]] == ["unknown_employee", "kind_invalid", "row_invalid"]
    assert (await mgr.post(f"{P}/items/import", json={"period": "2026-11", "csv": bad, "dry_run": False})).json()["detail"]["code"] == "import_invalid"
    r = await mgr.post(f"{P}/items/import", json={"period": "2026-11", "csv": csv, "dry_run": False})
    assert r.json()["imported"] == 2
    got = (await mgr.get(f"{P}/items", params={"period": "2026-11"})).json()["items"]
    assert sorted((i["kind"], i["amount"]) for i in got) == [("kpi", 90.0), ("sales", 12000000.0)]


async def test_run_matches_hand_calculation_and_lifecycle(new_client, clock):
    ctx = await setup(new_client)
    hr, mgr, boss = ctx["hr"], ctx["mgr"], ctx["boss"]
    await present_days(hr, ctx)
    await mk_scheme(hr, ctx)
    for kind, amount in (("sales", 20000000), ("plan", 18000000), ("kpi", 80), ("bonus", 100000), ("advance", 1000000), ("deduction", 50000)):
        await item(mgr, ctx, kind, amount)
    async with SessionLocal() as s:
        mk = lambda amt, day, st: Penalty(employee_id=ctx["eid"], incident_type="late", incident_date=date.fromisoformat(day), source="manual", step=1, label="x", amount=amt, status=st)  # noqa: E731
        s.add_all([mk(100000, "2026-11-10", "approved"), mk(50000, "2026-11-12", "approved"), mk(70000, "2026-11-13", "pending"), mk(999999, "2026-12-01", "approved"), mk(888888, "2026-11-14", "cancelled")])
        await s.commit()

    assert (await mgr.post(f"{P}/runs", json={"period": "2026-11"})).status_code == 403
    assert (await hr.post(f"{P}/runs", json={"period": "2027-05"})).json()["detail"]["code"] == "period_future"
    r = (await hr.post(f"{P}/runs", json={"period": "2026-11"})).json()
    rid = r["id"]
    assert r["status"] == "draft" and r["settings"]["tax_percent"] == "12"
    assert (await hr.post(f"{P}/runs", json={"period": "2026-11"})).status_code == 409
    assert (await hr.post(f"{P}/runs/{rid}/approve")).json()["detail"]["code"] == "not_calculated"
    res = (await hr.post(f"{P}/runs/{rid}/calculate")).json()
    assert res["lines"] == 1 and [m["name"] for m in res["missing_scheme"]] == ["Valiyev Vali"] and res["month_finished"] is True

    lines = (await boss.get(f"{P}/runs/{rid}/lines")).json()["items"]
    l = lines[0]  # noqa: E741
    assert (l["plan_days"], l["paid_days"]) == (25, 21)
    assert (l["fixed"], l["commission"], l["plan_bonus"], l["kpi_bonus"], l["bonus"]) == (4200000, 400000, 300000, 400000, 100000)
    assert (l["gross"], l["penalties"], l["deductions"], l["taxable"]) == (5400000, 150000, 50000, 5200000)
    assert (l["tax"], l["inps"], l["net"], l["advance"], l["to_pay"]) == (624000, 5200, 4570800, 1000000, 3570800)
    assert l["warnings"] == ["pending_penalties"]
    d = (await boss.get(f"{P}/runs/{rid}/lines/{l['id']}")).json()
    assert d["details"]["absent_days"] == 4 and d["details"]["scheme"]["fixed_amount"] == "5000000.00"
    run = (await boss.get(f"{P}/runs/{rid}")).json()
    assert run["totals"]["to_pay"] == 3570800 and run["totals"]["employees"] == 1 and run["flagged"] == 1 and run["missing_scheme"] == 1
    assert (await boss.get(f"{P}/runs/{rid}/lines", params={"flagged": "true"})).json()["total"] == 1

    # qayta hisoblash bir xil natija beradi (idempotent)
    await hr.post(f"{P}/runs/{rid}/calculate")
    assert (await hr.get(f"{P}/runs/{rid}")).json()["totals"]["to_pay"] == 3570800 and (await hr.get(f"{P}/runs/{rid}/lines")).json()["total"] == 1

    # xodim tasdiqdan oldin ko'rmaydi
    assert (await ctx["emp"].get(f"{P}/my")).json() == []
    # hisobdan keyin kiritma o'zgardi: tasdiqlab bo'lmaydi
    extra = await item(mgr, ctx, "bonus", 1)
    assert (await hr.post(f"{P}/runs/{rid}/approve")).json()["detail"]["code"] == "stale_calculation"
    await mgr.delete(f"{P}/items/{extra['id']}")
    await hr.post(f"{P}/runs/{rid}/calculate")
    assert (await hr.post(f"{P}/runs/{rid}/approve")).json()["status"] == "approved"
    # muhrlangan: kiritma, hisob, o'chirish yopiq
    assert (await mgr.post(f"{P}/items", json={"employee_id": ctx["eid"], "period": "2026-11", "kind": "bonus", "amount": 5})).json()["detail"]["code"] == "period_locked"
    assert (await hr.post(f"{P}/runs/{rid}/calculate")).json()["detail"]["code"] == "run_locked"
    assert (await hr.delete(f"{P}/runs/{rid}")).json()["detail"]["code"] == "run_locked"
    # sxema keyin o'zgarsa ham muhrlangan satr o'zgarmaydi
    await hr.post(f"{P}/schemes", json={"employee_id": ctx["eid"], "effective_from": "2026-11-25", "fixed_amount": 9000000})
    assert (await boss.get(f"{P}/runs/{rid}/lines")).json()["items"][0]["to_pay"] == 3570800
    # xodim o'z varaqasini ko'radi, boshqasinikini yo'q
    my = (await ctx["emp"].get(f"{P}/my")).json()
    assert len(my) == 1 and my[0]["to_pay"] == 3570800 and my[0]["period"] == "2026-11"
    assert (await ctx["emp"].get(f"{P}/my/{my[0]['id']}")).json()["details"]["scheme"]["id"]
    assert (await ctx["other"].get(f"{P}/my/{my[0]['id']}")).status_code == 404
    assert (await ctx["other"].get(f"{P}/runs/{rid}/lines")).status_code == 403
    assert "payroll.payslip" in [n["kind"] for n in (await ctx["emp"].get("/api/v1/notifications")).json()["items"]]
    # qayta ochish va to'lash
    assert (await hr.post(f"{P}/runs/{rid}/reopen")).json()["status"] == "draft"
    assert (await ctx["emp"].get(f"{P}/my")).json() == []
    await hr.post(f"{P}/runs/{rid}/calculate")
    await hr.post(f"{P}/runs/{rid}/approve")
    paid = (await hr.post(f"{P}/runs/{rid}/paid")).json()
    assert paid["status"] == "paid" and paid["paid_at"]
    assert (await hr.post(f"{P}/runs/{rid}/reopen")).json()["detail"]["code"] == "run_locked"
    # qayta hisoblashda yangi sxema (oklad 9 mln, komissiya/KPI yo'q) qo'llanadi: 7 560 000 + 100 000 - 200 000 = 7 460 000 baza
    assert (await hr.get(f"{P}/runs")).json()[0]["to_pay"] == 7_460_000 - 895_200 - 7_460 - 1_000_000
    for a in ("run_create", "calculate", "approve", "reopen", "paid"):
        assert (await hr.get("/api/v1/audit", params={"action": f"payroll.{a}"})).json()["total"] >= 1, a

    # eksport
    resp = await boss.get(f"{P}/runs/{rid}/export")
    assert resp.status_code == 200 and "spreadsheetml" in resp.headers["content-type"]
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(resp.content)).active
    row = [c.value for c in ws[2]]
    assert row[1] == "Aliyev Ali" and row[6] == 7_560_000 and row[-1] == 5_557_340 and ws[3][1].value == "JAMI"
    assert (await hr.get("/api/v1/audit", params={"action": "payroll.export"})).json()["total"] == 1
